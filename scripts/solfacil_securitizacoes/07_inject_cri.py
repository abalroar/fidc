# -*- coding: utf-8 -*-
import json, re, shutil, zipfile
from openpyxl import load_workbook
P="/home/user/fidc/outputs/solfacil/Solfacil_CRIs_Comparativo.xlsx"
C=json.load(open("/tmp/claude-0/-home-user-fidc/27891f9c-8b60-5562-8194-4e25a29cb16e/scratchpad/cache_cri.json"))
cached={}
for k,v in C.items():
    sh,coord=k.rsplit("!",1); cached[(sh,coord)]=v
wb=load_workbook(P); order=wb.sheetnames
faltando=[]
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith("=") and (ws.title,c.coordinate) not in cached:
                faltando.append((ws.title,c.coordinate,c.value))
assert not faltando, "fórmula sem valor em cache: %s" % faltando[:3]
del wb
CELL=r'(<c r="%s"(?: [^>]*)?>)(<f>[^<]*</f>)(\s*<v\s*/>|\s*<v>[^<]*</v>)?'
tmp=P+".tmp"; zin=zipfile.ZipFile(P); zout=zipfile.ZipFile(tmp,"w",zipfile.ZIP_DEFLATED); n=0
for item in zin.infolist():
    data=zin.read(item.filename)
    m=re.match(r"xl/worksheets/sheet(\d+)\.xml$", item.filename)
    if m:
        sheet=order[int(m.group(1))-1]; txt=data.decode("utf-8")
        for (sh,coord),v in cached.items():
            if sh!=sheet: continue
            if isinstance(v,str):
                def rep(mm,v=v):
                    tag=re.sub(r'\st="[^"]*"','',mm.group(1))[:-1]+' t="str">'
                    return tag+mm.group(2)+"<v>%s</v>"%v
            else:
                def rep(mm,v=v): return mm.group(1)+mm.group(2)+"<v>%s</v>"%repr(round(v,10))
            txt,k=re.subn(CELL%coord, rep, txt); n+=k
        data=txt.encode("utf-8")
    elif item.filename=="xl/workbook.xml":
        t=data.decode("utf-8"); t=re.sub(r"<calcPr[^>]*/>","",t)
        data=t.replace("</workbook>",'<calcPr calcId="0" fullCalcOnLoad="1"/></workbook>').encode("utf-8")
    zout.writestr(item.filename,data,zipfile.ZIP_DEFLATED)
zin.close(); zout.close(); shutil.move(tmp,P)
assert n==len(cached), "gravados %d de %d"%(n,len(cached))
print("fórmulas:",len(cached),"| cache gravado:",n)
