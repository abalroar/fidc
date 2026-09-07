# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "outputs/solfacil/Solfacil_Securitizacoes_Comparativo.xlsx"
AR = "Arial"; NAVY = "1F3864"; BLUE = "2E5C8A"; WHITE = "FFFFFF"
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = load_workbook(OUT)

C = ["174", "177"]
HDR = {"174": "CRI Vert 174ª emissão\n6 séries · mai/2026",
       "177": "CRI Vert 177ª emissão\n5 séries · jul-set/2026"}
D = {}
D["Emissora"] = {"174": "Vert Companhia Securitizadora\nCNPJ 25.005.683/0001-09 · registro CVM 680 (S2)",
                 "177": "Vert Companhia Securitizadora\nCNPJ 25.005.683/0001-09 · registro CVM 680 (S2)"}
D["Originadora / agente de cobrança"] = {"174": "Solfácil Energia Solar Tecnologia e Serviços Financeiros\nCNPJ 31.931.053/0001-50",
                                         "177": "Solfácil Energia Solar Tecnologia e Serviços Financeiros\nCNPJ 31.931.053/0001-50"}
D["Cedentes"] = {"174": "FIDC IS Green Solfácil VI e Solfácil\n(take-out da carteira do FIDC VI)",
                 "177": "Cedentes Fundos (FIDCs Solfácil) e Solfácil"}
D["Coordenadores"] = {"174": "XP Investimentos (líder)", "177": "XP Investimentos (líder) e Bradesco BBI"}
D["Agente fiduciário"] = {"174": "Oliveira Trust DTVM", "177": "Oliveira Trust DTVM"}
D["Custodiante / escriturador"] = {"174": "Vert DTVM", "177": "Vert DTVM"}
D["Registro CVM"] = {"174": "SRE/1197/2026 — registrado em 20/05/2026\nencerrada em 28/05/2026",
                     "177": "SRE/2316/2026 — registrado em 20/07/2026\nencerrada em 02/09/2026"}
D["Rito e público-alvo"] = {"174": "Rito automático, com bookbuilding\nInvestidor qualificado",
                            "177": "Rito automático, sem bookbuilding\nInvestidor profissional"}
D["Volume"] = {"174": "R$ 456,481 mm distribuídos publicamente\n+ R$ 14,118 mm da 6ª série em colocação privada\nTotal R$ 470,599 mm",
               "177": "R$ 627,647 mm distribuídos publicamente\n+ 5ª série em colocação privada (valor não divulgado)"}
D["Rating"] = {"174": "Moody's AAA (1ª série Super Sênior A)\ndemais séries conforme prospecto",
               "177": "Sem classificação de risco"}
D["Data de emissão"] = {"174": "20/05/2026", "177": "20/07/2026"}
D["Vencimento por série"] = {"174": "Super Sênior A e B: 20/05/2031 (1.826 dias)\nSênior: 22/05/2034\nMezanino e Subordinado: 20/05/2036",
                             "177": "Não divulgado (dispensa de prospecto e lâmina)"}
D["Estrutura de séries"] = {
 "174": "1ª Super Sênior A · 2ª Super Sênior B · 3ª Sênior\n4ª Mezanino · 5ª Subordinado\n6ª Subordinado Jr (privada, subscrita pela Solfácil)",
 "177": "1ª Sênior A · 2ª Sênior B\n3ª Mezanino I · 4ª Mezanino II\n5ª série privada (subordinada, Solfácil)"}
D["Remuneração por série"] = {
 "174": "Super Sênior A: pré 14,8064% a.a.\nSuper Sênior B: 104% do CDI\nSênior: pré 15,7760% a.a.\nMezanino: CDI + 5,50%\nSubordinado: CDI + 8,00%\n6ª série: atrelada ao CDI",
 "177": "Não divulgada (dispensa de prospecto e lâmina)"}
D["Subordinação mínima"] = {
 "174": "Razões de Cobertura apuradas diariamente\n(ativos com PDD ÷ saldo devedor acumulado):\nSuper Sênior 147,06% → 32,0% de subordinação\nSênior 120,48% → 17,0%\nMezanino 109,89% → 9,0%\nSubordinada 105,26% → 5,0%",
 "177": "Não divulgada"}
D["Tipo de amortização (como funciona)"] = {
 "174": "Duas mecânicas convivendo. As séries Super Sênior A e B têm\ncronograma de amortização estabelecida (Anexo I do Termo de\nSecuritização): pagamento mensal no dia 20, carência de 2 meses e,\na partir de 20/08/2026, amortização linear crescente (1/58, 1/57 …)\naté 20/05/2031. As séries Sênior, Mezanino e Subordinado não têm\ncronograma: amortizam apenas por Amortização Extraordinária\n— varredura do caixa disponível na cascata, condicionada às\nRazões de Cobertura.",
 "177": "Não divulgada"}
D["Carência"] = {"174": "Principal e juros: 2 meses para Super Sênior A/B e Sênior\n(1º pagamento em 20/08/2026); 3 meses para Mezanino e\nSubordinado (1º pagamento em 21/09/2026)",
                 "177": "Não divulgada"}
D["Datas de pagamento"] = {"174": "Mensais, dia 20 (ou dia útil seguinte)", "177": "Não divulgadas"}
D["Regime pro rata × sequencial"] = {
 "174": "Amortização Pro Rata da 1ª integralização até o 47º mês\n(inclusive). Amortização Sequencial a partir do 48º mês — ou antes,\nse ocorrer Evento de Desalavancagem. Havendo Evento de\nRealavancagem, volta ao Pro Rata apenas se a causa tiver sido\ndesalavancagem; a troca do 48º mês é definitiva.",
 "177": "Não divulgado"}
D["Gatilhos de amortização sequencial"] = {
 "174": "Evento de Desalavancagem:\n(i) Índice de Atraso de Estoque desenquadrado (> 15%) em 3 datas\nde verificação consecutivas;\n(ii) rebaixamento de 2 níveis do rating das séries Super Sênior/Sênior;\n(iii) não pagamento de remuneração ou amortização das 1ª e 2ª séries,\nnão sanado em 5 dias úteis;\n(iv) não divulgação do Relatório da Emissão;\n(v) desenquadramento das Razões de Cobertura em 2 Datas de\nPagamento consecutivas ou 4 alternadas em 12 meses.\nPassagem do 48º mês também dispara o sequencial, por prazo.",
 "177": "Não divulgados"}
D["Resgate antecipado obrigatório"] = {
 "174": "Quando amortizados 98% do valor nominal unitário da série e houver\ncaixa suficiente para o resgate integral", "177": "Não divulgado"}
D["Revolvência"] = {"174": "Não há", "177": "Não há"}
D["Regime fiduciário / patrimônio separado"] = {"174": "Sim", "177": "Sim"}
D["Ativo-lastro"] = {"174": "CCB pré-fixadas de financiamento solar originadas na\nPlataforma Solfácil, para aquisição de sistema solar em imóveis",
                     "177": "CCB pré-fixadas de financiamento solar originadas na\nPlataforma Solfácil"}
D["Critérios de elegibilidade"] = {
 "174": "Taxa pré-fixada, pagamento em reais, parcelas mensais sem balão\nPrazo máximo do recebível: 3.845 dias (≈126 meses)\nPrazo médio ponderado da carteira: máx. 2.000 dias (66 meses)\nCarência máxima da CCB: 185 dias\nIdade máxima do devedor PF: 71 anos\nPJ (Res. CMN 5.118): constituída há ≥ 2 anos\nValor presente máximo: PF R$ 350 mil · PJ R$ 700 mil\nConcentração por devedor: 0,15% do patrimônio separado,\ncaindo para 0,07% acima de 750.000 CRI integralizados\nTaxa de retorno da carteira ≥ Taxa Média Mínima de Retorno",
 "177": "Mesmo padrão da 174ª emissão (não publicado)"}
D["Índice de atraso (limite)"] = {"174": "Índice de Atraso de Estoque ≤ 15% do valor presente cedido\n(atrasos acima de 90 dias)", "177": "Não divulgado"}
D["Título verde"] = {"174": "Sim — Green Bond Principles (ICMA), com parecer de segunda opinião",
                     "177": "Sim — Green Bond Principles (ICMA), com parecer de segunda opinião"}
D["Distribuição alcançada"] = {"174": "2.505 pessoas físicas · 13 fundos · 4 instituições financeiras\n10 demais pessoas jurídicas",
                               "177": "7 fundos · 3 instituições financeiras"}
D["Observações de crédito"] = {
 "174": "É o take-out do FIDC Solfácil VI: a carteira do fundo foi cedida à\nsecuritizadora, e o PL do VI caiu de R$ 895,8 mm (set/25) para\nR$ 211,1 mm (jul/26). A estrutura tem seis camadas — mais\ngranular que qualquer FIDC da família — e move o funding do\nmercado profissional para o varejo qualificado (2.505 pessoas físicas).",
 "177": "Distribuída a investidores profissionais sob dispensa de prospecto e\nde lâmina (art. 9º, I e art. 23, §1º da Res. CVM 160), de modo que o\nTermo de Securitização não é público. As características estruturais\nsó podem ser confirmadas junto à emissora ou aos coordenadores."}

SER174 = [("1ª série — Super Sênior A", 103.870), ("2ª série — Super Sênior B", 225.550),
          ("3ª série — Sênior", 70.590), ("4ª série — Mezanino", 37.647),
          ("5ª série — Subordinado", 18.824), ("6ª série — Subordinado Jr (privada)", 14.118)]
SER177 = [("1ª série — Sênior A", 100.000), ("2ª série — Sênior B", 450.000),
          ("3ª série — Mezanino I", 51.765), ("4ª série — Mezanino II", 25.882),
          ("5ª série — privada (subordinada)", None), ("—", None)]

ws = wb.create_sheet("CRI")
ws["A1"] = "Securitizações Solfácil — CRI"
ws["A1"].font = Font(name=AR, size=14, bold=True, color=NAVY)
ws["A2"] = ("Certificados de recebíveis imobiliários lastreados em CCB originadas na Plataforma Solfácil. Fontes: lâmina e prospecto definitivo "
            "da 174ª emissão e anúncio de início da 177ª emissão (Sistema de Registro de Ofertas da CVM), e base de ofertas encerradas em dados.cvm.gov.br.")
ws["A2"].font = Font(name=AR, size=9, italic=True, color="595959")
ws.row_dimensions[1].height = 20; ws.row_dimensions[2].height = 26

SEC = "§"
rows = [
 (SEC, "IDENTIFICAÇÃO"),
 ("row", "Emissora"), ("row", "Originadora / agente de cobrança"), ("row", "Cedentes"),
 ("row", "Coordenadores"), ("row", "Agente fiduciário"), ("row", "Custodiante / escriturador"),
 ("row", "Registro CVM"), ("row", "Rito e público-alvo"), ("row", "Data de emissão"),
 ("row", "Volume"), ("row", "Rating"), ("row", "Vencimento por série"),
 (SEC, "ESTRUTURA DE CAPITAL"),
 ("row", "Estrutura de séries"), ("row", "Remuneração por série"),
 (SEC, "TAMANHO DAS SÉRIES (R$ mm)"),
 ("ser", 0), ("ser", 1), ("ser", 2), ("ser", 3), ("ser", 4), ("ser", 5),
 ("tot", "Total emitido (R$ mm)"),
 (SEC, "SUBORDINAÇÃO"),
 ("row", "Subordinação mínima"),
 (SEC, "AMORTIZAÇÃO"),
 ("row", "Tipo de amortização (como funciona)"), ("row", "Carência"), ("row", "Datas de pagamento"),
 ("row", "Regime pro rata × sequencial"), ("row", "Gatilhos de amortização sequencial"),
 ("row", "Resgate antecipado obrigatório"), ("row", "Revolvência"),
 (SEC, "LASTRO E ELEGIBILIDADE"),
 ("row", "Ativo-lastro"), ("row", "Critérios de elegibilidade"), ("row", "Índice de atraso (limite)"),
 ("row", "Regime fiduciário / patrimônio separado"), ("row", "Título verde"),
 (SEC, "DISTRIBUIÇÃO E LEITURA DE CRÉDITO"),
 ("row", "Distribuição alcançada"), ("row", "Observações de crédito"),
]

HR = 4
ws.cell(row=HR, column=1, value="Atributo")
for j, k in enumerate(C, start=2):
    ws.cell(row=HR, column=j, value=HDR[k])
for j in range(1, len(C) + 2):
    c = ws.cell(row=HR, column=j)
    c.font = Font(name=AR, size=10, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    c.border = BOX
ws.row_dimensions[HR].height = 34

r = HR + 1
ser_rows = []
tot_row = None
for kind, label in rows:
    if kind == SEC:
        ws.cell(row=r, column=1, value=label)
        for j in range(1, len(C) + 2):
            c = ws.cell(row=r, column=j)
            c.fill = PatternFill("solid", fgColor=BLUE)
            c.font = Font(name=AR, size=10, bold=True, color=WHITE)
            c.border = BOX
        ws.row_dimensions[r].height = 18
        r += 1; continue
    if kind == "ser":
        i = label
        ws.cell(row=r, column=1, value="%s  |  %s" % (SER174[i][0], SER177[i][0]))
        ws.cell(row=r, column=1).font = Font(name=AR, size=9, bold=True, color=NAVY)
        for j, (k, S) in enumerate(zip(C, [SER174, SER177]), start=2):
            c = ws.cell(row=r, column=j)
            if S[i][1] is None:
                c.value = "n.d." if S[i][0] != "—" else None
                c.font = Font(name=AR, size=9, italic=True)
                c.alignment = Alignment(horizontal="right")
            else:
                c.value = S[i][1]
                c.number_format = '#,##0.0;(#,##0.0);-'
                c.font = Font(name=AR, size=9, color="0000FF")
            c.border = BOX
        ws.cell(row=r, column=1).border = BOX
        ws.cell(row=r, column=1).alignment = Alignment(vertical="top", wrap_text=True)
        ser_rows.append(r); r += 1; continue
    if kind == "tot":
        ws.cell(row=r, column=1, value=label).font = Font(name=AR, size=9, bold=True, color=NAVY)
        ws.cell(row=r, column=1).border = BOX
        for j in range(2, len(C) + 2):
            col = get_column_letter(j)
            c = ws.cell(row=r, column=j, value="=SUM(%s%d:%s%d)" % (col, ser_rows[0], col, ser_rows[-1]))
            c.number_format = '#,##0.0;(#,##0.0);-'
            c.font = Font(name=AR, size=9, bold=True)
            c.alignment = Alignment(horizontal="right")
            c.border = BOX
        tot_row = r; r += 1; continue
    ws.cell(row=r, column=1, value=label).font = Font(name=AR, size=10, bold=True, color=NAVY)
    ws.cell(row=r, column=1).alignment = Alignment(vertical="top", wrap_text=True)
    ws.cell(row=r, column=1).border = BOX
    for j, k in enumerate(C, start=2):
        c = ws.cell(row=r, column=j, value=D[label][k])
        c.font = Font(name=AR, size=9)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = BOX
    r += 1

ws.column_dimensions["A"].width = 40
for j in range(2, len(C) + 2):
    ws.column_dimensions[get_column_letter(j)].width = 62
ws.freeze_panes = "B5"
ws.sheet_view.showGridLines = False

foot = r + 1
ws.cell(row=foot, column=1, value=(
 "Notas: (1) tamanhos das séries em azul são entradas, calculados como quantidade × R$ 1.000 de valor nominal unitário. "
 "(2) Na 177ª emissão a 5ª série é de colocação privada e sua quantidade não foi divulgada, de modo que o total soma apenas as séries públicas. "
 "(3) A subordinação da 174ª emissão está expressa como razão de cobertura; a subordinação equivalente em % é o complemento de 1 ÷ razão. "
 "(4) 'Não divulgado' na 177ª emissão decorre da dispensa de prospecto e de lâmina para ofertas a investidores profissionais."))
ws.cell(row=foot, column=1).font = Font(name=AR, size=8, italic=True, color="595959")
ws.cell(row=foot, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells(start_row=foot, start_column=1, end_row=foot, end_column=len(C) + 1)
ws.row_dimensions[foot].height = 46

wb.save(OUT)
print("CRI ok")
