# -*- coding: utf-8 -*-
"""LibreOffice is unavailable in this sandbox, so cached values are computed here
and injected into the saved workbook. Formulas are preserved verbatim."""
import re, shutil, zipfile, os
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

P = "outputs/solfacil/Solfacil_Securitizacoes_Comparativo.xlsx"

wb = load_workbook(P)
grid = {}   # (sheet, col, row) -> float
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, (int, float)):
                grid[(ws.title, c.column, c.row)] = float(c.value)

def val(sheet, ref):
    m = re.match(r"([A-Z]+)(\d+)$", ref)
    return grid.get((sheet, column_index_from_string(m.group(1)), int(m.group(2))), 0.0)

def evaluate(sheet, f):
    f = f.lstrip("=")
    m = re.match(r"SUM\(([A-Z]+\d+):([A-Z]+\d+)\)$", f)
    if m:
        c1, r1 = re.match(r"([A-Z]+)(\d+)", m.group(1)).groups()
        c2, r2 = re.match(r"([A-Z]+)(\d+)", m.group(2)).groups()
        tot = 0.0
        for ci in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
            for ri in range(int(r1), int(r2) + 1):
                tot += grid.get((sheet, ci, ri), 0.0)
        return tot
    m = re.match(r"IF\(([A-Z]+\d+)=0,0,\(([A-Z]+\d+)\+([A-Z]+\d+)\)/([A-Z]+\d+)\)$", f)
    if m:
        d = val(sheet, m.group(4))
        return 0.0 if d == 0 else (val(sheet, m.group(2)) + val(sheet, m.group(3))) / d
    m = re.match(r"IF\(([A-Z]+\d+)=0,0,([A-Z]+\d+)/([A-Z]+\d+)\)$", f)
    if m:
        d = val(sheet, m.group(3))
        return 0.0 if d == 0 else val(sheet, m.group(2)) / d
    raise ValueError("padrão de fórmula não previsto: " + f)

# passo 1: totais (SUM), passo 2: percentuais que dependem deles
cached = {}
for rounds in range(2):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    try:
                        v = evaluate(ws.title, c.value)
                    except ValueError:
                        raise
                    cached[(ws.title, c.coordinate)] = v
                    grid[(ws.title, c.column, c.row)] = v

order = wb.sheetnames
del wb

tmp = P + ".tmp"
zin = zipfile.ZipFile(P, "r")
zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
patched = 0
for item in zin.infolist():
    data = zin.read(item.filename)
    m = re.match(r"xl/worksheets/sheet(\d+)\.xml$", item.filename)
    if m:
        sheet = order[int(m.group(1)) - 1]
        txt = data.decode("utf-8")
        for (sh, coord), v in cached.items():
            if sh != sheet:
                continue
            pat = re.compile(r'(<c r="%s"[^>]*>)(<f>[^<]*</f>)(?!<v>)' % coord)
            txt, n = pat.subn(lambda mm: mm.group(1) + mm.group(2) + "<v>%r</v>" % round(v, 10), txt)
            patched += n
        data = txt.encode("utf-8")
    zout.writestr(item, data)
zin.close(); zout.close()
shutil.move(tmp, P)
print("fórmulas:", len(cached), "| valores injetados:", patched)
for k, v in sorted(cached.items()):
    print("  ", k, round(v, 4))
