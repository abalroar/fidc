"""Unify the portfolio document while preserving its native financial objects."""
from __future__ import annotations

from collections.abc import Iterable, Mapping, Sequence
from copy import deepcopy
from io import BytesIO
import re
from zipfile import ZIP_DEFLATED, ZipFile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from services.bba_deck import _set_cell_borders, _use_plain_table_style
from services.portfolio_presentation_context import PortfolioPresentationContext
from services.presentation_text import public_document_text
from services.pptx_merge import PptxMergeError


BLACK = "1F1F1F"
GRAY = "5D6369"
LIGHT = "F5F6F7"
ORANGE = "EC7000"
WHITE = "FFFFFF"
FONT = "Calibri"
_PROTECTED_PARTS = ("ppt/charts/", "ppt/embeddings/", "ppt/media/", "ppt/slideMasters/", "ppt/slideLayouts/", "ppt/notesMasters/", "ppt/theme/", "ppt/fonts/")
_SOURCE_LABELS = {
    "A": "CVM · Informe mensal",
    "B": "CVM · Informe mensal",
    "C": "CVM / Fundos.NET · Estrutura e emissões",
    "D": "Fundos.NET/B3 · Curadoria documental",
    "Capa": "CVM / Fundos.NET",
    "Referências": "CVM / Fundos.NET",
}
_BLOCK_LABELS = {"A": "Visão executiva", "B": "Histórico mensal", "C": "Estrutura e emissões", "D": "Curadoria documental"}


def _public_note(value: object) -> str:
    text = public_document_text(value)
    return re.sub(
        r"(?<![\w:/])(?:\.{1,2}/)?(?:data|reports)/[^\s|;]+",
        lambda match: match.group(0).rsplit("/", 1)[-1], text,
    )


def _font_size(shape) -> float:
    return max((size.pt for paragraph in shape.text_frame.paragraphs for size in [paragraph.font.size, *(run.font.size for run in paragraph.runs)] if size), default=0)


def _is_cover(slide, name: str) -> bool:
    if any(shape.has_chart or shape.has_table for shape in slide.shapes):
        return False
    return any(shape.has_text_frame and shape.text == name and shape.top / 914400 >= 1.4 and _font_size(shape) >= 22 for shape in slide.shapes)


def _text_box(slide, text: str, x: float, y: float, width: float, height: float, *, size: float = 11, color: str = BLACK, bold: bool = False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear(); frame.word_wrap = True; frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Inches(0)
    paragraph = frame.paragraphs[0]; paragraph.alignment = align
    paragraph.space_before = paragraph.space_after = Pt(0)
    run = paragraph.add_run(); run.text = text
    run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return shape


def _remove_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)


def _append_notes(slide, lines: Iterable[str]) -> None:
    current = _public_note(slide.notes_slide.notes_text_frame.text).strip()
    added = [line for line in lines if line]
    slide.notes_slide.notes_text_frame.text = "\n".join(([current] if current else []) + added)


def _fund_notes(funds: Iterable[object]) -> list[str]:
    notes = []
    for fund in funds:
        if isinstance(fund, Mapping):
            name = fund.get("name") or fund.get("display_name") or fund.get("short_name") or "Nome não informado"
            display = fund.get("display_name") or fund.get("short_name")
            cnpj = fund.get("cnpj") or "não informado"
        else:
            name = getattr(fund, "name", None) or getattr(fund, "display_name", "Nome não informado")
            display = getattr(fund, "display_name", None) or getattr(fund, "short_name", None)
            cnpj = getattr(fund, "cnpj", "não informado")
        alias = f" · Nome curto: {display}" if display and display != name else ""
        notes.append(f"{name} · CNPJ {cnpj}{alias}")
    return notes


def _blank_slide(prs):
    layout = next((layout for layout in prs.slide_layouts if layout.name.lower() == "blank"), prs.slide_layouts[0])
    return prs.slides.add_slide(layout)


def _compose_cover(slide, context: PortfolioPresentationContext, *, partial: bool) -> list[str]:
    previous = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.text]
    title = next((shape for shape in slide.shapes if shape.has_text_frame and shape.text == context.name), None)
    for shape in list(slide.shapes):
        if shape.has_text_frame and shape.text and (title is None or shape._element is not title._element):
            _remove_shape(shape)
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    if title is None:
        title = _text_box(slide, context.name, .95, 1.85, 11.35, 1.35, size=32, bold=True)
        marker = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(.60), Inches(1.85), Inches(.055), Inches(2.3))
        marker.fill.solid(); marker.fill.fore_color.rgb = RGBColor.from_string(ORANGE); marker.line.fill.background()
    size = 32 if len(context.name) <= 80 else 22 if len(context.name) <= 300 else 13
    for paragraph in title.text_frame.paragraphs:
        paragraph.font.name = FONT; paragraph.font.size = Pt(size); paragraph.font.color.rgb = RGBColor.from_string(BLACK)
        for run in paragraph.runs:
            run.font.name = FONT; run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(BLACK)
    title.text_frame.word_wrap = True
    # Reserve enough lines for long user names; never derive a date from them.
    capacity = max(1, int((title.width / 914400) * 72 / (size * .53)))
    lines = max(1, (len(context.name) + capacity - 1) // capacity)
    title.height = max(title.height, Inches(lines * size / 72 * 1.15))
    top = max(3.1, (title.top + title.height) / 914400 + .27)
    _text_box(slide, f"Carteira · {context.fund_count} fundos", title.left / 914400, top, 11.35, .24, size=11, color=GRAY)
    for index, line in enumerate((f"Posição comum · {context.common_position}", f"Histórico selecionado · {context.historical_window}", f"Consulta · {context.consulted_at}"), 1):
        _text_box(slide, line, title.left / 914400, top + .38 * index, 11.35, .28, size=11, color=GRAY)
    if partial:
        _text_box(slide, "Exportação parcial · Blocos indisponíveis e motivos nas referências e nas notas.", title.left / 914400, top + 1.70, 11.35, .4, size=11, color=GRAY)
    return previous


def _add_references(prs, context: PortfolioPresentationContext, *, included: Sequence[str], warnings: Sequence[str], compactions: Sequence[str], partial: bool):
    slide = _blank_slide(prs)
    _text_box(slide, "Referências da análise", .5, .25, 12.25, .5, size=22, bold=True)
    rows = [line.partition(": ")[::2] if ": " in line else ("Leitura", line) for line in context.reference_lines()]
    rows.append(("Unidades", "Unidades e denominadores conforme os rótulos originais de cada gráfico e tabela."))
    rows.append(("Blocos incluídos", ", ".join(_BLOCK_LABELS.get(name, name) for name in included)))
    missing = [name for name in _BLOCK_LABELS if name not in included]
    if missing:
        rows.append(("Blocos ausentes", ", ".join(_BLOCK_LABELS[name] for name in missing) + "; motivos disponíveis nas notas."))
    rows.append(("Apresentação", "Exportação parcial; motivos completos nas notas." if partial else "Blocos disponíveis incluídos nesta apresentação."))
    if compactions:
        rows.append(("Compactação", f"{len(compactions)} registros: repetições agrupadas ou componentes substituídos; detalhamento nas notas."))
    if warnings:
        rows.append(("Avisos", f"{len(warnings)} avisos de exportação; nenhum motivo foi descartado. Leitura integral nas notas."))
    graphic = slide.shapes.add_table(len(rows) + 1, 2, Inches(.5), Inches(1.10), Inches(12.25), Inches(.34 * (len(rows) + 1)))
    table = graphic.table; table.columns[0].width = Inches(2.25); table.columns[1].width = Inches(10)
    _use_plain_table_style(table)
    for row_index, row in enumerate([("Referência", "Contexto"), *rows]):
        for column, value in enumerate(row):
            cell = table.cell(row_index, column); cell.text = value
            cell.margin_left = cell.margin_right = Inches(.08); cell.margin_top = cell.margin_bottom = Inches(.04)
            cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(LIGHT if row_index == 0 else WHITE)
            _set_cell_borders(cell, bottom="D7DADD")
            for paragraph in cell.text_frame.paragraphs:
                paragraph.space_before = paragraph.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = FONT; run.font.size = Pt(10); run.font.bold = row_index == 0 or column == 0
                    run.font.color.rgb = RGBColor.from_string(BLACK)
        line_count = max((len(str(value)) + capacity - 1) // capacity for value, capacity in zip(row, (28, 133)))
        table.rows[row_index].height = Inches(max(.34, line_count * 10 / 72 * 1.2 + .08))
    graphic.height = sum(row.height for row in table.rows)
    _append_notes(slide, [*context.reference_lines(), "Compactação e substituições por bloco:", *compactions, "Avisos e omissões:", *(_public_note(warning) for warning in warnings)])
    return slide


def _restore_native_parts(original: bytes, finalized: bytes) -> bytes:
    """Keep native chart/workbook/media and visual master parts byte-for-byte."""
    with ZipFile(BytesIO(original)) as source, ZipFile(BytesIO(finalized)) as target:
        files = {item.filename: target.read(item) for item in target.infolist()}
        protected = {name for name in source.namelist() if name.startswith(_PROTECTED_PARTS)}
        for name in protected:
            files[name] = source.read(name)
        root = etree.fromstring(files["[Content_Types].xml"])
        existing = {(node.tag, node.get("PartName") or node.get("Extension")) for node in root}
        for node in etree.fromstring(source.read("[Content_Types].xml")):
            key = (node.tag, node.get("PartName") or node.get("Extension"))
            if key not in existing and (node.get("Extension") or str(node.get("PartName") or "").lstrip("/") in protected):
                root.append(deepcopy(node)); existing.add(key)
        files["[Content_Types].xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        output = BytesIO()
        with ZipFile(output, "w", compression=ZIP_DEFLATED) as package:
            for name, value in files.items():
                package.writestr(name, value)
        return output.getvalue()


def finalize_portfolio_pptx(
    payload: bytes, *, context: PortfolioPresentationContext,
    blocks: Sequence[tuple[str, int]], funds: Iterable[object], warnings: Sequence[str] = (),
) -> bytes:
    """Finalize only the portfolio flow; shared generators remain unchanged."""
    with ZipFile(BytesIO(payload)) as archive:
        master_parts = [node for node in etree.fromstring(archive.read("[Content_Types].xml")) if node.get("ContentType") == "application/vnd.openxmlformats-officedocument.presentationml.notesMaster+xml"]
        if len(master_parts) > 1:
            raise PptxMergeError("A apresentação contém vários notes masters; concatene cadeias compatíveis com o merger antes da finalização.")
    prs = Presentation(BytesIO(payload))
    if sum(count for _, count in blocks) != len(prs.slides) or any(count < 0 for _, count in blocks):
        raise ValueError("A contagem dos blocos não corresponde aos slides da apresentação.")
    records = []; offset = 0; first_by_block = {}
    compactions = []
    for name, count in blocks:
        slides = list(prs.slides)[offset:offset + count]
        if slides:
            first_by_block[name] = slides[0]
        for index, slide in enumerate(slides, offset + 1):
            records.append((slide, name))
            text = "\n".join([*(shape.text for shape in slide.shapes if shape.has_text_frame), slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""])
            heading = max((shape for shape in slide.shapes if shape.has_text_frame and shape.top / 914400 < 1.3), key=_font_size, default=None)
            identity = f" · {heading.text}" if heading is not None else ""
            for line in text.splitlines():
                if "Compactação:" in line or line.startswith(("RDB não localizado em ", "Colunas idênticas agrupadas;")):
                    compactions.append(f"Bloco {name}{identity}, slide original {index}: {_public_note(line)}")
        offset += count
    included = [name for name, count in blocks if count]
    partial = bool(warnings) or not {"A", "B", "C", "D"}.issubset(included)
    a_cover = first_by_block.get("A"); b_cover = first_by_block.get("B")
    a_cover = a_cover if a_cover is not None and _is_cover(a_cover, context.name) else None
    b_cover = b_cover if b_cover is not None and _is_cover(b_cover, context.name) else None
    main_cover = a_cover or b_cover
    if a_cover is not None and b_cover is not None:
        compactions.append("Bloco B · Compactação: capa repetida retirada; capa A preservada.")
    compactions = list(dict.fromkeys(compactions))
    # Add new parts before deleting a slide; python-pptx otherwise reuses an
    # occupied slide partname when there is a hole in the original numbering.
    references = _add_references(prs, context, included=included, warnings=warnings, compactions=compactions, partial=partial)
    records.append((references, "Referências"))
    if main_cover is None:
        main_cover = _blank_slide(prs)
        slide_id = prs.slides._sldIdLst[-1]
        prs.slides._sldIdLst.remove(slide_id); prs.slides._sldIdLst.insert(0, slide_id)
        records.insert(0, (main_cover, "Capa"))
    if a_cover is not None and b_cover is not None:
        previous = [shape.text for shape in b_cover.shapes if shape.has_text_frame and shape.text]
        original_notes = b_cover.notes_slide.notes_text_frame.text if b_cover.has_notes_slide else ""
        _append_notes(references, ["Compactação: capa repetida do bloco B retirada; capa A preservada.", "Texto original da capa B:", *previous, "Notas originais da capa B:", _public_note(original_notes)])
        records = [(slide, name) for slide, name in records if slide is not b_cover]
        slide_id = next(node for node in prs.slides._sldIdLst if prs.part.related_slide(node.rId) is b_cover)
        prs.part.drop_rel(slide_id.rId); prs.slides._sldIdLst.remove(slide_id)
    previous_cover = _compose_cover(main_cover, context, partial=partial)
    _append_notes(main_cover, [*context.reference_lines(), "Identificação integral da cesta:", *_fund_notes(funds), "Texto original da capa:", *(_public_note(line) for line in previous_cover), "Avisos de exportação:", *(_public_note(warning) for warning in warnings)])

    total = len(records)
    for page, (slide, block) in enumerate(records, 1):
        previous_footer = []
        for shape in list(slide.shapes):
            y = shape.top / 914400
            if shape.has_text_frame and y >= 6.90:
                if shape.text:
                    previous_footer.append(_public_note(shape.text))
                _remove_shape(shape)
            elif not shape.has_text_frame and not shape.has_chart and not shape.has_table and y >= 6.90 and shape.height / 914400 < .06:
                _remove_shape(shape)
            elif slide is not main_cover and block == "A" and shape.has_text_frame and shape.text == context.name and .5 <= y <= .75 and shape.height / 914400 <= .4:
                previous_footer.append(f"Identificação do cabeçalho: {context.name}"); _remove_shape(shape)
            elif slide is not main_cover and shape.has_text_frame and y < 1.3 and _font_size(shape) >= 17:
                for paragraph in shape.text_frame.paragraphs:
                    paragraph.font.name = FONT; paragraph.font.color.rgb = RGBColor.from_string(BLACK)
                    for run in paragraph.runs:
                        run.font.name = FONT; run.font.color.rgb = RGBColor.from_string(BLACK)
        if previous_footer:
            _append_notes(slide, ["Referências originais do bloco:", *previous_footer])
        elif slide.has_notes_slide:
            _append_notes(slide, ())
        source = _SOURCE_LABELS.get(block, _public_note(block))
        if block == "C" and slide.has_notes_slide:
            observed = re.search(r"(?m)^Posição CVM: ([^\n]+)$", slide.notes_slide.notes_text_frame.text)
            if observed:
                source = f"CVM · {observed.group(1)} / Fundos.NET"
        _text_box(slide, f"{source} · Toma Conta", .5, 7.26, 10.5, .16, size=9, color=GRAY)
        _text_box(slide, f"{page}/{total}", 11.75, 7.26, 1.05, .16, size=9, color=GRAY, align=PP_ALIGN.RIGHT)
    output = BytesIO(); prs.save(output)
    return _restore_native_parts(payload, output.getvalue())


__all__ = ["finalize_portfolio_pptx"]
