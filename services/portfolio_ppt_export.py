"""Native portfolio deck compositor. Inputs and financial rules belong to producers.

The renderer builds one presentation from a content plan. It does not restyle or
concatenate the old presentations. Tables, lines and sparklines remain editable.
"""

from __future__ import annotations

from dataclasses import dataclass, replace
from io import BytesIO
import math
import re
import textwrap
from typing import Any

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from services.portfolio_ppt_style import DEFAULT_STYLE, PortfolioPptStyle, series_color


def build_portfolio_pptx_bytes(*, style: PortfolioPptStyle = DEFAULT_STYLE,
                               presentation_context=None, export_warnings=(),
                               table_placement: str = "appendix", section_dividers: str = "auto",
                               **inputs: Any) -> bytes:
    """Compose the portfolio deck.

    ``table_placement``: ``"inline"`` puts each theme's series right after its
    charts (pre-read); ``"appendix"`` moves every series table to an appendix
    in the same theme order (live presentation).
    ``section_dividers``: ``"auto"`` adds dividers only to long decks,
    ``"always"`` or ``"never"`` force the choice.
    ``fund_annex`` (in ``inputs``): ``"auto"`` shows each key metric fund by
    fund (small multiples) up to ``FUND_ANNEX_MAX_FUNDS`` funds and one row
    per fund above; ``"metrics"``, ``"full"`` or ``"summary"`` force the choice.
    """
    from services.portfolio_ppt_content import build_portfolio_deck_spec

    reference_lines = ()
    if presentation_context is not None:
        inputs.setdefault("dashboard_summary_competencia", presentation_context.common_position)
        reference_lines = tuple(presentation_context.reference_lines())
    spec = build_portfolio_deck_spec(**inputs, table_placement=table_placement,
                                     reference_lines=reference_lines, export_warnings=tuple(export_warnings))
    if presentation_context is not None:
        spec = replace(spec, metadata=(*spec.metadata,
            f"Posição comum: {presentation_context.common_position}. Última observada: {presentation_context.latest_observed}."))
    if export_warnings:
        spec = replace(spec, metadata=(*spec.metadata, "Cobertura parcial. Partes indisponíveis identificadas em Referências e cobertura."))
    return render_portfolio_deck(spec, style=style, section_dividers=section_dividers)


def _number(value: object) -> float | None:
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except (TypeError, ValueError, OverflowError):
        return None


def _br(value: float, decimals: int) -> str:
    if round(value, decimals) == 0:
        value = 0.0  # never "-0,0"
    return f"{value:,.{decimals}f}".replace(",", "_").replace(".", ",").replace("_", ".")


def format_cell(value: object, token: str = "text") -> str:
    if value is None or value is pd.NA or value is pd.NaT:
        return "N/D"
    if not isinstance(value, str):
        try:
            if pd.isna(value):
                return "N/D"
        except (ValueError, TypeError):
            pass
    kind, _, places = token.partition(":")
    if kind in {"money", "money_value", "number", "percent", "integer"}:
        number = _number(value)
        if number is None:
            return "N/D" if str(value).strip() in {"", "nan", "None"} else str(value)
        decimals = int(places or ("0" if kind == "integer" else "2"))
        scaled = number / 1_000_000 if kind in {"money", "money_value"} else number
        if abs(round(scaled, decimals)) >= 1000:
            # Four or more integer digits: decimals add width, not information,
            # and push a 13-month table onto two slides.
            decimals = 0
        if kind == "money":
            return f"R$ {_br(number / 1_000_000, decimals)} mm"
        if kind == "money_value":
            return _br(number / 1_000_000, decimals)
        return _br(number, decimals) + ("%" if kind == "percent" else "")
    text = str(value).strip()
    return text if text and text not in {"nan", "None", "<NA>", "NaT"} else "N/D"


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#"))


def _text(slide, text: str, x: float, y: float, w: float, h: float, *, size: float,
          style: PortfolioPptStyle, bold: bool = False, color: str | None = None,
          align=PP_ALIGN.LEFT, wrap: bool = True):
    shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    shape.name = "Texto" if not bold else "Título"
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = wrap
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(str(text).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = paragraph.space_before = Pt(0)
        paragraph.line_spacing = 1.05
        run = paragraph.add_run()
        run.text = line
        run.font.name = style.font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _rgb(color or style.text)
    return shape


def _lines(text: object, width: float, size: float, *, bold: bool = False) -> list[str]:
    # Conservative fit, including legal names and long documentary clauses.
    # Bold labels (first column, headers) are wider, often in capitals.
    capacity = max(5, int((width - 12) / (size * (0.56 if bold else 0.49))))
    return [line for paragraph in str(text).split("\n")
            for line in (textwrap.wrap(paragraph, capacity, break_long_words=True,
                                      break_on_hyphens=False) or [""])]


def _header(prs, spec, style: PortfolioPptStyle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(style.white)
    _text(slide, spec.title, style.margin, style.title_top, style.content_width, 34,
          size=style.title_size, bold=True, style=style)
    _text(slide, spec.context, style.margin, style.context_top, style.content_width, 34,
          size=style.context_size, color=style.muted, style=style)
    source = str(spec.source or "CVM · Informe Mensal · Fundos.NET/B3")
    # The full source is retained in references and notes; this line is a reader aid.
    footer = source if len(source) <= 168 else "CVM · Informe Mensal · Fundos.NET/B3. Referências completas no anexo."
    _text(slide, footer.replace("\n", " "), style.margin, style.footer_top,
          style.content_width, 14, size=style.footer_size, color=style.muted, style=style, wrap=False)
    notes = [spec.context, source, *spec.notes]
    notes.extend(note for block in spec.blocks for note in getattr(block, "notes", ()))
    slide.notes_slide.notes_text_frame.text = "\n".join(dict.fromkeys(str(note) for note in notes if note))
    return slide


def _cover(prs, spec, style):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(style.white)
    title_lines = _lines(spec.title, style.content_width, 34)
    title_height = len(title_lines) * 39
    top = min(215, 300 - title_height)
    _text(slide, "\n".join(title_lines), style.margin, top, style.content_width,
          title_height, size=34, bold=True, style=style)
    _text(slide, spec.context, style.margin, top + title_height + 12, style.content_width,
          60, size=16, color=style.muted, style=style)
    metadata = spec.metadata if isinstance(spec.metadata, (tuple, list)) else (spec.metadata,)
    # No masthead, decorative shape or page counter.
    content = "\n".join(str(item) for item in metadata if item)
    lines = _lines(content, style.content_width, style.footer_size)
    height = len(lines) * 12
    _text(slide, "\n".join(lines), style.margin, style.footer_top - height + 12,
          style.content_width, height, size=style.footer_size, color=style.muted, style=style)
    slide.notes_slide.notes_text_frame.text = content


_NUMERIC_TOKENS = {"number", "money", "money_value", "percent", "integer"}


def _is_numeric_column(block, column) -> bool:
    token = block.formats.get(column, "text").split(":")[0]
    return (column in set(getattr(block, "numeric_columns", ())) or token in _NUMERIC_TOKENS
            or bool(re.fullmatch(r"[a-z]{3}/\d{2}", str(column))))


def _table_widths(block, width: float, style) -> list[float]:
    """Natural widths: about 56–110pt per numeric column and the label column
    sized by its text. Narrow tables are not stretched to the slide width;
    only wrapping text columns absorb free space."""
    columns = list(block.frame.columns)
    if len(columns) == 1:
        return [width]
    size = style.table_min_size if getattr(block, "compact", False) else style.table_size
    weights, wants = [], []
    for index, column in enumerate(columns):
        values = [str(column), *[format_cell(v, block.formats.get(column, "text")) for v in block.frame[column]]]
        length = max((len(value) for value in values), default=10)
        if index == 0:
            # Bold label column: same character width as the line estimate.
            natural = length * size * .56 + 14
            weights.append(min(225, max(110, natural)))
        elif _is_numeric_column(block, column):
            natural = length * size * .56 + 14
            weights.append(min(110, max(56, natural)))
        else:
            natural = length * size * .49 + 14
            weights.append(min(450, max(115, natural)))
        wants.append(max(0.0, natural - weights[-1]))
    total = sum(weights)
    if total > width:
        if all(_is_numeric_column(block, column) for column in columns[1:]):
            first = min(weights[0], width * .30)
            others = weights[1:]
            return [first, *[(width - first) * value / sum(others) for value in others]]
        return [width * value / total for value in weights]
    spare = width - total
    hungry = [index for index, want in enumerate(wants) if want > 0 and (index == 0 or not _is_numeric_column(block, columns[index]))]
    if hungry:
        need = sum(wants[index] for index in hungry)
        for index in hungry:
            weights[index] += spare * min(1.0, need / spare) * wants[index] / need
    return weights


def _table_layout(block, width: float, style, *, compact=False):
    compact = compact or getattr(block, "compact", False)
    size = style.table_min_size if compact or len(block.frame.columns) > 11 else style.table_size
    widths = _table_widths(block, width, style)
    header_height = max(24, max(len(_lines(col, w, size, bold=True)) for col, w in zip(block.frame.columns, widths)) * size * 1.2 + 8)
    heights = []
    for _, row in block.frame.iterrows():
        lines = max(len(_lines(format_cell(value, block.formats.get(col, "text")), w, size, bold=index == 0))
                    for index, (col, value, w) in enumerate(zip(block.frame.columns, row, widths)))
        # Line height ≈ 1.2 × size plus the 2pt cell margins, as rendered; a
        # lower estimate lets a long table run into the footer.
        height = max(16.5 if size <= 10 else 21, lines * size * 1.2 + 4.5)
        heights.append(height)
    return widths, header_height, heights, size


def _borders(cell, style, *, header=False):
    tcpr = cell._tc.get_or_add_tcPr()
    for name in ("lnL", "lnR", "lnT", "lnB"):
        for element in list(tcpr.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}" + name)):
            tcpr.remove(element)
        line = OxmlElement("a:" + name)
        line.set("w", "6350")  # 0.5pt, horizontal rule only.
        if name == "lnB" and not header:
            fill = OxmlElement("a:solidFill")
            rgb = OxmlElement("a:srgbClr")
            rgb.set("val", style.grid)
            fill.append(rgb)
        else:
            fill = OxmlElement("a:noFill")
        line.append(fill)
        tcpr.append(line)


def _table(slide, block, x, y, width, *, style, compact=False):
    widths, header_h, heights, size = _table_layout(block, width, style, compact=compact)
    width = sum(widths)
    if block.title:
        _text(slide, block.title, x, y, width, 19, size=style.section_size, bold=True, style=style)
        y += 22
    table = slide.shapes.add_table(len(block.frame) + 1, len(block.frame.columns),
                                   Pt(x), Pt(y), Pt(width), Pt(header_h + sum(heights))).table
    table.first_row = table.horz_banding = table.vert_banding = False
    for column, column_width in zip(table.columns, widths):
        column.width = Pt(column_width)
    for row, height in zip(table.rows, [header_h, *heights]):
        row.height = Pt(height)
    numeric = set(getattr(block, "numeric_columns", ()))
    colors = getattr(block, "cell_colors", {}) or {}
    values = [list(block.frame.columns), *block.frame.values.tolist()]
    for row_index, row_values in enumerate(values):
        for column_index, (column, value) in enumerate(zip(block.frame.columns, row_values)):
            cell = table.cell(row_index, column_index)
            token = block.formats.get(column, "text")
            cell.text = str(value) if row_index == 0 else format_cell(value, token)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(style.orange if row_index == 0 else style.white)
            cell.margin_left = cell.margin_right = Pt(6)
            cell.margin_top = cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.word_wrap = True
            cell.text_frame.auto_size = MSO_AUTO_SIZE.NONE
            is_numeric = column in numeric or token.split(":")[0] in {"money", "money_value", "number", "percent", "integer"}
            if row_index and is_numeric:
                cell.text_frame.word_wrap = False
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.RIGHT if column_index > 0 and is_numeric else PP_ALIGN.LEFT
                paragraph.space_before = paragraph.space_after = Pt(0)
                paragraph.line_spacing = 1.02
                for run in paragraph.runs:
                    run.font.name = style.font
                    run.font.size = Pt(size)
                    run.font.bold = row_index == 0 or column_index == 0 or bool(getattr(block, "snapshot", False) and column_index == 1)
                    color = colors.get((row_index - 1, column), style.text)
                    color = {"good": style.good, "bad": style.bad, "neutral": style.muted}.get(color, color)
                    run.font.color.rgb = _rgb(style.white if row_index == 0 else color)
            _borders(cell, style, header=row_index == 0)
    return header_h + sum(heights) + (22 if block.title else 0)


def _axis_bounds(values):
    lo, hi = min(values), max(values)
    raw = (hi - lo) / 5
    power = 10 ** math.floor(math.log10(raw))
    step = next(factor * power for factor in (1, 2, 2.5, 5, 10) if factor * power >= raw)
    lower, upper = math.floor(lo / step) * step, math.ceil(hi / step) * step
    if math.isclose(upper, hi):
        upper += step
    if math.isclose(lower, lo) and lo != 0:
        lower -= step
    if lo >= 0:
        lower = max(0, lower)
    decimals = max(0, -math.floor(math.log10(step)))
    if not math.isclose(step * 10 ** decimals, round(step * 10 ** decimals)):
        decimals += 1
    return lower, upper, step, decimals


def _chart_layout(chart, x, y, w, h):
    area = chart._chartSpace.xpath(".//c:plotArea")[0]
    layouts = area.xpath("./c:layout")
    layout = layouts[0] if layouts else OxmlElement("c:layout")
    if not layouts:
        area.insert(0, layout)
    manual = OxmlElement("c:manualLayout")
    for name, value in (("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"),
                        ("wMode", "factor"), ("hMode", "factor"),
                        ("x", x), ("y", y), ("w", w), ("h", h)):
        element = OxmlElement("c:" + name)
        element.set("val", str(value))
        manual.append(element)
    layout.append(manual)


def _chart_gaps(chart):
    main = chart._chartSpace.xpath("./c:chart")[0]
    elements = main.xpath("./c:dispBlanksAs")
    gaps = elements[0] if elements else OxmlElement("c:dispBlanksAs")
    gaps.set("val", "gap")
    if not elements:
        main.append(gaps)


def _axis_cap(values):
    """Cap the axis when one or two spikes flatten the rest of the series
    (e.g. a coverage ratio over a near-zero denominator). Outliers lie beyond
    the Tukey far-out fence (3×IQR) of the remaining values and are at least
    1.5× the highest of them. The fence is measured without the candidates,
    so two spikes in a short series cannot hide each other."""
    finite = sorted(v for v in values if v is not None)
    if len(finite) < 8:
        return None
    for count in (2, 1):
        regular, outliers = finite[:-count], finite[-count:]
        q1, q3 = regular[len(regular) // 4], regular[(3 * len(regular)) // 4]
        fence = q3 + 3 * (q3 - q1)
        if regular[-1] > 0 and outliers[0] > fence and outliers[0] >= 1.5 * regular[-1]:
            _, upper, _, _ = _axis_bounds([regular[0], regular[-1] * 1.1])
            return upper
    return None


def _chart_notes(block):
    """Sentences printed above the plot: constant series, empty series and a
    capped axis. They are computed before drawing so paired panels can share
    one band and start their plots at the same height."""
    series = [(s, [_number(v) for v in s.values]) for s in block.series]
    variable = [(s, vals) for s, vals in series
                if s.reference or len({_br(v, s.decimals) for v in vals if v is not None}) > 1
                or (block.kind == "bar" and any(v is not None for v in vals))]
    lines = []
    for spec, vals in series:
        if spec.reference or (spec, vals) in variable:
            continue
        valid = [v for v in vals if v is not None]
        suffix = "%" if spec.unit == "%" else " " + spec.unit
        if not valid:
            lines.append(f"Sem dados: {spec.label}.")
        else:
            qualifier = ("1 observação válida" if len(valid) == 1 else
                         "constante nas observações disponíveis" if any(v is None for v in vals) else "constante no período")
            lines.append(f"{spec.label}: {_br(valid[0], spec.decimals)}{suffix}, {qualifier}.")
    cap = None
    if block.kind == "line":
        drawn = [v for s, vals in variable if not s.reference for v in vals if v is not None]
        cap = _axis_cap(drawn)
        if cap is not None:
            peaks = [(v, block.categories[i], s) for s, vals in variable if not s.reference
                     for i, v in enumerate(vals) if v is not None and v > cap]
            text = "; ".join(f"{_br(v, s.decimals)}{'%' if s.unit == '%' else ''} em {month}" for v, month, s in peaks)
            places = 0 if cap >= 10 else 1 if cap >= 1 else 2
            lines.append(f"Eixo limitado a {_br(cap, places)}{'%' if block.unit == '%' else ''} para legibilidade; pico de {text}.")
    return variable, lines, cap


def _band_height(lines, width):
    return len(_lines("\n".join(lines), width, 10)) * 13 + 4 if lines else 0


def _kpi(slide, block, x, y, width, height, style):
    _text(slide, block.title, x, y, width, 22, size=style.section_size, bold=True, style=style)
    lines = _lines(block.text, width, 13)
    _text(slide, "\n".join(lines), x, y + 30, width, min(height - 30, len(lines) * 18),
          size=13, color=style.muted, style=style)


def _chart(slide, block, x, y, width, height, style, *, band=None, compact=False):
    """``compact``: a small multiple, with smaller type and narrower lanes."""
    _text(slide, block.title, x, y, width, 22, size=11 if compact else style.section_size, bold=True, style=style)
    variable, lines, cap = _chart_notes(block)
    own_band = _band_height(lines, width)
    if lines:
        _text(slide, "\n".join(_lines("\n".join(lines), width, 10)), x, y + 24, width, own_band,
              size=10, color=style.muted, style=style)
    band = own_band if band is None else max(band, own_band)
    if not [s for s, _ in variable if not s.reference]:
        return
    chart_y = y + 26 + band
    chart_h = height - 26 - band
    if block.kind == "bar":
        _bar_chart(slide, block, variable, x, chart_y, width, chart_h, style)
        return
    data = CategoryChartData()
    # Hierarchical native categories keep month positions tied to the data and
    # show each year once on the second axis row.
    parsed_categories = [str(label).split("/") for label in block.categories]
    hierarchical = bool(parsed_categories) and all(len(parts) == 2 for parts in parsed_categories)
    if hierarchical:
        previous_year = None
        year_category = None
        for month, year in parsed_categories:
            year = "20" + year if len(year) == 2 else year
            if year != previous_year:
                year_category = data.categories.add_category(year)
                previous_year = year
            year_category.add_sub_category(month)
    else:
        data.categories = list(block.categories)
    for spec, vals in variable:
        data.add_series(spec.label, vals)
    # The plot fills the column; only a narrow lane on the right is kept for
    # the end values.
    chart = slide.shapes.add_chart(XL_CHART_TYPE.LINE, Pt(x), Pt(chart_y), Pt(width), Pt(chart_h), data).chart
    chart.has_title = False
    drawn = [s for s, _ in variable]
    chart.has_legend = len(drawn) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(10)
        chart.legend.font.name = style.font
        chart.legend.font.color.rgb = _rgb(style.muted)
    type_size = 9 if compact else style.chart_size
    chart.font.name, chart.font.size = style.font, Pt(type_size)
    values = [v for s, vals in variable if not s.reference for v in vals if v is not None]
    references = [v for s, vals in variable if s.reference for v in vals if v is not None]
    if cap is not None:
        values = [v for v in values if v <= cap]
    if references and cap is None:
        values = [*values, *references]
    if cap is None and getattr(block, "axis_range", ()):
        values = [*values, *block.axis_range]
    lo, hi, major_unit, tick_decimals = _axis_bounds(values)
    if cap is not None:
        hi = max(hi, cap)
    axis = chart.value_axis
    axis.minimum_scale, axis.maximum_scale = lo, hi
    axis.major_unit = major_unit
    axis.tick_labels.font.name = style.font
    axis.tick_labels.font.size = Pt(type_size)
    axis.tick_labels.font.color.rgb = _rgb(style.muted)
    number_format = "[$-416]#,##0" + ("." + "0" * tick_decimals if tick_decimals else "")
    axis.tick_labels.number_format = number_format + ('"%"' if block.unit == "%" else "")
    axis.tick_labels.number_format_is_linked = False
    axis.has_major_gridlines = True
    axis.major_gridlines.format.line.color.rgb = _rgb(style.grid)
    axis.major_gridlines.format.line.width = Pt(.5)
    axis.format.line.fill.background()
    axis.major_tick_mark = axis.minor_tick_mark = XL_TICK_MARK.NONE
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    chart.category_axis.tick_labels.font.name = style.font
    chart.category_axis.tick_labels.font.size = Pt(8 if compact else 10)
    chart.category_axis.tick_labels.font.color.rgb = _rgb(style.muted)
    chart.category_axis.tick_label_spacing = max(1, math.ceil(len(block.categories) / (18 if width > 600 else 13 if width > 320 else 5)))
    chart.category_axis.format.line.fill.background()
    chart.category_axis.major_tick_mark = chart.category_axis.minor_tick_mark = XL_TICK_MARK.NONE
    _chart_gaps(chart)
    label_lane, left = (40.0, 34.0) if compact else (52.0, 42.0)
    top = 30.0 if chart.has_legend else 6.0
    _chart_layout(chart, left / width, top / chart_h, max(.3, 1 - (left + label_lane) / width), max(.3, 1 - (top + 34) / chart_h))
    for native, (spec, _) in zip(chart.series, variable):
        native.format.line.color.rgb = _rgb(style.muted if spec.reference else series_color(spec.key, style))
        native.format.line.width = Pt(1.0 if spec.reference else style.line_width)
        if spec.reference:
            native.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        native.marker.style = XL_MARKER_STYLE.NONE
        native.smooth = False
    endpoints = [(index, next(v for v in reversed(vals) if v is not None))
                 for index, (spec, vals) in enumerate(variable)
                 if not spec.reference and any(v is not None for v in vals)]
    span = (hi - lo) or 1
    for index, endpoint in endpoints:
        spec, vals = variable[index]
        if cap is not None and endpoint > cap:
            continue
        # Value and short unit only; the legend names the series.
        label = _br(endpoint, spec.decimals) + ("%" if spec.unit == "%" else "")
        last_index = max(i for i, v in enumerate(vals) if v is not None)
        if last_index < len(vals) - 1:
            label += f" ({block.categories[last_index]})"
        native_label = chart.series[index].points[last_index].data_label
        native_label.position = XL_DATA_LABEL_POSITION.RIGHT
        close = [other_index for other_index, other in endpoints
                 if other_index != index and abs(other - endpoint) / span < .08]
        if close:
            # Close final values: the higher one above, the lower one below.
            others = [dict(endpoints)[other] for other in close]
            higher = endpoint > max(others) or (endpoint == max(others) and index < min(close))
            native_label.position = XL_DATA_LABEL_POSITION.ABOVE if higher else XL_DATA_LABEL_POSITION.BELOW
        native_label.has_text_frame = True
        frame = native_label.text_frame
        frame.clear()
        frame.word_wrap = False
        paragraph = frame.paragraphs[0]
        run = paragraph.add_run()
        run.text = label
        run.font.name, run.font.size, run.font.bold = style.font, Pt(type_size), True
        run.font.color.rgb = _rgb(series_color(spec.key, style))
    # Native hierarchical labels remain horizontal, including narrow plots.
    for body in chart._chartSpace.xpath(".//c:catAx/c:txPr/a:bodyPr"):
        body.set("rot", "0")


def _bar_chart(slide, block, variable, x, y, width, height, style):
    data = CategoryChartData()
    data.categories = list(block.categories)
    for spec, vals in variable:
        data.add_series(spec.label, vals)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Pt(x), Pt(y), Pt(width), Pt(height), data).chart
    chart.has_title = False
    chart.has_legend = len(variable) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(10)
        chart.legend.font.name = style.font
    chart.font.name, chart.font.size = style.font, Pt(style.chart_size)
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.overlap = -10
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = '[$-416]0.0"%"'
    labels.number_format_is_linked = False
    labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(10)
    labels.font.name = style.font
    values = [v for _, vals in variable for v in vals if v is not None]
    lo, hi, major_unit, _ = _axis_bounds([min(0, *values), max(0, *values) * 1.18 or 1])
    axis = chart.value_axis
    axis.minimum_scale, axis.maximum_scale, axis.major_unit = lo, hi, major_unit
    axis.visible = False
    axis.has_major_gridlines = False
    category = chart.category_axis
    category.reverse_order = True  # first series on top, as in the table
    category.tick_labels.font.size = Pt(10)
    category.tick_labels.font.name = style.font
    category.format.line.color.rgb = _rgb(style.grid)
    category.major_tick_mark = XL_TICK_MARK.NONE
    colors = {"return_12m": style.orange, "cdi_12m": "8B9199"}
    for native, (spec, _) in zip(chart.series, variable):
        native.format.fill.solid()
        native.format.fill.fore_color.rgb = _rgb(colors.get(spec.key, style.muted))
        for index, category in enumerate(block.categories):
            if str(category).startswith("CDI"):
                point = native.points[index]
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = _rgb("8B9199")


def _short_series_label(label):
    return {
        "Inadimplência / direitos creditórios": "Vencidos",
        "Provisão / direitos creditórios": "Provisão",
        "Cobertura provisão / vencidos": "PDD / vencidos",
        "Cobertura PDD / NPL > 90d ex-360": "PDD / NPL >90d",
        "NPL > 1d ex-360": ">1d", "NPL > 30d ex-360": ">30d",
        "NPL > 60d ex-360": ">60d", "NPL > 90d ex-360": ">90d",
        "PL total": "PL", "Carteira bruta ex-360": "Carteira ex-360",
    }.get(label, label)


def _fits_horizontally(block, width, style):
    widths, _, _, size = _table_layout(block, width, style)
    for index, (column, allotted) in enumerate(zip(block.frame.columns, widths)):
        token = block.formats.get(column, "text").split(":")[0]
        numeric = column in block.numeric_columns or token in {"number", "money", "money_value", "percent", "integer"}
        month = bool(re.fullmatch(r"[a-z]{3}/\d{2}", str(column)))
        if numeric or month:
            font_size = size
            length = max((len(format_cell(value, block.formats.get(column, "text"))) for value in block.frame[column]), default=0)
            required = max(48 if month else 0, length * font_size * .56 + 12)
            if allotted + .01 < required:
                return False
    return True


def _table_chunks(block, style, width=None, height=None):
    """Pieces of a table that fit one slide each. A continuation of one or
    two rows is avoided by trying the compact font first."""
    chunks = list(_split_table(block, style, width, height))
    if (len(chunks) >= 2 and len(chunks[-1].frame) <= 2 and not getattr(block, "compact", False)
            and list(chunks[-1].frame.columns) == list(block.frame.columns)):
        compact = list(_split_table(replace(block, compact=True), style, width, height))
        if len(compact) < len(chunks):
            return compact
    return chunks


def _split_table(block, style, width=None, height=None):
    width = width or style.content_width
    height = height or style.content_bottom - style.content_top
    columns = list(block.frame.columns)
    # Repeat the identity column when a dense matrix needs another page.
    # One short text column (e.g. a position month) does not make a table textual.
    textual = [column for column in columns[1:] if not (column in block.numeric_columns or block.formats.get(column, "text").split(":")[0] in {"number", "money", "money_value", "percent", "integer"})]
    mostly_numeric = not textual or (len(textual) == 1 and block.frame[textual[0]].map(lambda value: len(str(value))).max() <= 16)
    max_columns = style.month_columns + 1 if len(columns) > 10 or mostly_numeric else 7
    while max_columns > 2 and not _fits_horizontally(replace(block, frame=block.frame.loc[:, columns[:max_columns]]), width, style):
        max_columns -= 1
    if len(columns) <= max_columns:
        groups = [columns]
    else:
        # Even groups (7 + 6, never 12 + 1); the identity column repeats.
        body = columns[1:]
        count = math.ceil(len(body) / (max_columns - 1))
        size = math.ceil(len(body) / count)
        groups = [[columns[0], *body[i:i + size]] for i in range(0, len(body), size)]
    yielded = 0

    def continued(segment):
        # A table that goes on in another place says so; two pieces of the
        # same table never repeat the same subtitle.
        nonlocal yielded
        yielded += 1
        if yielded > 1 and segment.title and not segment.title.endswith("(cont.)"):
            return replace(segment, title=f"{segment.title} (cont.)")
        return segment

    for cols in groups:
        selected = replace(block, frame=block.frame.loc[:, cols])
        start = 0
        while start < len(selected.frame):
            end = start
            for candidate_end in range(start + 1, len(selected.frame) + 1):
                candidate = replace(selected, frame=selected.frame.iloc[start:candidate_end])
                widths, header_h, heights, _ = _table_layout(candidate, width, style)
                if header_h + sum(heights) + (22 if block.title else 0) <= height:
                    end = candidate_end
                else:
                    break
            if end > start:
                yield continued(replace(selected, frame=selected.frame.iloc[start:end]))
                start = end
            else:
                # Split clauses too long for one page using that page's actual
                # column widths. Re-measuring a segment avoids height drift.
                single = replace(selected, frame=selected.frame.iloc[start:start + 1])
                widths, header_h, _, _ = _table_layout(single, width, style)
                available = height - header_h - (22 if block.title else 0)
                row = selected.frame.iloc[start]
                chunks = [textwrap.wrap(format_cell(value, selected.formats.get(col, "text")),
                                       max(12, int((w - 12) / (style.table_size * .49) * max(1, int(available / 15) - 3))),
                                       break_long_words=True, break_on_hyphens=False) or [""]
                          for col, value, w in zip(cols, row, widths)]
                for part in range(max(map(len, chunks))):
                    values = [pieces[part] if part < len(pieces) else "" for pieces in chunks]
                    values[0] = str(row.iloc[0]) if part else values[0]
                    yield continued(replace(selected, frame=pd.DataFrame([values], columns=cols), formats={}))
                start += 1
        if selected.frame.empty:
            yield continued(selected)


#: Decks longer than this get section dividers in ``section_dividers="auto"``.
DIVIDER_MIN_SLIDES = 40
#: ...and only before sections with at least this many slides.
DIVIDER_MIN_SECTION = 3
SECTION_DIVIDER_MODES = ("auto", "always", "never")


def _divider(prs, label, size, style):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(style.white)
    _text(slide, label, style.margin, 228, style.content_width, 40, size=30, bold=True, style=style)
    _text(slide, f"{size} slides", style.margin, 276, style.content_width, 20, size=13, color=style.muted, style=style)
    slide.notes_slide.notes_text_frame.text = f"Divisória de seção: {label}."


#: Smallest plot a chart keeps (plus its 22pt title and note band); free
#: space on a slide is then given back to the charts.
CHART_PLOT_MIN = 130.0
#: A chart with more series than this takes the full width.
HALF_WIDTH_SERIES = 3
#: Vertical gap between stacked blocks.
BLOCK_GAP = 12.0


@dataclass
class _Item:
    """One placeable piece: a chart, a table chunk or a text."""
    block: Any
    page: Any
    full: bool
    height: float
    chart: bool = False
    band: float = 0.0
    compact: bool = False
    group: str = ""
    x: float = 0.0
    y: float = 0.0
    width: float = 0.0


def _half(style):
    return (style.content_width - style.gap) / 2


def _merged_texts(blocks):
    from services.portfolio_ppt_content import TextBlock

    merged = []
    for block in blocks:
        if merged and isinstance(block, TextBlock) and isinstance(merged[-1], TextBlock) and block.title == merged[-1].title:
            merged[-1] = replace(merged[-1], text=f"{merged[-1].text}\n{block.text}")
        else:
            merged.append(block)
    return merged


def _text_height(block, width, *, full):
    """As drawn: ``_kpi`` beside other blocks (30pt title band, 13pt type on
    18pt lines); the 14pt-line flow text across the slide."""
    if full:
        return (22 if block.title else 0) + len(_lines(block.text, width, DEFAULT_STYLE.table_size)) * 14
    return 30 + len(_lines(block.text, width, 13)) * 18


def _items(page, style):
    """Size every block of a page. In ``panels`` pages every block takes half
    the width (side by side); elsewhere a table takes half the width when its
    natural width fits there."""
    from services.portfolio_ppt_content import ChartBlock, TableBlock

    half, full_width = _half(style), style.content_width
    height = style.content_bottom - style.content_top
    panels = page.layout == "panels"
    # Tables beside charts use the small type; tables paired with tables keep theirs.
    beside_charts = panels and any(isinstance(block, ChartBlock) for block in page.blocks)
    for block in _merged_texts(page.blocks):
        if isinstance(block, ChartBlock):
            drawn = [s for s in block.series if not s.reference]
            wide = len(drawn) > HALF_WIDTH_SERIES and not panels
            width = full_width if wide else half
            band = _band_height(_chart_notes(block)[1], width)
            yield _Item(block, page, wide, 22 + band + CHART_PLOT_MIN, chart=True, band=band)
        elif isinstance(block, TableBlock):
            compact = beside_charts or block.compact
            if not panels and _narrow_text_table(block, style):
                # Wrapped at half the width, beside the next one.
                for chunk in _table_chunks(replace(block, compact=True), style, half, height):
                    yield _Item(chunk, page, False, _measure(chunk, half, style, True), compact=True)
                continue
            for chunk in _table_chunks(replace(block, compact=compact), style, half if panels else full_width, height):
                widths, _, _, _ = _table_layout(chunk, full_width, style)
                if panels or sum(widths) <= half:
                    yield _Item(chunk, page, False, _measure(chunk, half, style, compact), compact=compact)
                else:
                    yield _Item(chunk, page, True, _measure(chunk, full_width, style, False))
        elif _text_height(block, half, full=False) <= height / 2:
            yield _Item(block, page, False, _text_height(block, half, full=False))
        else:
            # Long text runs across the slide and continues on the next one.
            lines = _lines(block.text, full_width, style.table_size)
            per_slide = max(1, int((height - 22) // 14))
            for start in range(0, len(lines), per_slide):
                title = block.title if start == 0 or not block.title else f"{block.title} (cont.)"
                piece = replace(block, title=title, text="\n".join(lines[start:start + per_slide]))
                yield _Item(piece, page, True, _text_height(piece, full_width, full=True))


def _measure(block, width, style, compact=False):
    _, header_h, rows, _ = _table_layout(block, width, style, compact=compact)
    return header_h + sum(rows) + (22 if block.title else 0)


def _narrow_text_table(block, style):
    """A two- or three-column text table (documentary fields) reads as well
    wrapped in half the width, and then pairs with the next one."""
    if block.snapshot or len(block.frame.columns) > 3 or len(block.frame) == 0:
        return False
    if any(_is_numeric_column(block, column) for column in block.frame.columns[1:]):
        return False
    widths, _, _, _ = _table_layout(block, style.content_width, style)
    half = _half(style)
    if sum(widths) > .72 * style.content_width:
        return False
    return _measure(block, half, style, True) <= 1.8 * _measure(block, style.content_width, style, block.compact)


def _split_to_fit(item, available, style):
    """Rows of a full-width table that fill the rest of a slide; the remaining
    rows continue on the next one. None when the piece would be too small."""
    from services.portfolio_ppt_content import TableBlock

    from services.portfolio_ppt_content import SECTION_APPENDIX, SECTION_COVERAGE, SECTION_DOCUMENTS

    block = item.block
    # Only reference material is split; a table read in the body stays whole.
    if (item.page.section not in {SECTION_APPENDIX, SECTION_COVERAGE, SECTION_DOCUMENTS} or not isinstance(block, TableBlock)
            or block.snapshot or len(block.frame) < 6 or available < 110):
        return None
    width = style.content_width if item.full else _half(style)
    head = next(iter(_split_table(replace(block, compact=item.compact), style, width, available)), None)
    if head is None or list(head.frame.columns) != list(block.frame.columns):
        return None
    rows = len(head.frame)
    if rows < 3 or len(block.frame) - rows < 3 or _measure(head, width, style, item.compact) > available:
        return None
    title = block.title if not block.title or block.title.endswith("(cont.)") else f"{block.title} (cont.)"
    rest = _table_chunks(replace(block, frame=block.frame.iloc[rows:], title=title, compact=item.compact), style, width,
                         style.content_bottom - style.content_top)
    return [replace(item, block=piece, height=_measure(piece, width, style, item.compact)) for piece in (head, *rest)]


def _pack(items, style):
    """Two-column shelves, filled in reading order. A half item goes to the
    shorter column of the open shelf; a full item closes it. A slide ends only
    when the next piece does not fit. Returns slides as lists of shelves."""
    limit = style.content_bottom - style.content_top + .5
    slides, shelves, top, cols = [], [], 0.0, None

    def close():
        nonlocal top, cols
        if cols is not None:
            top = max(sum(i.height for i in col) + BLOCK_GAP * len(col) for col in cols) + shelves[-1]["top"]
            cols = None

    def new_slide():
        nonlocal shelves, top, cols
        if shelves:
            slides.append(shelves)
        shelves, top, cols = [], 0.0, None

    queue = list(items)
    while queue:
        item = queue.pop(0)
        if item.block is None:  # a theme boundary inside a shared slide
            close()
            continue
        if item.full:
            close()
            if shelves and top + item.height > limit:
                pieces = _split_to_fit(item, limit - top - .5, style)
                if pieces:
                    item, queue = pieces[0], [*pieces[1:], *queue]
                else:
                    new_slide()
            shelves.append({"top": top, "full": item, "cols": None})
            top += item.height + BLOCK_GAP
            continue
        if cols is None:
            if shelves and top + item.height > limit:
                new_slide()
            cols = [[], []]
            shelves.append({"top": top, "full": None, "cols": cols})
        bottoms = [shelves[-1]["top"] + sum(i.height + BLOCK_GAP for i in col) for col in cols]
        side = 0 if bottoms[0] <= bottoms[1] else 1
        if bottoms[side] + item.height > limit:
            pieces = _split_to_fit(item, limit - bottoms[side] - .5, style)
            if pieces:
                cols[side].append(pieces[0])
                queue = [*pieces[1:], *queue]
                continue
            close()
            new_slide()
            cols = [[], []]
            shelves.append({"top": 0.0, "full": None, "cols": cols})
            side = 0
        cols[side].append(item)
    if shelves:
        slides.append(shelves)
    return slides


def _place(shelves, style):
    """Positions for one slide. Free height goes to shelves whose columns hold
    one block each with a chart among them, so charts grow instead of leaving
    a band of white below them; paired charts keep one plot height."""
    half = _half(style)
    limit = style.content_bottom - style.content_top

    def height_of(shelf):
        if shelf["full"] is not None:
            return shelf["full"].height
        return max((sum(i.height for i in col) + BLOCK_GAP * (len(col) - 1) for col in shelf["cols"] if col), default=0)

    used = sum(height_of(shelf) for shelf in shelves) + BLOCK_GAP * (len(shelves) - 1)
    growable = [shelf for shelf in shelves
                if (shelf["full"] is not None and shelf["full"].chart)
                or (shelf["cols"] is not None and all(len(col) <= 1 for col in shelf["cols"])
                    and any(col and col[0].chart for col in shelf["cols"]))]
    extra = max(0.0, limit - used) / len(growable) if growable else 0.0
    placed, y = [], style.content_top
    for shelf in shelves:
        grow = extra if any(shelf is g for g in growable) else 0.0
        if shelf["full"] is not None:
            item = shelf["full"]
            item.x, item.y, item.width = style.margin, y, style.content_width
            if item.chart:
                item.height += grow
            placed.append(item)
            y += item.height + BLOCK_GAP
            continue
        charts = [col[0] for col in shelf["cols"] if col and col[0].chart and len(col) == 1]
        if grow or len(charts) == 2:
            # Side-by-side charts share one height and one note band.
            target = max(height_of(shelf) + grow, *(c.height for c in charts)) if charts else 0
            band = max((c.band for c in charts), default=0)
            for chart in charts:
                chart.height, chart.band = target, band
        for side, col in enumerate(shelf["cols"]):
            cy = y
            for item in col:
                item.x, item.y, item.width = style.margin + side * (half + style.gap), cy, half
                placed.append(item)
                cy += item.height + BLOCK_GAP
        y += height_of(shelf) + BLOCK_GAP
    return placed


def _slide_spec(title, items, fallback):
    """Header of one slide: its own pages' context, source and notes only, so a
    long section no longer repeats every note on every slide."""
    pages = list(dict.fromkeys(id(item.page) for item in items))
    by_id = {id(item.page): item.page for item in items}
    first = by_id[pages[0]] if pages else fallback
    notes = [note for page_id in pages for note in by_id[page_id].notes]
    return replace(first, title=title, notes=tuple(dict.fromkeys(notes)), blocks=tuple(item.block for item in items))


def _draw_items(prs, title, items, style, fallback):
    from services.portfolio_ppt_content import ChartBlock, TableBlock

    slide = _header(prs, _slide_spec(title, items, fallback), style)
    for item in items:
        if isinstance(item.block, ChartBlock):
            _chart(slide, item.block, item.x, item.y, item.width, item.height, style, band=item.band)
        elif isinstance(item.block, TableBlock):
            _table(slide, item.block, item.x, item.y, item.width, style=style, compact=item.compact)
        elif item.full:
            lines = _lines(item.block.text, item.width, style.table_size)
            y = item.y
            if item.block.title:
                _text(slide, item.block.title, item.x, y, item.width, 22, size=style.section_size, bold=True, style=style)
                y += 22
            _text(slide, "\n".join(lines), item.x, y, item.width, item.height - (y - item.y), size=style.table_size,
                  color=style.muted if item.page.kind == "appendix" else style.text, style=style)
        else:
            _kpi(slide, item.block, item.x, item.y, item.width, item.height, style)


#: Small multiples per slide: (columns, rows) by number of charts.
def _grid_shape(count):
    return (2, 1) if count <= 2 else (2, 2) if count <= 4 else (3, 2) if count <= 6 else (3, 3) if count <= 9 else (4, 3)


def _draw_grid(prs, page, style):
    """One small chart per fund, same size and, for ratios, the same scale."""
    charts = list(page.blocks)
    count = 0
    for start in range(0, len(charts), 12):
        group = charts[start:start + 12]
        columns, rows = _grid_shape(len(group))
        gap = 16.0
        width = (style.content_width - gap * (columns - 1)) / columns
        height = (style.content_bottom - style.content_top - BLOCK_GAP * (rows - 1)) / rows
        title = page.title if start == 0 else f"{page.title} (cont.)"
        slide = _header(prs, replace(page, title=title, blocks=tuple(group)), style)
        for index, block in enumerate(group):
            x = style.margin + (index % columns) * (width + gap)
            y = style.content_top + (index // columns) * (height + BLOCK_GAP)
            _chart(slide, block, x, y, width, height, style, compact=True)
        count += 1
    return count


def _groups(spec):
    """Pages that read as one section are packed together: a theme's charts,
    its tables and its texts; the whole appendix; the whole documentation.
    Small-multiple pages and per-fund flows keep their own slides."""
    from services.portfolio_ppt_content import SECTION_APPENDIX, SECTION_DOCUMENTS, SECTION_FUNDS

    titles = {SECTION_APPENDIX: "Apêndice · séries", SECTION_DOCUMENTS: "Documental"}
    groups = []
    for page in spec.slides:
        own = page.layout == "grid" or page.section == SECTION_FUNDS
        key = (page.section or page.title, page.title if own else "")
        if own or not groups or groups[-1]["key"] != key:
            groups.append({"key": key, "label": page.section or page.title,
                           "title": titles.get(page.section, page.title), "pages": []})
        groups[-1]["pages"].append(page)
    return groups


def _render(spec, style, dividers=None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(style.width), Pt(style.height)
    prs.core_properties.title = spec.title
    prs.core_properties.author = "Toma Conta"
    _cover(prs, spec, style)
    counts = []
    for index, group in enumerate(_merge_small_themes(_groups(spec), style)):
        if dividers and index in dividers:
            _divider(prs, *dividers[index], style)
        before = len(prs.slides._sldIdLst)
        pages = group["pages"]
        if len(pages) == 1 and pages[0].layout == "grid":
            _draw_grid(prs, pages[0], style)
        else:
            for shelves in _pack(group["items"], style):
                placed = _place(shelves, style)
                # A slide shared by two sections names both.
                title = " · ".join(dict.fromkeys(item.group for item in placed if item.group)) or group["title"]
                _draw_items(prs, title, placed, style, pages[0])
        counts.append((group["label"], len(prs.slides._sldIdLst) - before))
    return prs, counts


def _merge_small_themes(groups, style):
    """Analysis themes keep a slide each, except when consecutive themes fit on
    one slide together (e.g. constant series that became sentences)."""
    from services.portfolio_ppt_content import SECTION_APPENDIX, SECTION_COVERAGE, SECTION_DOCUMENTS, THEME_TITLES

    themes = set(THEME_TITLES.values())
    closing = {SECTION_APPENDIX, SECTION_DOCUMENTS, SECTION_COVERAGE}
    merged = []
    for group in groups:
        group = {**group, "items": [] if group["pages"][0].layout == "grid" else
                 [replace(item, group=group["title"]) for page in group["pages"] for item in _items(page, style)]}
        previous = merged[-1] if merged else None
        if (previous is not None and group["label"] in themes and previous["themes"]
                and len(_pack(previous["items"], style)) == 1 and len(_pack(group["items"], style)) == 1):
            candidate = [*previous["items"], _Item(None, None, False, 0.0), *group["items"]]
            if len(_pack(candidate, style)) == 1:
                previous.update(items=candidate, pages=[*previous["pages"], *group["pages"]])
                continue
        if (previous is not None and previous["label"] in closing and group["label"] in closing and group["items"]):
            # Series appendix, documentation and references close the deck;
            # neighbours share a slide when that saves one.
            candidate = [*previous["items"], _Item(None, None, False, 0.0), *group["items"]]
            if len(_pack(candidate, style)) < len(_pack(previous["items"], style)) + len(_pack(group["items"], style)):
                previous.update(items=candidate, pages=[*previous["pages"], *group["pages"]])
                continue
        merged.append({**group, "themes": group["label"] in themes})
    return merged


def _section_dividers(counts, mode):
    if mode == "never" or (mode == "auto" and 1 + sum(count for _, count in counts) <= DIVIDER_MIN_SLIDES):
        return {}
    sections = []
    for index, (label, count) in enumerate(counts):
        if sections and sections[-1][0] == label:
            sections[-1][2] += count
        else:
            sections.append([label, index, count])
    # Short sections do not need a divider.
    return {index: (label, size) for label, index, size in sections if label and size >= DIVIDER_MIN_SECTION}


def render_portfolio_deck(spec, *, style=DEFAULT_STYLE, section_dividers="auto") -> bytes:
    if section_dividers not in SECTION_DIVIDER_MODES:
        raise ValueError(f"section_dividers deve ser um de {SECTION_DIVIDER_MODES}.")
    prs, counts = _render(spec, style)
    dividers = _section_dividers(counts, section_dividers)
    if dividers:
        prs, _ = _render(spec, style, dividers)
    stream = BytesIO()
    prs.save(stream)
    return stream.getvalue()
