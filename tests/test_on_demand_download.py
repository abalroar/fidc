from __future__ import annotations

from contextlib import nullcontext
from dataclasses import replace
from datetime import date
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import Mock, patch
import unittest
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook
import pandas as pd

from services.on_demand_download import (
    PreparedDownload, current_download_store, download_context_signature,
    prepare_download, render_on_demand_download,
)
from services.portfolio_store import PortfolioFund, PortfolioRecord
from services.fidc_model.b3_cdi import B3CdiMonthlyRate
from services.mercado_livre_dashboard import MercadoLivreOutputs, build_consolidated_snapshot_excel_bytes
from tabs import portfolio_page


def _portfolio_download_test_app(portfolio, analysis) -> None:
    import streamlit as st
    from types import SimpleNamespace
    from unittest.mock import patch
    from tabs import portfolio_page

    selected = st.selectbox("Período do teste", ["original", "revisto"])
    period = SimpleNamespace(cache_key=selected, label=selected)
    st.toggle("Somente fundo escolhido na tela", key="scope_filter")

    def build_snapshot(*args, **kwargs):
        st.session_state["snapshot_calls"] = st.session_state.get("snapshot_calls", 0) + 1
        return b"snapshot-original"

    with (
        patch("tabs.portfolio_page._portfolio_download_context_signature", return_value=selected),
        patch("tabs.portfolio_page.credit_tab.resolve_fund_return_export_inputs", return_value=({}, {})),
        patch("tabs.portfolio_page.build_consolidated_snapshot_excel_bytes", side_effect=build_snapshot),
    ):
        portfolio_page._render_unified_portfolio_download(
            analysis=analysis, selected_portfolio=portfolio, period=period,
        )


class OnDemandDownloadTests(unittest.TestCase):
    def test_render_and_download_reruns_do_not_generate_until_requested(self) -> None:
        state = {}
        producer = Mock(return_value=b"original-file")
        kwargs = dict(
            namespace="test", artifact="xlsx", label="Resumo (Excel)",
            file_name="resumo.xlsx", mime="application/xlsx", producer=producer,
        )
        with (
            patch("services.on_demand_download.st.button", side_effect=[False, True, False]),
            patch("services.on_demand_download.st.spinner", return_value=nullcontext()),
            patch("services.on_demand_download.st.download_button") as download,
        ):
            for request in range(3):
                store = current_download_store(state, namespace="test", signature="same-data")
                render_on_demand_download(store=store, **kwargs)
                self.assertEqual(0 if request == 0 else 1, producer.call_count)
        self.assertEqual(2, download.call_count)
        self.assertEqual(b"original-file", download.call_args.kwargs["data"])
        self.assertEqual("ignore", download.call_args.kwargs["on_click"])

    def test_filter_or_data_change_discards_prepared_file(self) -> None:
        state = {}
        for version in ("initial", "other-portfolio", "other-period", "retified-data"):
            store = current_download_store(state, namespace="test", signature=version)
            self.assertEqual({}, store)
            prepare_download(store, artifact="xlsx", producer=lambda: b"file")
            self.assertEqual({"xlsx": PreparedDownload(b"file")}, current_download_store(
                state, namespace="test", signature=version,
            ))

    def test_failed_refresh_removes_previous_file_and_next_request_recovers(self) -> None:
        store = {"xlsx": PreparedDownload(b"old")}
        producer = Mock(side_effect=[RuntimeError("failed"), b"recovered"])
        with self.assertRaises(RuntimeError):
            prepare_download(store, artifact="xlsx", producer=producer)
        self.assertNotIn("xlsx", store)
        self.assertEqual(b"recovered", prepare_download(store, artifact="xlsx", producer=producer).data)

    def test_signature_includes_all_rows_zero_missing_and_updated_source(self) -> None:
        frame = pd.DataFrame({"value": [0.0, float("nan"), 3.0]})
        baseline = download_context_signature(frame)
        self.assertEqual(baseline, download_context_signature(frame.copy(deep=True)))
        revised = frame.copy(deep=True)
        revised.iloc[0, 0] = float("nan")
        self.assertNotEqual(baseline, download_context_signature(revised))
        revised = frame.copy(deep=True)
        revised.iloc[-1, 0] = 4.0
        self.assertNotEqual(baseline, download_context_signature(revised))
        with TemporaryDirectory() as directory:
            source = Path(directory) / "manifest.json"
            missing = download_context_signature(frame, source_paths=(source,))
            source.write_text("version-one", encoding="utf-8")
            present = download_context_signature(frame, source_paths=(source,))
            self.assertNotEqual(missing, present)
            source.write_text("version-two", encoding="utf-8")
            self.assertNotEqual(present, download_context_signature(frame, source_paths=(source,)))


class PortfolioDemandDownloadTests(unittest.TestCase):
    def setUp(self) -> None:
        self.portfolio = PortfolioRecord(
            id="7681d350f2994945ada0590b0e3703c2", name="Cloudwalk",
            funds=tuple(PortfolioFund(cnpj=cnpj, display_name=f"Cloudwalk {cnpj}") for cnpj in (
                "42085816000105", "42085830000109", "42102603000144", "44124617000194",
                "54218673000141", "54218941000125", "54219179000100", "54248022000102",
                "57609282000146", "60356171000180", "62393679000183",
            )),
            created_at="2026-10-11", updated_at="2026-10-11",
        )
        self.analysis = portfolio_page.PortfolioAnalysisData(
            scopes=(), aggregate_bundle=None,
            outputs=SimpleNamespace(fund_monthly={}),
            monitor_outputs=SimpleNamespace(), research_outputs=SimpleNamespace(),
            verification_report=None, dashboard_errors={}, load_errors={},
        )
        self.period = SimpleNamespace(cache_key="2025-10::2026-09", label="10/2025 a 09/2026")

    def test_render_does_not_prepare_any_file_on_render(self) -> None:
        state = {}
        with (
            patch("tabs.portfolio_page.st.session_state", state),
            patch("tabs.portfolio_page.st.caption"),
            patch("tabs.portfolio_page.st.expander", return_value=nullcontext()),
            patch("services.on_demand_download.st.button", return_value=False),
            patch("tabs.portfolio_page.credit_tab.resolve_fund_return_export_inputs") as return_inputs,
            patch("tabs.portfolio_page._prepare_complete_portfolio_pptx") as ppt,
            patch("tabs.portfolio_page.build_consolidated_snapshot_excel_bytes") as snapshot,
            patch("tabs.portfolio_page.build_full_variable_excel_export_bytes") as full,
            patch("tabs.portfolio_page.build_full_variable_csv_zip_bytes") as csv,
            patch("tabs.portfolio_page._prepare_analytical_portfolio_xlsx") as analytical,
        ):
            for _ in range(2):
                portfolio_page._render_unified_portfolio_download(
                    analysis=self.analysis, selected_portfolio=self.portfolio, period=self.period,
                )
        for producer in (ppt, snapshot, full, csv, analytical):
            producer.assert_not_called()
        return_inputs.assert_not_called()

    def test_request_prepares_only_snapshot_and_changes_invalidate_it(self) -> None:
        state = {}
        with (
            patch("tabs.portfolio_page.st.session_state", state),
            patch("tabs.portfolio_page.st.caption"),
            patch("tabs.portfolio_page.st.expander", return_value=nullcontext()),
            patch("services.on_demand_download.st.spinner", return_value=nullcontext()),
            patch("services.on_demand_download.st.download_button") as download,
            patch("services.on_demand_download.st.button", side_effect=lambda *a, **kw: kw["key"].endswith("::snapshot_excel")),
            patch("tabs.portfolio_page._prepare_complete_portfolio_pptx") as ppt,
            patch("tabs.portfolio_page.build_consolidated_snapshot_excel_bytes", return_value=b"snapshot") as snapshot,
            patch("tabs.portfolio_page.build_full_variable_excel_export_bytes") as full,
            patch("tabs.portfolio_page.build_full_variable_csv_zip_bytes") as csv,
            patch("tabs.portfolio_page._prepare_analytical_portfolio_xlsx") as analytical,
        ):
            portfolio_page._render_unified_portfolio_download(
                analysis=self.analysis, selected_portfolio=self.portfolio, period=self.period,
            )
            snapshot.assert_called_once()
            self.assertIs(self.analysis.outputs, snapshot.call_args.args[0])
            context = snapshot.call_args.kwargs["export_context"]
            self.assertEqual(len(self.portfolio.funds), context["Fundos solicitados"])
            self.assertEqual(self.period.label, context["Período solicitado"])
            self.assertEqual(len(self.portfolio.funds), len(context["CNPJs da cesta solicitada"].split("; ")))
            self.assertEqual(b"snapshot", download.call_args.kwargs["data"])
            for producer in (ppt, full, csv, analytical):
                producer.assert_not_called()
            with patch("services.on_demand_download.st.button", return_value=False):
                download.reset_mock()
                self.period.cache_key = "2026-01::2026-09"
                portfolio_page._render_unified_portfolio_download(
                    analysis=self.analysis, selected_portfolio=self.portfolio, period=self.period,
                )
                download.assert_not_called()
            self.assertEqual(1, snapshot.call_count)

    def test_other_portfolios_prepare_separate_files_without_calling_legacy(self) -> None:
        state = {}
        other_portfolio = replace(self.portfolio, id="other", name="Outra carteira")
        with (
            patch("tabs.portfolio_page.st.session_state", state),
            patch("tabs.portfolio_page.st.caption"),
            patch("tabs.portfolio_page.st.expander", return_value=nullcontext()),
            patch("services.on_demand_download.st.spinner", return_value=nullcontext()),
            patch("services.on_demand_download.st.download_button"),
            patch("services.on_demand_download.st.button", side_effect=lambda *a, **kw: kw["key"].endswith("::snapshot_excel")),
            patch("tabs.portfolio_page.build_consolidated_snapshot_excel_bytes", side_effect=[b"first", b"other"]),
            patch("tabs.portfolio_page._render_legacy_portfolio_download") as legacy,
        ):
            for portfolio in (self.portfolio, other_portfolio):
                portfolio_page._render_unified_portfolio_download(
                    analysis=self.analysis, selected_portfolio=portfolio, period=self.period,
                )
        legacy.assert_not_called()
        first = state[f"prepared_downloads::portfolio::{self.portfolio.id}"]["files"]
        other = state["prepared_downloads::portfolio::other"]["files"]
        self.assertEqual(b"first", first["snapshot_excel"].data)
        self.assertEqual(b"other", other["snapshot_excel"].data)
        self.assertIsNot(first, other)

    def test_streamlit_request_rerun_help_and_period_invalidation(self) -> None:
        from streamlit.testing.v1 import AppTest

        app = AppTest.from_function(_portfolio_download_test_app, args=(self.portfolio, self.analysis)).run()
        self.assertEqual([], list(app.exception))
        prepare = next(button for button in app.button if button.label == "Preparar: Resumo histórico da carteira (Excel)")
        self.assertIn("cobertura variável", prepare.proto.help)
        prepare.click().run()
        self.assertEqual([], list(app.exception))
        store_key = f"prepared_downloads::portfolio::{self.portfolio.id}"
        self.assertEqual(b"snapshot-original", app.session_state[store_key]["files"]["snapshot_excel"].data)
        self.assertEqual(1, app.session_state["snapshot_calls"])
        download = next(element for element in app.get("download_button") if element.proto.label == "Baixar: Resumo histórico da carteira (Excel)")
        self.assertIn("posição comum", download.proto.help)
        app.run()
        self.assertEqual(b"snapshot-original", app.session_state[store_key]["files"]["snapshot_excel"].data)
        self.assertEqual(1, app.session_state["snapshot_calls"])
        app.toggle(key="scope_filter").set_value(True).run()
        self.assertEqual([], list(app.exception))
        self.assertEqual(b"snapshot-original", app.session_state[store_key]["files"]["snapshot_excel"].data)
        self.assertEqual(1, app.session_state["snapshot_calls"])
        app.selectbox[0].set_value("revisto").run()
        self.assertEqual([], list(app.exception))
        self.assertEqual({}, app.session_state[store_key]["files"])
        self.assertEqual(1, app.session_state["snapshot_calls"])
        self.assertEqual(0, len(app.get("download_button")))

    def test_cdi_or_benchmark_correction_invalidates_only_ppt_without_regenerating(self) -> None:
        state = {}
        cnpj = self.portfolio.funds[0].cnpj
        initial_rate = B3CdiMonthlyRate(
            mes="2026-02", cdi_mensal=0.01, dias_uteis=18,
            data_inicio=date(2026, 2, 1), data_fim=date(2026, 2, 28), source="B3",
        )
        initial_inputs = ({cnpj: (initial_rate,)}, {cnpj: {"senior:1": 0.02}})
        cdi_corrected = ({cnpj: (replace(initial_rate, cdi_mensal=0.011),)}, initial_inputs[1])
        benchmark_corrected = (cdi_corrected[0], {cnpj: {"senior:1": 0.025}})
        with (
            patch("tabs.portfolio_page.st.session_state", state),
            patch("tabs.portfolio_page.st.caption"),
            patch("tabs.portfolio_page.st.expander", return_value=nullcontext()),
            patch("services.on_demand_download.st.spinner", return_value=nullcontext()),
            patch("services.on_demand_download.st.download_button") as download,
            patch("services.on_demand_download.st.button", side_effect=lambda *a, **kw: kw["key"].endswith(("::complete_pptx", "::snapshot_excel"))),
            patch("tabs.portfolio_page.credit_tab.resolve_fund_return_export_inputs", side_effect=[initial_inputs, initial_inputs, cdi_corrected, cdi_corrected, benchmark_corrected]) as resolver,
            patch("tabs.portfolio_page._prepare_complete_portfolio_pptx", return_value=PreparedDownload(b"pptx")) as ppt,
            patch("tabs.portfolio_page.build_consolidated_snapshot_excel_bytes", return_value=b"snapshot") as snapshot,
        ):
            portfolio_page._render_unified_portfolio_download(
                analysis=self.analysis, selected_portfolio=self.portfolio, period=self.period,
            )
            self.assertEqual(initial_inputs, ppt.call_args.kwargs["return_inputs"])
            with patch("services.on_demand_download.st.button", return_value=False):
                download.reset_mock()
                portfolio_page._render_unified_portfolio_download(
                    analysis=self.analysis, selected_portfolio=self.portfolio, period=self.period,
                )
                self.assertEqual({b"pptx", b"snapshot"}, {call.kwargs["data"] for call in download.call_args_list})
                download.reset_mock()
                portfolio_page._render_unified_portfolio_download(
                    analysis=self.analysis, selected_portfolio=self.portfolio, period=self.period,
                )
                self.assertEqual([b"snapshot"], [call.kwargs["data"] for call in download.call_args_list])
            self.assertEqual(1, ppt.call_count)  # Correction invalidates without generating.
            with patch("services.on_demand_download.st.button", side_effect=lambda *a, **kw: kw["key"].endswith("::complete_pptx")):
                portfolio_page._render_unified_portfolio_download(
                    analysis=self.analysis, selected_portfolio=self.portfolio, period=self.period,
                )
            self.assertEqual(cdi_corrected, ppt.call_args.kwargs["return_inputs"])
            with patch("services.on_demand_download.st.button", return_value=False):
                download.reset_mock()
                portfolio_page._render_unified_portfolio_download(
                    analysis=self.analysis, selected_portfolio=self.portfolio, period=self.period,
                )
                self.assertEqual([b"snapshot"], [call.kwargs["data"] for call in download.call_args_list])
            self.assertEqual(2, ppt.call_count)
            self.assertEqual(1, snapshot.call_count)
            self.assertEqual(5, resolver.call_count)

    def test_return_input_failure_keeps_page_and_excel_downloads_available(self) -> None:
        state = {
            f"prepared_downloads::portfolio::{self.portfolio.id}::pptx": {
                "signature": "previous-context", "files": {"complete_pptx": PreparedDownload(b"old-pptx")},
            },
        }
        with (
            patch("tabs.portfolio_page.st.session_state", state),
            patch("tabs.portfolio_page.st.caption"),
            patch("tabs.portfolio_page.st.expander", return_value=nullcontext()),
            patch("services.on_demand_download.st.spinner", return_value=nullcontext()),
            patch("services.on_demand_download.st.download_button") as download,
            patch("services.on_demand_download.st.button", side_effect=lambda *a, **kw: kw["key"].endswith("::snapshot_excel")),
            patch("tabs.portfolio_page.credit_tab.resolve_fund_return_export_inputs", side_effect=RuntimeError("offline")),
            patch("tabs.portfolio_page._prepare_complete_portfolio_pptx") as ppt,
            patch("tabs.portfolio_page.build_consolidated_snapshot_excel_bytes", return_value=b"snapshot") as snapshot,
        ):
            portfolio_page._render_unified_portfolio_download(
                analysis=self.analysis, selected_portfolio=self.portfolio, period=self.period,
            )
        ppt.assert_not_called()
        snapshot.assert_called_once()
        self.assertEqual(b"snapshot", download.call_args.kwargs["data"])

    def test_document_source_revision_discards_prepared_file_without_regenerating(self) -> None:
        state = {}
        with TemporaryDirectory() as directory:
            source = Path(directory) / "rdb-source.csv"
            source.write_text("source-version-one", encoding="utf-8")
            with (
                patch("pathlib.Path.rglob", side_effect=lambda *args: iter((source,))),
                patch("tabs.portfolio_page.st.session_state", state),
                patch("tabs.portfolio_page.st.caption"),
                patch("tabs.portfolio_page.st.expander", return_value=nullcontext()),
                patch("services.on_demand_download.st.spinner", return_value=nullcontext()),
                patch("services.on_demand_download.st.download_button") as download,
                patch("services.on_demand_download.st.button", side_effect=lambda *a, **kw: kw["key"].endswith("::snapshot_excel")),
                patch("tabs.portfolio_page.build_consolidated_snapshot_excel_bytes", return_value=b"snapshot") as snapshot,
            ):
                portfolio_page._render_unified_portfolio_download(
                    analysis=self.analysis, selected_portfolio=self.portfolio, period=self.period,
                )
                self.assertEqual(b"snapshot", download.call_args.kwargs["data"])
                source.write_text("source-version-two", encoding="utf-8")
                with patch("services.on_demand_download.st.button", return_value=False):
                    download.reset_mock()
                    portfolio_page._render_unified_portfolio_download(
                        analysis=self.analysis, selected_portfolio=self.portfolio, period=self.period,
                    )
                download.assert_not_called()
                snapshot.assert_called_once()
                self.assertEqual({}, state[f"prepared_downloads::portfolio::{self.portfolio.id}"]["files"])

    def test_partial_source_scope_preserves_requested_basket_and_warns_on_omitted_blocks(self) -> None:
        from pptx import Presentation
        from pptx.util import Inches
        native = Presentation()
        native.slide_width = Inches(13.333)
        native.slide_height = Inches(7.5)
        native.slides.add_slide(native.slide_layouts[6])
        stream = BytesIO()
        native.save(stream)
        cnpj = self.portfolio.funds[0].cnpj
        analysis = replace(
            self.analysis,
            outputs=SimpleNamespace(fund_monthly={cnpj: pd.DataFrame()}),
            scopes=(),
        )
        with (
            patch("services.fundonet_ppt_export.build_dashboard_pptx_bytes") as dashboard,
            patch("services.somatorio_fidcs_ppt_export.build_somatorio_fidcs_pptx_bytes") as monthly,
            patch("services.fidc_analytical_slide.build_fidc_analytical_pptx_bytes", return_value=stream.getvalue()) as analytical,
            patch("tabs.portfolio_page._build_documentary_portfolio_deck", return_value=None),
            patch("tabs.portfolio_page.credit_tab.resolve_fund_return_export_inputs") as cdi,
        ):
            prepared = portfolio_page._prepare_complete_portfolio_pptx(
                analysis=analysis, selected_portfolio=self.portfolio, period=self.period,
            )
        analytical.assert_called_once_with(self.portfolio.funds, portfolio_name=self.portfolio.name, compact=True)
        dashboard.assert_not_called()
        monthly.assert_not_called()
        cdi.assert_not_called()
        finalized = Presentation(BytesIO(prepared.data))
        self.assertEqual(3, len(finalized.slides))  # Cover, available block, references.
        self.assertIn(self.portfolio.name, finalized.slides[0].shapes[0].text)
        self.assertIn("Visões Fundos.NET e base mensal: fundos inconsistentes entre as fontes carregadas.", prepared.warnings)
        self.assertIn("Curadoria documental: pacote comparativo não localizado para esta cesta.", prepared.warnings)

    def test_native_ppt_composition_matches_previous_pipeline(self) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        def native_deck(title: str, value: str) -> bytes:
            deck = Presentation()
            deck.slide_width = Inches(13.333)
            deck.slide_height = Inches(7.5)
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2)).table
            table.cell(0, 0).text = title
            table.cell(1, 1).text = value
            slide.notes_slide.notes_text_frame.text = "Fonte documental · 02/2026 · R$"
            stream = BytesIO()
            deck.save(stream)
            return stream.getvalue()

        cnpj = self.portfolio.funds[0].cnpj
        analysis = replace(self.analysis,
            outputs=SimpleNamespace(fund_monthly={cnpj: pd.DataFrame()}),
            scopes=(portfolio_page.PortfolioAnalysisScope(
                value=cnpj, label="Cloudwalk", kind="fund", cnpj=cnpj, dashboard=SimpleNamespace(),
            ),),
        )
        with (
            patch("tabs.portfolio_page.credit_tab.resolve_fund_return_export_inputs", return_value=({}, {})) as resolver,
            patch("services.somatorio_fidcs_ppt_export.build_somatorio_fidcs_pptx_bytes", return_value=native_deck("Base mensal", "100.000.000")),
            patch("services.fidc_analytical_slide.build_fidc_analytical_pptx_bytes", return_value=native_deck("Emissões", "N/D")),
            patch("services.fidc_analytical_slide.build_fidc_analytical_xlsx_bytes", return_value=b"xlsx"),
            patch("tabs.portfolio_page._build_documentary_portfolio_deck", return_value=native_deck("Regulamento", "0")),
            patch("tabs.portfolio_page.build_consolidated_snapshot_excel_bytes", return_value=b"xlsx"),
            patch("tabs.portfolio_page.build_full_variable_excel_export_bytes", return_value=b"xlsx"),
            patch("tabs.portfolio_page.build_full_variable_csv_zip_bytes", return_value=b"csv"),
            patch("tabs.portfolio_page.st.expander", return_value=nullcontext()),
            patch("tabs.portfolio_page.st.download_button") as download,
        ):
            portfolio_page._render_legacy_portfolio_download(
                analysis=analysis, selected_portfolio=self.portfolio, period=self.period,
            )
            baseline = next(call.kwargs["data"] for call in download.call_args_list
                if call.kwargs["mime"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation")
            prepared = portfolio_page._prepare_complete_portfolio_pptx(
                analysis=analysis, selected_portfolio=self.portfolio, period=self.period,
                return_inputs=({}, {}),
            )
            resolver.assert_called_once()  # Baseline only; the requested PPT reuses resolved inputs.
        before = Presentation(BytesIO(baseline))
        after = Presentation(BytesIO(prepared.data))
        self.assertEqual(3, len(before.slides))
        self.assertEqual(5, len(after.slides))  # Native blocks plus portfolio cover and references.
        for left, right in zip(before.slides, list(after.slides)[1:-1], strict=True):
            self.assertIn(left.notes_slide.notes_text_frame.text, right.notes_slide.notes_text_frame.text)
            left_tables = [shape.table for shape in left.shapes if shape.has_table]
            right_tables = [shape.table for shape in right.shapes if shape.has_table]
            self.assertEqual(len(left_tables), len(right_tables))
            for old, new in zip(left_tables, right_tables, strict=True):
                self.assertEqual([[cell.text for cell in row.cells] for row in old.rows],
                                 [[cell.text for cell in row.cells] for row in new.rows])

    def test_snapshot_values_types_charts_and_missing_values_are_preserved(self) -> None:
        monthly = pd.DataFrame({
            "competencia": ["01/2026", "02/2026"],
            "competencia_dt": pd.to_datetime(["2026-01-01", "2026-02-01"]),
            "pl_total": [100_000_000.0, 0.0], "pl_senior": [75_000_000.0, 0.0],
            "pl_subordinada_mezz_ex360": [25_000_000.0, float("nan")],
            "subordinacao_total_ex360_pct": [25.0, float("nan")],
        })
        outputs = MercadoLivreOutputs(
            fund_monthly={}, fund_wide={}, consolidated_monthly=monthly,
            consolidated_wide=pd.DataFrame(), warnings_df=pd.DataFrame(), metadata={},
        )
        original = build_consolidated_snapshot_excel_bytes(outputs)
        store = {}
        prepared = prepare_download(
            store, artifact="snapshot", producer=lambda: build_consolidated_snapshot_excel_bytes(
                outputs, export_context={"Fonte": "Informe Mensal CVM/Fundos.NET", "Fundos solicitados": 11},
            ),
        )
        before = load_workbook(BytesIO(original), data_only=True)
        after = load_workbook(BytesIO(prepared.data), data_only=True)
        self.assertEqual(before.sheetnames, after.sheetnames[:3])
        self.assertEqual("Contexto da exportação", after.sheetnames[3])
        self.assertEqual("Informe Mensal CVM/Fundos.NET", after["Contexto da exportação"]["B2"].value)
        self.assertEqual(11, after["Contexto da exportação"]["B3"].value)
        for name in before.sheetnames:
            self.assertEqual(list(before[name].values), list(after[name].values))
        pl_row = next(row for row in after["Resumo exibido"].values if row[0] == "PL FIDC total")
        self.assertEqual(("PL FIDC total", 0, 100_000_000), pl_row)
        sub_row = next(row for row in after["Resumo exibido"].values if row[0] == "Subordinada + Mezanino ex-360")
        self.assertIsNone(sub_row[1])
        self.assertEqual(len(before["Gráficos"]._charts), len(after["Gráficos"]._charts))
        self.assertGreaterEqual(len(after["Gráficos"]._charts), 2)
        for name in ("Resumo exibido", "Dados gráficos"):
            self.assertEqual("B2", after[name].freeze_panes)
            for left_row, right_row in zip(before[name], after[name], strict=True):
                for left, right in zip(left_row, right_row, strict=True):
                    self.assertEqual(left.data_type, right.data_type)
                    self.assertEqual(left.number_format, right.number_format)
        namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with ZipFile(BytesIO(prepared.data)) as archive:
            for path in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
                root = ElementTree.fromstring(archive.read(path))
                view = root.find("x:sheetViews/x:sheetView", namespace)
                pane = view.find("x:pane", namespace)
                self.assertEqual("1", pane.get("xSplit"))
                self.assertEqual("1", pane.get("ySplit"))
                self.assertEqual("B2", pane.get("topLeftCell"))
                selections = view.findall("x:selection", namespace)
                names = [selection.get("pane") for selection in selections]
                self.assertEqual(len(names), len(set(names)))
                self.assertEqual({"topRight", "bottomLeft", "bottomRight"}, set(names))
                self.assertIn(pane.get("activePane"), names)


class MeliDemandDownloadTests(PortfolioDemandDownloadTests):
    """Exercise the complete download lifecycle with the approved Seller basket."""

    def setUp(self) -> None:
        super().setUp()
        self.portfolio = PortfolioRecord(
            id="47b0eb1b7bea42ef85233f9b1186129c", name="MELI (FIDC Sellers I, II e III)",
            funds=tuple(PortfolioFund(cnpj=cnpj, display_name=name) for cnpj, name in (
                ("50473039000102", "SELLER FIDC SEGMENTO MEIOS DE PAGAMENTO DE RESPONSABILIDADE LIMITADA"),
                ("55471753000177", "SELLER II FUNDO DE INVESTIMENTO EM DIREITOS CREDITÓRIOS SEGMENTO MEIOS DE PAGAMENTO RESP LTDA"),
                ("63572282000111", "SELLER 3 FUNDO DE INVESTIMENTO EM DIREITOS CREDITÓRIOS SEGMENTO MEIOS DE PAGAMENTO DE RESP LIMITADA"),
            )),
            created_at="2026-04-22", updated_at="2026-04-22",
        )


if __name__ == "__main__":
    unittest.main()
