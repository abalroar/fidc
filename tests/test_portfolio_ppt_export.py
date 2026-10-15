from __future__ import annotations

from dataclasses import fields, is_dataclass
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pandas as pd
import pytest
from openpyxl import load_workbook
from pptx import Presentation

from services.document_curation_comparison import DocumentComparisonPage
from services.portfolio_store import PortfolioFund


def _inputs(*, months=13, funds=1, series=3, missing=False, long_document=False):
    dates = pd.date_range('2025-06-01', periods=months, freq='MS')
    fund_objects = tuple(PortfolioFund(cnpj=f'{index+1:014d}', display_name=f'Fundo {index+1}') for index in range(funds))
    monthly, monitor, histories, summaries = {}, {}, {}, {}
    for index, fund in enumerate(fund_objects):
        frame = pd.DataFrame({
            'competencia': dates.strftime('%m/%Y'), 'competencia_dt': dates,
            'cnpj': fund.cnpj, 'fund_name': fund.display_name,
            'pl_total': [1_003_456_789.12 + index*100_000_000 + month*10_000_000 for month in range(months)],
            'pl_senior': [700_000_000 + index*10_000_000 + month*1_000_000 for month in range(months)],
            'pl_subordinada_mezz_ex360': [300_000_000 + index*10_000_000 + month*2_000_000 for month in range(months)],
            'carteira_bruta': [900_000_000 + index*100_000_000 + month*3_000_000 for month in range(months)],
            'carteira_ex360': [850_000_000 + index*100_000_000 + month*3_000_000 for month in range(months)],
            'subordinacao_total_ex360_pct': 30.0,
            'npl_over90_ex360_pct': [0.0 if month==0 else 2.345678+index+month/100 for month in range(months)],
            'pdd_npl_over90_ex360_pct': 200.0,
            'npl_over1_ex360_pct': 10.0, 'npl_over30_ex360_pct': 8.0, 'npl_over60_ex360_pct': 5.0,
            'carteira_ex360_yoy_pct': [float('nan')]*months,
            'npl_1_90_pct': 2.0, 'npl_91_360_pct': 1.0,
            'duration_months': [6.0+index+month/10 for month in range(months)],
            'roll_61_90_m3_pct': [float('nan') if month<3 else 1.2345+index/10+month/100 for month in range(months)],
            'roll_91_120_m4_pct': [float('nan') if month<4 else 2.2345+index/10+month/100 for month in range(months)],
            'roll_121_150_m5_pct': [float('nan') if month<5 else 3.2345+index/10+month/100 for month in range(months)],
            'roll_151_180_m6_pct': [float('nan') if month<6 else 4.2345+index/10+month/100 for month in range(months)],
        })
        if missing and months>1:
            frame.loc[1,'npl_over90_ex360_pct'] = float('nan')
            frame.loc[1,'pl_total'] = float('nan')
            frame['npl_over90_ex360_source_status'] = ['reported_zero' if month==0 else 'not_reported' if month==1 else 'reported_value' for month in range(months)]
        monthly[fund.cnpj] = frame
        monitor[fund.cnpj] = frame.copy()
        histories[fund.cnpj] = pd.DataFrame([
            dict(competencia=timestamp.strftime('%m/%Y'), competencia_dt=timestamp,
                 class_kind='senior', class_key=f'senior:{row:02d}',
                 class_label=f'{row+1:02d}ª série',
                 retorno_mensal_pct=0.0 if row==0 else float('nan') if missing and month==1 else 1.0+row/100)
            for row in range(series) for month,timestamp in enumerate(dates)
        ])
        summaries[fund.cnpj] = pd.DataFrame([
            dict(class_kind='senior', class_key=f'senior:{row:02d}', class_label=f'{row+1:02d}ª série',
                 latest_competencia=dates[-1].strftime('%m/%Y'), retorno_ano_pct=row/10,
                 retorno_12m_pct=row/5, ytd_status='incompleto' if missing else 'completo',
                 ytd_competencias_ausentes='07/2025' if missing else '')
            for row in range(series)
        ])
    consolidated = monthly[fund_objects[0].cnpj].copy()
    if funds>1:
        consolidated['fund_name'] = 'Carteira'
        for column in ('pl_total','pl_senior','pl_subordinada_mezz_ex360','carteira_bruta','carteira_ex360'):
            consolidated[column] = sum(monthly[fund.cnpj][column] for fund in fund_objects)
    outputs = SimpleNamespace(fund_monthly=monthly, consolidated_monthly=consolidated,
        fund_return_history=histories, fund_return_summary=summaries, warnings_df=pd.DataFrame(),
        metadata={'portfolio_name':'Carteira', 'funds':[fund.to_dict() for fund in fund_objects]})
    monitor_outputs = SimpleNamespace(fund_monitor=monitor, consolidated_monitor=consolidated.copy(),
        fund_cohorts={}, consolidated_cohorts=pd.DataFrame(), audit_table=pd.DataFrame(), warnings=[])
    document = ' '.join(f'cláusula{index}' for index in range(220)) if long_document else 'Mensal conforme suplemento'
    page = DocumentComparisonPage('Condições', pd.DataFrame({'Critério':['Amortização','Quantidade','Taxa'],
        fund_objects[0].display_name:[document,0,None]}, dtype=object),
        notes=('Cobertura parcial; comparação inválida sem associação documental.',),
        sources=('Regulamento ID CVM 123, cláusula 7.1, página 18.',))
    return dict(portfolio_name='Carteira', funds=fund_objects, outputs=outputs,
        monitor_outputs=monitor_outputs, documentary_pages=(page,),
        requested_period_label=f"06/2025 a {dates[-1]:%m/%Y}",
        monthly_cdi_rates_by_fund={fund.cnpj:() for fund in fund_objects},
        benchmark_spreads_by_fund={fund.cnpj:{} for fund in fund_objects},
        generated_at=datetime(2026,10,12,16,31))


def _build(kwargs):
    from services.portfolio_ppt_export import build_portfolio_pptx_bytes
    return build_portfolio_pptx_bytes(**kwargs)


def _tables(payload):
    presentation = Presentation(BytesIO(payload))
    return [shape.table for slide in presentation.slides for shape in slide.shapes if shape.has_table]


def _visible_text(payload):
    presentation = Presentation(BytesIO(payload))
    result = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_table:
                result.extend(cell.text for row in shape.table.rows for cell in row.cells)
            elif shape.has_text_frame:
                result.append(shape.text)
        result.append(slide.notes_slide.notes_text_frame.text)
    return '\n'.join(result)


def _walk_frames(value):
    if isinstance(value,pd.DataFrame):
        yield value
    elif is_dataclass(value):
        for field in fields(value):
            yield from _walk_frames(getattr(value,field.name))
    elif isinstance(value,dict):
        for item in value.values():
            yield from _walk_frames(item)
    elif isinstance(value,(tuple,list)):
        for item in value:
            yield from _walk_frames(item)


def _spec(kwargs):
    from services.portfolio_ppt_content import build_portfolio_deck_spec
    return build_portfolio_deck_spec(**kwargs)


def test_fund_comparison_distinguishes_zero_denominator_from_missing_coverage():
    kwargs = _inputs(months=1, funds=3)
    for fund, npl in zip(kwargs['funds'], (0.0, float('nan'), 1.0)):
        for mapping in (kwargs['outputs'].fund_monthly, kwargs['monitor_outputs'].fund_monitor):
            mapping[fund.cnpj]['npl_over90_ex360_pct'] = npl
            mapping[fund.cnpj]['npl_over90_ex360'] = npl * 1_000_000
            mapping[fund.cnpj]['pdd_npl_over90_ex360_pct'] = float('nan')
    tables = _tables(_build(kwargs))
    # One row per fund; the coverage column keeps N/A distinct from N/D.
    table = next(table for table in tables if table.cell(0, 0).text == 'Fundo')
    header = [cell.text for cell in table.rows[0].cells]
    column = header.index('Cobertura PDD / NPL > 90d (%)')
    assert [row.cells[column].text for row in list(table.rows)[1:]] == ['N/A · NPL zero', 'N/D', 'N/D']


def test_complete_monthly_values_are_preserved_beyond_old_twelve_month_cut_and_inputs_are_unchanged():
    kwargs = _inputs(months=13)
    source = kwargs['outputs'].fund_monthly['00000000000001'].copy(deep=True)
    spec = _spec(kwargs)
    frames = list(_walk_frames(spec))
    all_numbers = {float(value) for frame in frames for value in frame.to_numpy().ravel()
                   if isinstance(value,(int,float)) and pd.notna(value)}
    assert set(source['pl_total']).issubset(all_numbers)
    assert set(source['npl_over90_ex360_pct']).issubset(all_numbers)
    payload = _build(kwargs)
    text = _visible_text(payload)
    assert 'jun/25' in text and 'jun/26' in text
    # The first and last observations must be present in native table cells.
    cells = {cell.text for table in _tables(payload) for row in table.rows for cell in row.cells}
    money_cells={value.removeprefix('R$ ').removesuffix(' mm') for value in cells}
    # Four integer digits drop the decimals on screen; the exact values are
    # in the plan (checked above) and in the notes.
    assert '1.003' in money_cells
    assert '1.123' in money_cells
    assert 'R$ milhões' in text or 'R$ mm' in text
    pd.testing.assert_frame_equal(source,kwargs['outputs'].fund_monthly['00000000000001'])


def test_single_fund_uses_one_scope_and_retains_zero_absence_and_partial_states():
    kwargs = _inputs(months=13,missing=True)
    payload = _build(kwargs)
    text = _visible_text(payload)
    assert 'N/D' in text
    assert '0,00%' in text
    assert 'Cobertura parcial' in text and 'comparação inválida' in text
    npl_rows = [(table,row) for table in _tables(payload) for row in table.rows
                if row.cells[0].text=='NPL > 90d ex-360']
    assert npl_rows
    for table,row in npl_rows:
        headers = [cell.text for cell in table.rows[0].cells]
        if 'jun/25' in headers and 'jul/25' in headers:
            assert row.cells[headers.index('jun/25')].text == '0,00%'
            assert row.cells[headers.index('jul/25')].text == 'N/D'
    assert any(state in text for state in ('reported_zero','Zero reportado','zero reportado'))
    assert any(state in text for state in ('not_reported','Não reportado','não reportado','Não informado','não informado'))
    presentation = Presentation(BytesIO(payload))
    titles = [next((shape.text for shape in slide.shapes if shape.has_text_frame and shape.top<914400), '')
              for slide in presentation.slides]
    assert titles.count('Snapshot') == 1
    assert not any('Consolidado' in title for title in titles)
    chart_ns = {'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    with ZipFile(BytesIO(payload)) as archive:
        series = [ET.fromstring(archive.read(path)) for path in archive.namelist()
                  if path.startswith('ppt/charts/chart') and path.endswith('.xml')]
    npl = [root for root in series if any('90' in (item.findtext('c:v',namespaces=chart_ns) or '')
              for item in root.findall('.//c:tx//c:pt',chart_ns))]
    for root in npl:
        assert root.find('.//c:dispBlanksAs',chart_ns).attrib['val'] == 'gap'
        for serie in root.findall('.//c:ser',chart_ns):
            label=serie.findtext('c:tx/c:strRef/c:strCache/c:pt/c:v',namespaces=chart_ns)
            if label=='NPL > 90d ex-360':
                cache=serie.find('c:val/c:numRef/c:numCache',chart_ns)
                points={int(point.attrib['idx']):point.findtext('c:v',namespaces=chart_ns)
                        for point in cache.findall('c:pt',chart_ns)}
                assert float(points[0])==0.0
                assert 1 not in points


def test_twenty_three_return_series_and_all_months_remain_native_and_editable():
    kwargs = _inputs(series=23,months=13,missing=True)
    payload = _build(kwargs)
    cells = [cell.text for table in _tables(payload) for row in table.rows for cell in row.cells]
    for index in range(1,24):
        assert f'{index:02d}ª série' in cells
    text = _visible_text(payload)
    assert 'senior:00' in text and 'senior:22' in text
    assert 'jun/25' in cells and 'jun/26' in cells
    with ZipFile(BytesIO(payload)) as archive:
        chart_paths = [path for path in archive.namelist() if path.startswith('ppt/charts/') and path.endswith('.xml')]
        assert chart_paths
        workbooks = [path for path in archive.namelist() if path.startswith('ppt/embeddings/') and path.endswith('.xlsx')]
        assert workbooks
        for path in workbooks:
            workbook = load_workbook(BytesIO(archive.read(path)),data_only=False)
            assert any(isinstance(cell.value,(int,float)) for sheet in workbook for row in sheet for cell in row)


def test_all_five_funds_duration_survives_shared_plot_or_appendix():
    kwargs = _inputs(funds=5)
    payload = _build(kwargs)
    text = _visible_text(payload)
    for index in range(1,6):
        assert f'Fundo {index}' in text
        assert f'{index:014d}' in text
    frames = list(_walk_frames(_spec(kwargs)))
    numbers = {float(value) for frame in frames for value in frame.to_numpy().ravel()
               if isinstance(value,(int,float)) and pd.notna(value)}
    assert {6.0,7.0,8.0,9.0,10.0,10.9,11.2}.issubset(numbers)


@pytest.mark.parametrize('months',[1,2])
def test_short_period_constant_series_and_missing_lags_open_without_repair(months):
    payload = _build(_inputs(months=months,missing=True))
    presentation = Presentation(BytesIO(payload))
    assert len(presentation.slides)>=3
    assert 'jun/25' in _visible_text(payload)
    buffer = BytesIO(); presentation.save(buffer)
    assert len(Presentation(BytesIO(buffer.getvalue())).slides)==len(presentation.slides)
    with ZipFile(BytesIO(payload)) as archive:
        names = archive.namelist()
        assert len(names)==len(set(names))
        for path in names:
            if path.endswith('.xml'):
                ET.fromstring(archive.read(path))


def test_long_document_is_paginated_with_all_terms_sources_and_no_geometry_overflow():
    payload = _build(_inputs(long_document=True))
    text = _visible_text(payload)
    assert all(f'cláusula{index}' in text for index in range(220))
    assert 'Regulamento ID CVM 123, cláusula 7.1, página 18.' in text
    assert '...' not in text and '…' not in text
    presentation = Presentation(BytesIO(payload))
    for slide in presentation.slides:
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0
            assert shape.left+shape.width<=presentation.slide_width
            assert shape.top+shape.height<=presentation.slide_height


def test_identical_class_names_preserve_independent_class_keys():
    kwargs = _inputs(series=2)
    cnpj='00000000000001'
    kwargs['outputs'].fund_return_history[cnpj]['class_label'] = 'Sênior'
    kwargs['outputs'].fund_return_summary[cnpj]['class_label'] = 'Sênior'
    payload = _build(kwargs)
    text = _visible_text(payload)
    assert 'senior:00' in text and 'senior:01' in text
    # One zero-return class and another positive-return class retain both values.
    cells = {cell.text for table in _tables(payload) for row in table.rows for cell in row.cells}
    assert '0,00%' in cells and '1,01%' in cells


def test_internal_pct_name_uses_percent_unit_in_the_monthly_annex():
    kwargs = _inputs(months=2)
    kwargs['dashboard'] = SimpleNamespace(
        latest_competencia='07/2025', summary={},
        default_history_df=pd.DataFrame({'competencia':['06/2025','07/2025'],
                                        'provisao_pct_direitos':[17.25,18.10]}),
    )
    spec = _spec(kwargs)
    owners = [(block, row) for page in spec.slides for block in page.blocks
              if hasattr(block, 'frame') for _, row in block.frame.iterrows()
              if str(row.iloc[0]) == 'Provisão / direitos creditórios']
    assert owners and all(block.formats.get('jul/25') == 'percent:2' for block, _ in owners)
    assert '18,10%' in _visible_text(_build(kwargs))


def test_original_summary_cut_is_visible_and_does_not_use_another_month_delta():
    kwargs = _inputs(months=2)
    kwargs['dashboard'] = SimpleNamespace(
        latest_competencia='07/2025', summary={'cobertura_pct':200.0},
        default_history_df=pd.DataFrame({'competencia':['06/2025','07/2025'],
                                        'cobertura_pct':[90.0,100.0]}),
    )
    kwargs['dashboard_summary_competencia'] = '08/2025'
    spec = _spec(kwargs)
    snapshot = next(page for page in spec.slides if page.title == 'Snapshot').blocks[0]
    row = snapshot.frame[snapshot.frame['Indicador'].str.contains('Cobertura provisão')].iloc[0]
    assert 'ago/25' in row['Indicador'] and row['Valor'] == '200,0%'
    assert row['Var. mensal'] == 'N/D'
    assert kwargs['dashboard'].summary['cobertura_pct'] == 200.0


def test_positive_small_balance_is_not_displayed_as_zero_in_snapshot():
    kwargs = _inputs(months=1)
    cnpj = kwargs['funds'][0].cnpj
    kwargs['outputs'].fund_monthly[cnpj]['pl_total'] = 100.0
    kwargs['monitor_outputs'].fund_monitor[cnpj]['pl_total'] = 100.0
    snapshot = next(page for page in _spec(kwargs).slides if page.title == 'Snapshot').blocks[0]
    assert snapshot.frame.loc[snapshot.frame['Indicador'].eq('PL total'), 'Valor'].iloc[0] == 'R$ 100,00'


def test_cdi_benchmark_comparison_keeps_accumulated_spreads_and_bps_precision():
    from services.fidc_model.b3_cdi import B3CdiMonthlyRate
    from services.fund_return_matrix import build_fund_return_matrix, format_fund_return_matrix
    kwargs=_inputs(months=12,series=2)
    cnpj='00000000000001'
    dates=pd.date_range('2025-06-01',periods=12,freq='MS')
    rates=tuple(B3CdiMonthlyRate(mes=month.strftime('%Y-%m'),cdi_mensal=0.01,
        dias_uteis=20,data_inicio=month.date(),data_fim=(month+pd.offsets.MonthEnd(1)).date(),
        source='B3 fixture completa',expected_dias_uteis=20) for month in dates)
    kwargs['monthly_cdi_rates_by_fund']={cnpj:rates}
    kwargs['benchmark_spreads_by_fund']={cnpj:{'senior:00':0.035,'senior:01':0.05}}
    expected=format_fund_return_matrix(build_fund_return_matrix(kwargs['outputs'],cnpj,months=12,
        monthly_cdi_rates=rates,benchmark_spreads=kwargs['benchmark_spreads_by_fund'][cnpj]))
    payload=_build(kwargs)
    observed={cell.text for table in _tables(payload) for row in table.rows for cell in row.cells}
    for column in expected:
        if any(term in column for term in ('CDI','Spread','Gap')):
            assert column in observed
            assert set(expected[column]).issubset(observed)
    assert 'B3 fixture completa' in _visible_text(payload)


def test_return_calculation_lookback_does_not_expand_selected_visible_months():
    kwargs=_inputs(months=2,series=2)
    cnpj='00000000000001'
    history=kwargs['outputs'].fund_return_history[cnpj]
    older=[]
    for key,label in [('senior:00','01ª série'),('senior:01','02ª série')]:
        for timestamp in pd.date_range('2024-07-01','2025-05-01',freq='MS'):
            older.append(dict(competencia=timestamp.strftime('%m/%Y'),competencia_dt=timestamp,
                class_kind='senior',class_key=key,class_label=label,retorno_mensal_pct=9.99))
    kwargs['outputs'].fund_return_history[cnpj]=pd.concat([pd.DataFrame(older),history],ignore_index=True)
    payload=_build(kwargs)
    cells={cell.text for table in _tables(payload) for row in table.rows for cell in row.cells}
    assert 'jun/25' in cells and 'jul/25' in cells
    assert 'jul/24' not in cells and 'mai/25' not in cells
    assert '9,99%' not in cells
    # Source-provided accumulated returns retain their value despite the visual cut.
    assert '0,20%' in cells


def test_dense_table_pagination_rechecks_width_and_height_after_each_slice():
    from services.portfolio_ppt_content import DeckSpec,SlideSpec,TableBlock
    from services.portfolio_ppt_export import render_portfolio_deck
    rows=[]
    for index in range(8):
        rows.append({'Indicador':f'Indicador financeiro {index}',
            'Definição e denominador':' '.join(f'm{index}_p{term}' for term in range(90 if index%2 else 160)),
            'Fonte':'CVM Informe Mensal, tabela I; base canônica e faixa de atraso identificada.'})
    spec=DeckSpec('Carteira','Fundo · jun/26',(),
        (SlideSpec('Referências','Fundo · jun/26',(TableBlock('Definições',pd.DataFrame(rows)),),kind='appendix'),))
    payload=render_portfolio_deck(spec)
    text=_visible_text(payload)
    assert all(f'm{index}_p{term}' in text for index in range(8) for term in range(90 if index%2 else 160))
    presentation=Presentation(BytesIO(payload))
    assert len(presentation.slides)>2
    for slide in presentation.slides:
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0
            assert shape.left+shape.width<=presentation.slide_width
            assert shape.top+shape.height<=presentation.slide_height
