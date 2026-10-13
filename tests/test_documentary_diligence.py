import csv
from io import BytesIO, StringIO
import json
from hashlib import sha256
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from pptx import Presentation

from services.deep_dive_ppt_export import build_document_comparison_pptx_bytes
from services.documentary_diligence import (
    build_documentary_comparison_pages, criterion_display, criterion_has_answer,
    documentary_common_criteria, documentary_coverage, documentary_csv_bytes,
    documentary_manifest, documentary_table_html, documentary_themes,
    has_documentary_diligence, load_documentary_portfolio, public_document_payload, source_is_traceable,
    quarterly_documentary_table_html,
    load_documentary_fund, documentary_evidence_zip_bytes,
    source_display,
)
from services.portfolio_store import PortfolioFund, PortfolioRecord


def source(**changes):
    return {"document_id": "123", "title": "Regulamento consolidado", "url": "https://example.org/doc/123",
            "reference_date": "2026-10-01", "version": "2026-10-01", "page": 8,
            "clause": "4.2", "excerpt": "Subordinação mínima de 5% do PL", **changes}


def criterion(key="subordination", **changes):
    return {"criterion_id": key, "label": "Subordinação mínima", "theme": "protection", "applicability": "applicable",
            "status": "confirmed", "value": "≥ 5% do PL; apuração diária", "sources": [source()], "rule": {}, "risk_note": "", **changes}


def portfolio():
    return PortfolioRecord("selected", "Carteira", (PortfolioFund("54871427000194", "Fundo A"), PortfolioFund("60010416000112", "Fundo B")), "2026-10-01", "2026-10-01")


def funds():
    return [{"schema_version": 2, "cnpj": fund.cnpj, "name": fund.display_name, "reviewed_at": "2026-10-13T00:00:00-03:00", "criteria": [criterion(value="≥ 5% do PL" if index == 0 else "≥ 10% do PL")], "lookthrough": []} for index, fund in enumerate(portfolio().funds)]


def publish(tmp_path, fund):
    folder = tmp_path / "data/documentary_diligence/funds"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / f"{fund['cnpj']}.json").write_text(json.dumps(fund))


def test_coverage_preserves_unresolved_in_denominator_and_zero_as_answer():
    fund = {"cnpj": "x", "criteria": [criterion(value=0), criterion("unresolved", applicability="unresolved", status="extraction_pending", value=None, sources=[]), criterion("none", applicability="not_applicable", status="not_applicable", value=None, reason="Classe com cota única, sem subordinação.")]}
    coverage = documentary_coverage(fund)
    assert coverage == {"cnpj": "x", "total": 3, "applicable": 1, "not_applicable": 1, "unresolved": 1, "answered": 1, "denominator": 2, "ratio": .5, "pending": 1}
    assert criterion_display(fund["criteria"][0]) == "0"


def test_unsupported_not_applicable_is_an_unresolved_pending_cell():
    item = criterion(applicability="not_applicable", status="not_applicable", value=None, sources=[], reason="Cota única")
    assert documentary_coverage({"criteria": [item]})["denominator"] == 1
    assert documentary_coverage({"criteria": [item]})["unresolved"] == 1
    assert criterion_display(item) == "Aplicabilidade sem comprovação"


def test_corpus_or_source_without_locator_never_counts_as_documentary_answer():
    assert not source_is_traceable(source(page=None, clause=None))
    assert not source_is_traceable(source(title="Corpus acessível CNPJ: consulta sem documento"))
    assert not source_is_traceable(source(identity_verified=False))
    assert not source_is_traceable(source(reference_date="data desconhecida"))
    assert not source_is_traceable(source(reference_date="00/13/2026"))
    assert not source_is_traceable(source(page=0, clause=None))
    assert not criterion_has_answer(criterion(sources=[source(page=None, clause=None)]))
    assert not criterion_has_answer(criterion(value="Sem informação documental para remuneração"))
    assert not criterion_has_answer(criterion(value="-"))
    assert criterion_display(criterion(sources=[])) == "Fonte ou extração incompleta"


def test_missing_fund_file_retains_basket_column_and_union_criteria(tmp_path):
    publish(tmp_path, funds()[0])
    loaded = load_documentary_portfolio(portfolio(), tmp_path)
    assert has_documentary_diligence(portfolio(), tmp_path)
    assert [fund["cnpj"] for fund in loaded] == [f.cnpj for f in portfolio().funds]
    assert loaded[1]["criteria"][0]["status"] == "document_not_obtained"
    assert loaded[1]["criteria"][0]["applicability"] == "unresolved"
    assert documentary_coverage(loaded[1])["denominator"] == 1
    markup = documentary_table_html(documentary_themes(loaded)[0], loaded)
    assert markup.index("Fundo A") < markup.index("Fundo B")
    assert "Documento não obtido" in markup and "N/D" not in markup
    assert "3 fundos" not in markup and "Todos os fundos" not in markup


def test_default_loader_prunes_history_and_zip_preserves_complete_evidence(tmp_path):
    fund = funds()[0]
    body = "EVIDÊNCIA HISTÓRICA INTEGRAL " * 200
    fund["criteria"][0]["history"] = [{"value": body, "reference_date": "2025-01-01", "sources": [source(document_id=str(100000 + i), excerpt=body, local_path="/Users/analista/fidc/raw.pdf")]} for i in range(20)]
    publish(tmp_path, fund)
    loaded = load_documentary_fund(fund["cnpj"], source_root=tmp_path)
    assert "history" not in loaded["criteria"][0]
    assert loaded["criteria"][0]["history_summary"]["count"] == 20
    assert len(loaded["criteria"][0]["history_summary"]["sources"]) == 8
    assert body not in json.dumps(loaded, ensure_ascii=False)
    sidecar = tmp_path / "data/documentary_diligence/history" / f"{fund['cnpj']}.json"
    sidecar.parent.mkdir(parents=True)
    published_sidecar = public_document_payload({"entries": {"1": {"value": body}}, "passages": {"1": source(excerpt=body, local_path="/Users/analista/fidc/raw.pdf")}})
    sidecar.write_text(json.dumps(published_sidecar, ensure_ascii=False, indent=2) + "\n")
    expected_hash = sha256(sidecar.read_bytes()).hexdigest()
    fund["criteria"][0]["history_ref"] = {"path": f"history/{fund['cnpj']}.json", "sha256": expected_hash}
    publish(tmp_path, fund)
    payload = documentary_evidence_zip_bytes(portfolio(), [fund["cnpj"]], source_root=tmp_path)
    with ZipFile(BytesIO(payload)) as archive:
        exported = json.loads(archive.read(f"funds/{fund['cnpj']}.json"))
        assert exported["criteria"][0]["history"][0]["value"] == body.strip()
        assert f"history/{fund['cnpj']}.json" in archive.namelist()
        assert sha256(archive.read(f"history/{fund['cnpj']}.json")).hexdigest() == expected_hash
        assert b"/Users/" not in archive.read(f"funds/{fund['cnpj']}.json")


def test_corrupt_json_is_visible_as_parse_failure(tmp_path):
    folder = tmp_path / "data/documentary_diligence/funds"; folder.mkdir(parents=True)
    (folder / f"{portfolio().funds[0].cnpj}.json").write_text("invalid")
    publish(tmp_path, funds()[1])
    loaded = load_documentary_portfolio(portfolio(), tmp_path)
    assert loaded[0]["corpus_status"]["status"] == "parse_failed"
    assert loaded[0]["criteria"][0]["status"] == "parse_failed"


def test_stable_headers_across_themes_and_missing_rows_go_to_pending_records():
    basket = funds()
    for fund in basket:
        fund["criteria"].append(criterion("eligibility", theme="eligibility", label="Elegibilidade", value=None, status="clause_not_found", sources=[]))
    themes = documentary_themes(basket)
    assert documentary_table_html(themes[0], basket) == ""
    table = documentary_table_html(themes[1], basket)
    assert table.index("Fundo A") < table.index("Fundo B")
    rows = list(csv.DictReader(StringIO(documentary_csv_bytes(basket).decode("utf-8-sig"))))
    assert len(rows) == 4
    assert len([row for row in rows if row["Status"] == "Cláusula não localizada"]) == 2


def test_criterion_theme_typos_cannot_split_a_comparison_across_sections():
    basket = funds()
    basket[0]["criteria"][0]["theme"] = "Proteção por camada"
    basket[1]["criteria"][0]["theme"] = "Proteção porcamada"
    themes = documentary_themes(basket)
    assert len(themes) == 1 and themes[0].label == "Proteção por camada"
    assert len(themes[0].criteria) == 1
    html = documentary_table_html(themes[0], basket)
    assert "≥ 5% do PL" in html and "≥ 10% do PL" in html


def test_only_confirmed_identical_rules_are_consolidated():
    basket = funds()
    basket[1]["criteria"][0]["value"] = basket[0]["criteria"][0]["value"]
    assert list(documentary_common_criteria(basket)) == ["subordination"]
    assert documentary_table_html(documentary_themes(basket)[0], basket) == ""
    basket[1]["criteria"][0]["status"] = "document_not_obtained"
    assert documentary_common_criteria(basket) == {}


def test_identical_value_with_different_cure_is_not_a_common_rule():
    basket = funds()
    basket[1]["criteria"][0]["value"] = basket[0]["criteria"][0]["value"]
    basket[1]["criteria"][0]["rule"] = {"cure_days": 10}
    assert documentary_common_criteria(basket) == {}


def test_conflict_displays_both_documented_values_without_counting_coverage():
    item = criterion(status="version_conflict", value="15% numeral; 10% por extenso")
    assert "15% numeral; 10% por extenso" in criterion_display(item)
    assert "Conflito" in criterion_display(item)
    assert documentary_coverage({"criteria": [item]})["answered"] == 0


def test_public_download_keeps_document_evidence_without_private_paths():
    original = {"sources": [source(local_path="/Users/analista/fidc/data/raw/doc.pdf", text_path="/Users/analista/temp/doc.txt", sha256="abc")], "reason": "Consultar /Users/analista/fidc/data/raw/doc.pdf", "observed_value": 0}
    clean = public_document_payload(original)
    assert "local_path" not in clean["sources"][0] and "text_path" not in clean["sources"][0]
    assert clean["sources"][0]["sha256"] == "abc" and clean["sources"][0]["url"] == "https://example.org/doc/123"
    assert clean["reason"] == "Consultar data/raw/doc.pdf" and clean["observed_value"] == 0
    assert "local_path" in original["sources"][0]


def test_historical_investee_position_keeps_date_and_uncertainty_in_exports():
    basket = funds()
    relation = {"cnpj": "41970012000126", "status": "verified_position", "reference_date": "2025-12-31", "current_position_confirmed": False,
                "reason": "Posição declarada nas DFs de 31/12/2025; posição atual sem comprovação.", "sources": [source(reference_date="2025-12-31")]}
    detail = {"parent_cnpj": basket[0]["cnpj"], "parent_name": "FIC", "relation": relation, "fund": {**basket[0], "cnpj": relation["cnpj"], "name": "Investido"}}
    rows = list(csv.DictReader(StringIO(documentary_csv_bytes(basket, lookthrough=[detail]).decode("utf-8-sig"))))
    exported = json.loads(rows[-1]["Vínculo documental"])
    assert exported["reference_date"] == "2025-12-31" and exported["current_position_confirmed"] is False
    assert exported["sources"][0]["document_id"] == "123"
    with patch("services.documentary_diligence.load_documentary_lookthrough", return_value=[detail]):
        pages = build_documentary_comparison_pages(portfolio(), basket)
    assert any("current_position_confirmed" in text and "false" in text and "2025-12-31" in text for text in pages[-1].sources)


def test_sources_and_values_escape_markup_and_reject_script_url():
    basket = funds()
    basket[0]["criteria"][0]["sources"] = [source(url="javascript:alert(1)", excerpt="<script>bad</script>")]
    basket[0]["criteria"][0]["value"] = "<svg onload='bad'>5%"
    html = documentary_table_html(documentary_themes(basket)[0], basket)
    assert "javascript:" not in html and "<script>" not in html and "<svg" not in html
    assert "&lt;script&gt;" in html and "ID 123" in html and "2026-10-01" in html and "p. 8" in html


def test_ocr_source_preserves_automatic_method_with_original_pdf_reference():
    basket = funds()
    extracted = source(parse_method="apple_vision_ocr", verification="ocr_primary_passage")
    basket[0]["criteria"][0]["sources"] = [extracted]
    assert "Texto por OCR local; fonte PDF original" in source_display(extracted)
    html = documentary_table_html(documentary_themes(basket)[0], basket)
    assert "p. 8 · OCR local" in html and "https://example.org/doc/123" in html
    pages = build_documentary_comparison_pages(portfolio(), basket)
    assert any("Texto por OCR local; fonte PDF original" in text for text in pages[0].sources)
    visual = {**extracted, "verification": "visual_directed"}
    assert "Revisão visual da página; fonte PDF original" in source_display(visual)
    assert "Texto por OCR local" not in source_display(visual)


def test_ppt_preserves_typed_absence_and_all_fund_columns():
    basket = funds()
    basket[1]["criteria"][0].update(status="document_not_obtained", value=None, sources=[])
    manifest = documentary_manifest(portfolio(), basket, "06/2026")
    pages = build_documentary_comparison_pages(portfolio(), basket)
    deck = Presentation(BytesIO(build_document_comparison_pptx_bytes(manifest, pages, portfolio_mode=True, compact=False)))
    table = next(shape.table for shape in deck.slides[0].shapes if shape.has_table)
    assert len(table.columns) == 3
    assert "Documento não obtido" in table.cell(1, 2).text
    assert "ID 123" in deck.slides[0].notes_slide.notes_text_frame.text
    assert "CNPJ 54871427000194" in deck.slides[0].notes_slide.notes_text_frame.text


def test_ambiguous_rule_keeps_aggregate_ime_reference_separate_in_all_exports():
    basket = funds()
    observed = {"criterion_id": "subordination", "comparison_required": True, "observed_value": None,
                "context_value": 12.5, "context_unit": "percent", "context_formula": "PL subordinado / PL total × 100",
                "competence": "08/2026", "status": "not_comparable", "reason": "O regulamento exige outra base por subclasse.",
                "source": {"type": "ime_document", "document_id": "789", "competence": "08/2026"}}
    observations = {basket[0]["cnpj"]: {"subordination": observed}}
    markup = documentary_table_html(documentary_themes(basket)[0], basket, observations)
    assert "Referência agregada: 12,50%" in markup and "Comparação não reproduzível" in markup
    assert "PL subordinado / PL total × 100" in markup
    records = list(csv.DictReader(StringIO(documentary_csv_bytes(basket, observations).decode("utf-8-sig"))))
    assert "Referência agregada: 12,50%" in records[0]["IME"]
    assert json.loads(records[0]["Base IME"])["observed_value"] is None
    pages = build_documentary_comparison_pages(portfolio(), basket, observations)
    assert "Referência agregada: 12,50%" in pages[0].frame.iloc[0, 1]
    assert any("context_formula" in note for note in pages[0].sources)


def test_verified_header_inversion_exposes_its_evidence_without_changing_basket_identity():
    basket = funds()
    mapping = {"class_cnpj": basket[0]["cnpj"], "fund_cnpj": "63953619000130", "effective_from": "2025-12-05",
               "allowed_header_pair": {"NR_CNPJ_FUNDO": basket[0]["cnpj"], "NR_CNPJ_CLASSE": "63953619000130"},
               "reason": "Cadastro CVM vincula os CNPJs por registro exato.", "sources": [source(url="https://dados.cvm.gov.br/registro.zip")]}
    observed = {"comparison_required": True, "competence": "08/2026", "status": "not_comparable",
                "source": {"cnpj": basket[0]["cnpj"], "identity_basis": "verified_fund_class_inversion", "identity_mapping": mapping,
                           "warnings": ["Inversão exata comprovada."], "field_paths": {"net_asset_value": ["PATRLIQ/VL_PATRIM_LIQ"]}}}
    html = documentary_table_html(documentary_themes(basket)[0], basket, {basket[0]["cnpj"]: {"subordination": observed}})
    assert "Inversão fundo/classe comprovada" in html and "Inversão exata comprovada." in html
    assert "https://dados.cvm.gov.br/registro.zip" in html and "PATRLIQ/VL_PATRIM_LIQ" in html
    assert html.index("Fundo A") < html.index("Fundo B")


def test_quarterly_declaration_keeps_its_period_sources_and_separate_coverage():
    basket = funds()
    before = documentary_coverage(basket[0])
    basket[0]["quarterly_observations"] = {"status": "parsed", "quarter": "2026-Q2", "competence": "06/2026",
                                         "source": source(identity_verified=True), "records": [{"criterion_id": "credit_policy_changes", "label": "Alteração da concessão",
                                         "value": "Houve alteração da política de concessão", "status": "declared", "source": source(page=4, identity_verified=True)}]}
    html = quarterly_documentary_table_html(basket)
    assert html.index("Fundo A") < html.index("Fundo B") and "2026-Q2" in html
    assert "Houve alteração da política de concessão" in html and "Informe trimestral não disponível" in html
    assert "p. 4" in html and "Resposta declarada no informe" in html
    assert documentary_coverage(basket[0]) == before
    records = list(csv.DictReader(StringIO(documentary_csv_bytes(basket, lookthrough=[]).decode("utf-8-sig"))))
    quarter = next(row for row in records if row["Camada"] == "Declaração trimestral")
    assert quarter["Cobertura válida"] == "False"
    pages = build_documentary_comparison_pages(portfolio(), basket, include_lookthrough=False)
    assert pages[-1].title == "Declarações trimestrais" and "2026-Q2" in pages[-1].frame.iloc[0, 1]


def test_quarterly_conflict_preserves_declaration_and_both_sources():
    basket = funds()
    quarter_source = source(document_id="456", title="Informe Trimestral", page=3, identity_verified=True)
    rule_source = source(document_id="123", page=19)
    basket[0]["quarterly_observations"] = {"status": "reported", "quarter": "2026-Q2", "source": quarter_source,
        "records": [{"criterion_id": "precatórios", "label": "Aquisição de precatórios", "status": "documentary_conflict", "value": "Informe declara aquisição de precatórios federais", "risk_note": "A norma veda esses créditos; operação e vigência não comprovadas.", "source": quarter_source, "sources": [quarter_source, rule_source]}]}
    html = quarterly_documentary_table_html(basket)
    assert "Informe declara aquisição" in html and "Declaração em conflito" in html
    assert "ID 456" in html and "ID 123" in html and "operação e vigência não comprovadas" in html
    rows = list(csv.DictReader(StringIO(documentary_csv_bytes(basket, lookthrough=[]).decode("utf-8-sig"))))
    quarter = next(row for row in rows if row["Camada"] == "Declaração trimestral")
    assert quarter["Status"] == "Declaração em conflito com o regulamento" and len(json.loads(quarter["Fontes"])) == 2


def test_v2_tab_works_without_legacy_manifest():
    from tabs import tab_deep_dive
    with patch.object(tab_deep_dive, "list_deep_dives", return_value=[]), patch("services.documentary_diligence.has_documentary_diligence", return_value=True), patch("services.documentary_diligence.render_documentary_diligence") as render:
        tab_deep_dive.render_tab_deep_dive(selected_portfolio=portfolio(), show_portfolio_selector=False, compact=True, portfolio_mode=True, reference_competence="06/2026", dashboards={})
    render.assert_called_once_with(portfolio(), reference_competence="06/2026", dashboards={})
