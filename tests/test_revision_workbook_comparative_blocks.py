"""Authored XLSX detail and summary tables have separate comparative contracts."""
from io import BytesIO
import json

from openpyxl import Workbook, load_workbook
import pytest

from scripts.publish_fidc_revision_bundle import _publication_staging_directory
from services.industry_comparative_period import ComparisonCut
from services.industry_revision_export import (
    RevisionExportUnavailable,
    _validate_workbook_offer_periods,
)


SPECS = (
    ("Comparativo renda fixa", "fixed_income_offer_comparison"),
    ("Regime de colocação", "closed_offer_placement_regime"),
    ("Histograma ofertas", "closed_offer_ticket_distribution"),
    ("Top 15 ofertas", "closed_offer_top15"),
    ("Validação emissões", "market_offer_reconciliation"),
    ("Público-alvo ofertas", "offer_target_public_shares"),
)
SUMMARY_SPECS = (
    ("Período", "period_label"), ("Ofertas no período", "period_closed_offers"),
    ("Volume do período", "period_registered_volume_brl"), ("Subtotal Top 15", "top15_registered_volume_brl"),
    ("% do total", "top15_share_of_period_volume"), ("IBBA líder · ofertas", "ibba_lead_offers_top15"),
    ("IBBA líder · volume", "ibba_lead_volume_top15_brl"), ("Garantia firme · ofertas", "firm_commitment_offers_top15"),
    ("Garantia firme · volume", "firm_commitment_volume_top15_brl"), ("% rito automático · volume", "automatic_rite_registered_volume_share"),
    ("Comparabilidade", "comparability_status"),
)


def _workbook_and_payload():
    cut = ComparisonCut(2026, 8)
    periods = ["2022 FY parcial", "2023 FY", "2024 FY", "2025 FY", cut.period_id()]
    payload = {"latest_complete": cut.competence, "offers_comparison_meta": cut.to_meta(),
               "anbima_market_offers_manifest": {"comparison_meta": cut.to_meta()}}
    book = Workbook()
    book.remove(book.active)
    for sheet_name, key in SPECS:
        rows = [{"period_label": period} for period in periods]
        if key == "closed_offer_top15":
            rows = [{"period_label": period} for index, period in enumerate(periods) for _ in range(7 if index == 0 else 15)]
        payload[key] = rows
        sheet = book.create_sheet(sheet_name)
        sheet.append([sheet_name]); sheet.append(["Até agosto"]); sheet.append([])
        sheet.append(["Período", "Valor"])
        for row in rows:
            sheet.append([row["period_label"], 1])
        if key != "closed_offer_top15":
            sheet.cell(7 + len(rows), 1, "Fonte: CVM; nota metodológica após a tabela.")
    summaries = []
    for index, period in enumerate(periods, 1):
        row = {"period_label": period, "period_order": index, "period_closed_offers": index * 100,
               "period_registered_volume_brl": index * 10_000_000_000.123, "top15_registered_volume_brl": index * 5_000_000_000.123,
               "top15_share_of_period_volume": 0.5, "ibba_lead_offers_top15": 2,
               "ibba_lead_volume_top15_brl": 1_000_000_000.12, "firm_commitment_offers_top15": 0,
               "firm_commitment_volume_top15_brl": 0.0, "automatic_rite_registered_volume_share": 0.75,
               "comparability_status": "parcial_não_comparável" if index == 1 else "comparável"}
        summaries.append(row)
    payload["closed_offer_top15_summary"] = summaries
    sheet = book["Top 15 ofertas"]
    assert len(payload["closed_offer_top15"]) == 67
    header_row = 74
    for column, (header, _) in enumerate(SUMMARY_SPECS, 1):
        sheet.cell(header_row, column, header)
    for index, row in enumerate(summaries, header_row + 1):
        for column, (_, key) in enumerate(SUMMARY_SPECS, 1):
            sheet.cell(index, column, row[key])
    sheet.cell(82, 1, "Nota: 2022 parcial, com sete ofertas; fonte documental mantida.")
    payload["issuance_taxonomy_table"] = [{"Categoria": "Financeiro", "jan–ago/26 (R$ bi)": 1.0}]
    book.create_sheet("Emissões por categoria").append([])
    for column, header in enumerate(payload["issuance_taxonomy_table"][0], 1):
        book["Emissões por categoria"].cell(4, column, header)
    book.create_sheet("Cedentes · presença").cell(4, 1, "PL ago/26 (R$)")
    return book, payload


def _roundtrip(book):
    buffer = BytesIO(); book.save(buffer); book.close()
    return load_workbook(BytesIO(buffer.getvalue()), read_only=True, data_only=True)


def test_real_openpyxl_detail_summary_and_notes_validate_separately():
    book, payload = _workbook_and_payload()
    book = _roundtrip(book)
    try:
        _validate_workbook_offer_periods(book, payload)
        assert book["Top 15 ofertas"]["A74"].value == "Período"
        assert book["Top 15 ofertas"]["A79"].value == "2026 jan-ago"
    finally:
        book.close()


@pytest.mark.parametrize("mutation", ["detail_period", "summary_period", "summary_count", "summary_volume", "summary_header", "summary_extra", "detail_extra", "stale_meta", "secondary_cut", "stale_taxonomy_header", "stale_presence_header"])
def test_comparative_blocks_reject_divergent_periods_and_summaries(mutation):
    book, payload = _workbook_and_payload()
    if mutation == "detail_period": book["Top 15 ofertas"]["A71"] = "2026 jan-jun"
    elif mutation == "summary_period": book["Top 15 ofertas"]["A79"] = "2026 jan-jun"
    elif mutation == "summary_count": book["Top 15 ofertas"]["B79"] = 999
    elif mutation == "summary_volume": book["Top 15 ofertas"]["D79"] = 123.0
    elif mutation == "summary_header": book["Top 15 ofertas"]["C74"] = "Outro volume"
    elif mutation == "summary_extra": book["Top 15 ofertas"]["A80"] = "2026 jan-ago"
    elif mutation == "detail_extra": book["Público-alvo ofertas"]["A10"] = "2026 jan-ago"
    elif mutation == "stale_meta": payload["offers_comparison_meta"]["current_period_end"] = "2026-06-30"
    elif mutation == "stale_taxonomy_header": book["Emissões por categoria"]["B4"] = "jan–jun/26 (R$ bi)"
    elif mutation == "stale_presence_header": book["Cedentes · presença"]["A4"] = "PL jun/26 (R$)"
    else: payload["anbima_market_offers_manifest"]["comparison_meta"] = ComparisonCut(2026, 6).to_meta()
    book = _roundtrip(book)
    try:
        with pytest.raises(RevisionExportUnavailable):
            _validate_workbook_offer_periods(book, payload)
    finally:
        book.close()


def test_publication_stage_survives_failure_with_all_completed_exports(tmp_path):
    stage = None
    with pytest.raises(RuntimeError, match="validation failed"):
        with _publication_staging_directory(tmp_path) as stage:
            (stage / "industry_data_revised.xlsx").write_bytes(b"completed Office bytes")
            raise RuntimeError("validation failed")
    assert stage.is_dir()
    assert (stage / "industry_data_revised.xlsx").read_bytes() == b"completed Office bytes"
    assert json.loads((stage / "publication_failure.json").read_text()) == {"error_type": "RuntimeError", "error": "validation failed"}


def test_successful_publication_stage_is_removed(tmp_path):
    with _publication_staging_directory(tmp_path) as stage:
        (stage / "complete.txt").write_text("success")
    assert not stage.exists()
