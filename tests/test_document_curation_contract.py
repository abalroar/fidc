from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path

import pandas as pd
from pptx import Presentation
import pytest

from services.document_curation_contract import (
    DocumentCurationContractError,
    EXPORT_SOURCE_FILE,
    REQUIRED_COMPARISON_IDS,
    finalize_document_curation,
    find_document_curation_package,
    validate_document_curation_contract,
)
from services.portfolio_store import PortfolioFund, PortfolioRecord, portfolio_basket_signature


READING_AT = "2026-10-07T11:56:39-03:00"


@pytest.fixture
def reviewed_package(tmp_path: Path):
    root = tmp_path / "deep_dives"
    package = root / "portfolio_a"
    (package / "tables").mkdir(parents=True)
    (package / "evidence").mkdir()
    portfolio = PortfolioRecord(
        id="portfolio-a", name="Carteira A",
        funds=(PortfolioFund("11111111000111", "Fundo completo A"), PortfolioFund("22222222000122", "Fundo completo B")),
        created_at=READING_AT, updated_at=READING_AT,
    )
    payload = {
        "deep_dive_id": "portfolio_a", "title": "Carteira A",
        "portfolio_id": portfolio.id, "portfolio_signature": portfolio_basket_signature(portfolio.funds),
        "generated_at": READING_AT, "source": "CVM / Fundos.NET",
        "funds": [{"cnpj": fund.cnpj, "name": fund.display_name, "short_name": fund.display_name} for fund in portfolio.funds],
        "tables": [], "comparison_columns": {"A": portfolio.funds[0].cnpj, "B": portfolio.funds[1].cnpj},
        "custom_metadata": {"preserve": True},
    }
    (package / "manifest.json").write_text(json.dumps(payload), encoding="utf-8")
    rows = []
    for table_id in REQUIRED_COMPARISON_IDS:
        pd.DataFrame({"Critério": ["Regra"], "A": ["15% do PL"], "B": ["Limite não localizado nos documentos acessíveis"]}).to_csv(package / f"tables/{table_id}.csv", index=False)
        rows.extend([
            {"Tabela": table_id, "Critério": "Regra", "CNPJ": portfolio.funds[0].cnpj, "Valor": "15% do PL", "Fonte": "Regulamento, ID CVM 101, p. 7", "Nota": ""},
            {"Tabela": table_id, "Critério": "Regra", "CNPJ": portfolio.funds[1].cnpj, "Valor": "Limite não localizado nos documentos acessíveis", "Fonte": "Lacuna documental: regulamento não acessível para este CNPJ", "Nota": ""},
            {"Tabela": table_id, "Critério": "Nota", "CNPJ": "Carteira", "Valor": "Qualificação documental", "Fonte": "", "Nota": "A regra exige verificação documental."},
        ])
    pd.DataFrame(rows).to_csv(package / "evidence/comparison_sources.csv", index=False)
    pd.DataFrame({"Tema": ["Recebíveis", "Gatilhos", "Pagamentos"], "Conclusão": ["CCBs de veículos", "15% do PL", "Calendário fechado não localizado"]}).to_csv(package / "tables/key_findings.csv", index=False)
    index = {"private_index_metadata": "preserve", "deep_dives": [{"deep_dive_id": "other", "custom": {"value": 7}, "generated_at": "2000-01-01"}]}
    (root / "index.json").write_text(json.dumps(index), encoding="utf-8")
    return root, package, portfolio


def _snapshot(root: Path) -> dict[str, bytes]:
    return {path.relative_to(root).as_posix(): path.read_bytes() for path in root.rglob("*") if path.is_file()}


def _manifest(package: Path) -> dict:
    return json.loads((package / "manifest.json").read_text())


def _write_manifest(package: Path, payload: dict) -> None:
    (package / "manifest.json").write_text(json.dumps(payload), encoding="utf-8")


def test_initial_check_registers_canonical_specs_in_memory_and_keeps_files_unchanged(reviewed_package):
    root, package, portfolio = reviewed_package
    before = _snapshot(root)
    result = finalize_document_curation(package, portfolio, check_only=True)
    assert result.check_only and not result.export_exists
    assert result.table_count == 7 and result.cell_count == 14 and result.slide_count == 7
    assert _snapshot(root) == before
    validated = validate_document_curation_contract(package, portfolio)
    assert [spec.id for spec in validated.manifest.tables if spec.kind == "document_comparison"] == list(REQUIRED_COMPARISON_IDS)


def test_missing_fund_column_is_rejected_without_any_mutation(reviewed_package):
    root, package, portfolio = reviewed_package
    path = package / f"tables/{REQUIRED_COMPARISON_IDS[0]}.csv"
    frame = pd.read_csv(path, keep_default_na=False).drop(columns="B")
    frame.to_csv(path, index=False)
    before = _snapshot(root)
    with pytest.raises(DocumentCurationContractError, match="chaves exatas dos cabeçalhos"):
        finalize_document_curation(package, portfolio)
    assert _snapshot(root) == before


def test_missing_comparison_theme_is_rejected_without_fallback(reviewed_package):
    _, package, portfolio = reviewed_package
    (package / "tables/comparison_payments.csv").unlink()
    with pytest.raises(DocumentCurationContractError, match="comparison_payments.csv"):
        validate_document_curation_contract(package, portfolio)


def test_evidence_value_must_match_the_exact_cell(reviewed_package):
    _, package, portfolio = reviewed_package
    path = package / "evidence/comparison_sources.csv"
    evidence = pd.read_csv(path, dtype=str, keep_default_na=False)
    evidence.loc[0, "Valor"] = "20% do PL"
    evidence.to_csv(path, index=False)
    with pytest.raises(DocumentCurationContractError, match="Valor diverge"):
        validate_document_curation_contract(package, portfolio)


@pytest.mark.parametrize("value", ["", "N/D", "NaN", "-", "—"])
def test_implicit_gaps_are_rejected(reviewed_package, value):
    _, package, portfolio = reviewed_package
    path = package / "tables/comparison_costs.csv"
    frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    frame.loc[0, "B"] = value
    frame.to_csv(path, index=False)
    with pytest.raises(DocumentCurationContractError, match="lacuna implícita"):
        validate_document_curation_contract(package, portfolio)


def test_explicit_documentary_gap_and_common_notes_are_accepted(reviewed_package):
    _, package, portfolio = reviewed_package
    validated = validate_document_curation_contract(package, portfolio)
    assert validated.pages[0].frame.iloc[0]["B"] == "Limite não localizado nos documentos acessíveis"
    assert validated.pages[0].notes == ("A regra exige verificação documental.",)
    assert "Lacuna documental" in validated.pages[0].sources[1]


def test_common_note_does_not_replace_the_cnpj_specific_evidence(reviewed_package):
    _, package, portfolio = reviewed_package
    path = package / "evidence/comparison_sources.csv"
    evidence = pd.read_csv(path, dtype=str, keep_default_na=False).drop(index=0)
    evidence.to_csv(path, index=False)
    with pytest.raises(DocumentCurationContractError, match="Célula sem evidência"):
        validate_document_curation_contract(package, portfolio)


def test_short_label_without_explicit_mapping_is_rejected(reviewed_package):
    _, package, portfolio = reviewed_package
    payload = _manifest(package)
    del payload["comparison_columns"]
    _write_manifest(package, payload)
    with pytest.raises(DocumentCurationContractError, match="comparison_columns"):
        validate_document_curation_contract(package, portfolio)


def test_all_comparisons_keep_the_same_column_order(reviewed_package):
    _, package, portfolio = reviewed_package
    path = package / "tables/comparison_protection.csv"
    frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    frame[["Critério", "B", "A"]].to_csv(path, index=False)
    with pytest.raises(DocumentCurationContractError, match="mesma ordem de fundos"):
        validate_document_curation_contract(package, portfolio)


def test_explicit_mapping_cannot_silently_ignore_an_extra_fund_label(reviewed_package):
    _, package, portfolio = reviewed_package
    payload = _manifest(package)
    payload["comparison_columns"]["Outro fundo"] = "33333333000133"
    _write_manifest(package, payload)
    with pytest.raises(DocumentCurationContractError, match="chaves exatas dos cabeçalhos"):
        validate_document_curation_contract(package, portfolio)


def test_signature_and_fund_identity_must_match_saved_portfolio(reviewed_package):
    _, package, portfolio = reviewed_package
    payload = _manifest(package)
    payload["portfolio_signature"] = "wrong-basket"
    _write_manifest(package, payload)
    with pytest.raises(DocumentCurationContractError, match="Assinatura"):
        validate_document_curation_contract(package, portfolio)
    payload["portfolio_signature"] = portfolio_basket_signature(portfolio.funds)
    payload["funds"].pop()
    _write_manifest(package, payload)
    with pytest.raises(DocumentCurationContractError, match="CNPJs do manifest"):
        validate_document_curation_contract(package, portfolio)


def test_finalization_preserves_reading_date_and_other_index_entries(reviewed_package):
    root, package, portfolio = reviewed_package
    original_index = json.loads((root / "index.json").read_text())
    result = finalize_document_curation(package, portfolio, output_root=root)
    manifest = _manifest(package)
    index = json.loads((root / "index.json").read_text())
    assert result.generated_at == manifest["generated_at"] == READING_AT
    assert manifest["custom_metadata"] == {"preserve": True}
    assert manifest["documentary_export"]["sha256"] == result.pptx_sha256
    assert index["deep_dives"][0] == original_index["deep_dives"][0]
    assert index["private_index_metadata"] == "preserve"
    assert index["deep_dives"][1]["portfolio_id"] == portfolio.id
    presentation = Presentation(BytesIO((package / EXPORT_SOURCE_FILE).read_bytes()))
    assert len(presentation.slides) == 7
    assert all(any(shape.has_table for shape in slide.shapes) for slide in presentation.slides)
    assert "Regulamento, ID CVM 101, p. 7" in presentation.slides[0].notes_slide.notes_text_frame.text
    check = finalize_document_curation(package, portfolio, check_only=True)
    assert check.export_exists and check.pptx_sha256 == result.pptx_sha256


def test_reading_date_changes_only_when_explicitly_provided(reviewed_package):
    _, package, portfolio = reviewed_package
    new_date = "2026-10-08T09:30:00-03:00"
    result = finalize_document_curation(package, portfolio, reading_at=new_date)
    assert result.generated_at == _manifest(package)["generated_at"] == new_date
    with pytest.raises(DocumentCurationContractError, match="fuso horário"):
        validate_document_curation_contract(package, portfolio, reading_at="2026-10-08T09:30:00")


def test_check_detects_stale_export_after_a_cell_and_its_evidence_change(reviewed_package):
    _, package, portfolio = reviewed_package
    finalize_document_curation(package, portfolio)
    path = package / "tables/comparison_protection.csv"
    frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    frame.loc[0, "A"] = "20% do PL"
    frame.to_csv(path, index=False)
    evidence_path = package / "evidence/comparison_sources.csv"
    evidence = pd.read_csv(evidence_path, dtype=str, keep_default_na=False)
    evidence.loc[evidence["Tabela"].eq("comparison_protection") & evidence["CNPJ"].eq("11111111000111"), "Valor"] = "20% do PL"
    evidence.to_csv(evidence_path, index=False)
    with pytest.raises(DocumentCurationContractError, match="PPTX diverge dos comparativos"):
        finalize_document_curation(package, portfolio, check_only=True)


def test_check_detects_modified_pptx_hash_and_missing_registered_export(reviewed_package):
    _, package, portfolio = reviewed_package
    finalize_document_curation(package, portfolio)
    export = package / EXPORT_SOURCE_FILE
    export.write_bytes(export.read_bytes() + b"modified")
    with pytest.raises(DocumentCurationContractError, match="Hash do PPTX diverge"):
        finalize_document_curation(package, portfolio, check_only=True)
    export.unlink()
    with pytest.raises(DocumentCurationContractError, match="PPTX registrado"):
        finalize_document_curation(package, portfolio, check_only=True)


def test_invalid_key_findings_is_rejected(reviewed_package):
    _, package, portfolio = reviewed_package
    path = package / "tables/key_findings.csv"
    pd.DataFrame({"Tema": ["Recebíveis"], "Conclusão": ["CCBs"]}).to_csv(path, index=False)
    with pytest.raises(DocumentCurationContractError, match="3 a 5 conclusões"):
        validate_document_curation_contract(package, portfolio)


def test_unrelated_malformed_manifest_does_not_block_target_lookup(reviewed_package):
    root, package, portfolio = reviewed_package
    unrelated = root / "unrelated"
    unrelated.mkdir()
    (unrelated / "manifest.json").write_text("{broken JSON", encoding="utf-8")
    assert find_document_curation_package(portfolio, root) == package


def test_finalized_check_requires_specs_actually_consumed_by_the_site(reviewed_package):
    _, package, portfolio = reviewed_package
    finalize_document_curation(package, portfolio)
    payload = _manifest(package)
    payload["tables"] = [spec for spec in payload["tables"] if spec["id"] != "comparison_protection"]
    _write_manifest(package, payload)
    with pytest.raises(DocumentCurationContractError, match="a tela carregaria um contrato diferente"):
        finalize_document_curation(package, portfolio, check_only=True)


def test_finalized_check_detects_changed_input_bytes_even_with_identical_rendering(reviewed_package):
    _, package, portfolio = reviewed_package
    # Change only unrelated input formatting: rendered data remains identical,
    # while a persisted deck must still identify its precise reviewed inputs.
    finalize_document_curation(package, portfolio)
    path = package / "tables/key_findings.csv"
    path.write_text(path.read_text() + "\n", encoding="utf-8")
    with pytest.raises(DocumentCurationContractError, match="Digest dos inputs diverge"):
        finalize_document_curation(package, portfolio, check_only=True)
