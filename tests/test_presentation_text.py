from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
from pptx import Presentation
import pytest

from services.deep_dive_models import DeepDiveManifest
from services.deep_dive_ppt_export import build_document_comparison_pptx_bytes
from services.document_curation_comparison import DocumentComparisonPage
from services.presentation_text import document_comparison_value, public_document_text, public_pptx_bytes


@pytest.mark.parametrize("text", [
    None, "", "N/D", "Não localizado nas fontes acessíveis desta leitura",
    "Prazo remetido ao suplemento de cada série.",
    "Cronograma fechado não identificado; modelo de apêndice não preenchido",
    "Taxa remetida ao anexo; percentual/mínimo não localizados na parte geral.",
    "Sobretaxa definida em bookbuilding", "Sem competência IME carregada em cache local",
    "Prazo da série por suplemento", "Taxas por suplemento", "Por apêndice",
    "Índice DI + sobretaxa remetida ao apêndice", "dado não certificado",
    "Sênior: CDI + spread definido emsuplemento",
    "Finalidade e limite operacional de hedge não localizados; referências genéricas não autorizam conclusão de uso",
    "Spread final da 2ª série não localizado; atos 1325817/1324299 inacessíveisHTTP 520",
    "Remetido ao CCCB. Taxa final não localizada",
])
def test_unanswered_cells_use_hyphen(text):
    assert document_comparison_value(text) == "-"


@pytest.mark.parametrize(("text", "expected"), [
    ("0,05% a.a./PL; mínimo mensal não identificado.", "0,05% a.a./PL"),
    ("VNU inicial de R$ 1.000; volume efetivo não localizado", "VNU inicial de R$ 1.000"),
    ("Incluída na taxa de administração; tarifa isolada não confirmada", "Incluída na taxa de administração"),
    ("Cronograma não identificado; amortização pro rata/sequencial em regime de caixa", "amortização pro rata/sequencial em regime de caixa"),
    ("DF set/25: cotas sem rating. Atualização não localizada", "DF set/25: cotas sem rating."),
    ("Efetiva não confirmada.25/06/2026 é indicativa", "25/06/2026 é indicativa"),
    ("Efetiva não confirmada. 25/06/2026 é indicativa", "25/06/2026 é indicativa"),
    ("Emissão autorizada de subordinadas até R$ 1 bi é autorização, com volume efetivo/VNU não confirmados.", "Emissão autorizada de subordinadas até R$ 1 bi é autorização"),
    ("PL/cotas sênior≥125%; índice ponderado sênior≥1,00, fator específico remetido ao apêndice", "PL/cotas sênior≥125%; índice ponderado sênior≥1,00"),
    ("Sênior, mezanino e júnior; suplementos preenchidos no regulamento de 2024", "Sênior, mezanino e júnior (regulamento de 2024)"),
    ("2ª emissão:1 cota sênior deR$ 150 mi; consolidação aprovada, volume final consolidado não localizado", "2ª emissão:1 cota sênior deR$ 150 mi; consolidação aprovada"),
    ("Sênior/A/B: pagamento mensal no 7º dia útil; amortização após carência segue item 14.2 do regulamento ausente.", "Sênior/A/B: pagamento mensal no 7º dia útil; amortização após carência"),
    ("Proxy por cotas. Confirmar divergência15%/dez", "Proxy por cotas; Divergência15%/dez"),
    ("Maior entre 90 dias de despesas e 0,1% do PL; validar divergência com extenso", "Maior entre 90 dias de despesas e 0,1% do PL; Divergência com extenso"),
    ("Classe fechada; séries têm prazos próprios nos suplementos; público profissional nos termos do regulamento.", "Classe fechada; público profissional nos termos do regulamento."),
])
def test_partial_answers_keep_material_terms(text, expected):
    assert document_comparison_value(text) == expected


@pytest.mark.parametrize("text", [
    "Sem coobrigação", "Sem carência", "Residual sem benchmark",
    "Estático; sem revolvência", "Derivativos vedados", "0", "0% a.a.",
    "15% numeral; 20% por extenso; divergência documental",
    "Até 96 parcelas, exceto FGTS; por devedor ≤ R$ 220 mil",
    "IME agregado não valida elegibilidade jurídica", "Mensal conforme suplemento",
    "PL da classe, conforme faixas do regulamento",
    "Facta, conforme contratos de endosso",
    "CDB-V e dação em pagamento de créditos de leasing conforme Contrato CDB",
    "Sênior, mezanino e júnior conforme suplemento existente.",
    "PF/PJ conforme tipo de crédito: duplicatas, cheques, CCB, CPR-F e outros títulos",
    "Pode haver ou não coobrigação, conforme contrato.",
    "Cessões podem ter ou não coobrigação conforme contrato",
    "Cedentes podem assumir coobrigação conforme contrato de cessão",
    "Total, parcial ou ausente, conforme cada contrato de cessão",
    "Recompra/compra por endossante, credor original ou originador conforme eventos do contrato",
    "Cedente e conveniados comerciais conforme contratos",
    "Pro rata ou sequencial conforme fase/suplemento",
    "Crédito adimplente na aquisição; devedor sem atraso >12 dias úteis; prazos médios de 540 DU (CCB/CPR/NCC) e 70 DU (demais).",
    "Cedentes bancos/instituições financeiras e instituições de ensino, conforme contrato; Pravaler presta gestão/originação operacional.",
    "Gestão substituída por Augme Capital; administração Genial Investimentos, conforme ato.",
    "1ª emissão sênior e mezanino; subordinadas júnior conforme regulamento",
    "Classe fechada, indeterminada, qualificados; suplementos sênior destinados a profissionais",
    "Proteção patrimonial ou troca de indexador sem risco de capital; suplemento prevê opções de compra DI",
    "Mezanino 1ª série: Taxa DI; sobretaxa não aplicável no suplemento.",
])
def test_material_negatives_conditions_conflicts_and_zero_are_preserved(text):
    assert document_comparison_value(text) == text


def test_public_sources_are_portable_and_keep_primary_identification():
    text = "Regulamento ID CVM 123, p. 7; /Users/matheusjprates/fidc/data/raw/123/regulamento.pdf"
    assert public_document_text(text) == "Regulamento ID CVM 123, p. 7; data/raw/123/regulamento.pdf"
    assert public_document_text("Imagem IMG_1713.jpg fornecida pelo usuário") == ""
    assert public_document_text("Auto IV da imagem: proposta sem CNPJ. Condições não verificadas.") == ""
    assert public_document_text("Fotografia da indústria; CVM / Fundos.NET") == "Fotografia da indústria; CVM / Fundos.NET"


def test_ppt_table_and_notes_are_clean_without_losing_primary_sources():
    manifest = DeepDiveManifest.from_dict({"deep_dive_id": "clean", "title": "FIDC BV", "generated_at": "2026-10-07T12:00:00-03:00", "funds": []}, package_dir=Path("unused"))
    page = DocumentComparisonPage("Custos", pd.DataFrame({"Critério": ["Taxa", "Prazo"], "Fundo": ["0,05% a.a.; mínimo não identificado", "Prazo remetido ao suplemento"]}),
        notes=("Auto IV da imagem: proposta sem CNPJ",),
        sources=("ID CVM 123 p.7 /Users/matheusjprates/fidc/data/raw/123/regulamento.pdf", "Imagem IMG_1713.jpg fornecida pelo usuário"))
    deck = Presentation(BytesIO(build_document_comparison_pptx_bytes(manifest, [page])))
    table = next(shape.table for shape in deck.slides[0].shapes if shape.has_table)
    assert table.cell(1, 1).text == "0,05% a.a."
    assert table.cell(2, 1).text == "-"
    assert tuple(table.cell(0, 1).fill.fore_color.rgb) == (255, 98, 0)
    assert tuple(table.cell(0, 1).text_frame.paragraphs[0].runs[0].font.color.rgb) == (255, 255, 255)
    notes = deck.slides[0].notes_slide.notes_text_frame.text
    assert "ID CVM 123 p.7" in notes
    assert "/Users/" not in notes and "IMG_1713" not in notes and "da imagem" not in notes


def test_export_boundary_removes_split_personal_paths_and_keeps_other_parts():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    paragraph = slide.shapes.add_textbox(0, 0, 1000000, 1000000).text_frame.paragraphs[0]
    paragraph.add_run().text = "Fonte: /Users/"
    paragraph.add_run().text = "matheusjprates/fidc/data/raw/regulamento.pdf"
    slide.notes_slide.notes_text_frame.text = "Regulamento ID CVM 123, p. 7"
    output = BytesIO(); deck.save(output)
    original = output.getvalue(); clean = public_pptx_bytes(original)
    with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(clean)) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "ppt/slides/slide1.xml":
                assert before.read(name) == after.read(name), name
    reopened = Presentation(BytesIO(clean))
    assert reopened.slides[0].shapes[0].text == "Fonte: data/raw/regulamento.pdf"
    assert reopened.slides[0].notes_slide.notes_text_frame.text == "Regulamento ID CVM 123, p. 7"
    assert public_pptx_bytes(clean) == clean
