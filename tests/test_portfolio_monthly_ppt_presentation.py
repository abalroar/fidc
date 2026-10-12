"""Portfolio-only presentation rules retain the shared export defaults and data."""
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pandas as pd
import pytest
from pptx import Presentation

from services.somatorio_fidcs_ppt_export import build_somatorio_fidcs_pptx_bytes


def _inputs(*, fund_count=1, credit=3.0, fund_ends_before_portfolio=False):
    dates = pd.date_range('2026-01-01', periods=3, freq='MS')
    consolidated = pd.DataFrame({
        'competencia_dt': dates, 'competencia': [date.strftime('%m/%Y') for date in dates],
        'fund_name': ['Carteira de teste'] * 3, 'pl_senior': [700.0] * 3,
        'pl_subordinada_mezz_ex360': [300.0] * 3, 'subordinacao_total_ex360_pct': [30.0] * 3,
        'carteira_ex360': [1000.0] * 3, 'carteira_ex360_yoy_pct': [10.0] * 3,
        'npl_over1_ex360_pct': [credit + 3] * 3, 'npl_over30_ex360_pct': [credit + 2] * 3,
        'npl_over60_ex360_pct': [credit + 1] * 3, 'npl_over90_ex360_pct': [credit] * 3,
        'npl_over90_ex360': [credit * 10] * 3,
        'pdd_npl_over90_ex360_pct': [200.0 if credit else float('nan')] * 3,
        'npl_1_90_pct': [credit] * 3, 'npl_91_360_pct': [credit] * 3,
        'duration_months': [6.0] * 3,
        'roll_61_90_m3_pct': [1.0] * 3, 'roll_91_120_m4_pct': [2.0] * 3,
        'roll_121_150_m5_pct': [3.0] * 3, 'roll_151_180_m6_pct': [4.0] * 3,
    })
    funds = {}
    for index in range(fund_count):
        fund = consolidated.copy()
        fund['fund_name'] = f'CLOUDWALK KICK ASS {index + 1} FUNDO DE INVESTIMENTO EM DIREITOS CREDITÓRIOS SEGMENTO MEIOS DE PAGAMENTO'
        fund['duration_months'] = 4.0 + index
        if fund_ends_before_portfolio:
            fund = fund.head(1).copy()
        funds[f'{index:014d}'] = fund
    outputs = SimpleNamespace(consolidated_monthly=consolidated, fund_monthly=funds,
                              fund_return_history={}, fund_return_summary={},
                              metadata={'portfolio_name': 'Nome do usuário 1T26', 'funds': []})
    monitor = SimpleNamespace(consolidated_monitor=consolidated, fund_monitor=funds,
                              consolidated_cohorts=pd.DataFrame(), fund_cohorts={})
    return outputs, monitor


def _deck(outputs, monitor, **kwargs):
    payload = build_somatorio_fidcs_pptx_bytes(outputs=outputs, monitor_outputs=monitor, **kwargs)
    with ZipFile(BytesIO(payload)) as package:
        assert package.testzip() is None
        assert any(name.startswith('ppt/charts/chart') for name in package.namelist())
        assert any(name.startswith('ppt/embeddings/') for name in package.namelist())
    return Presentation(BytesIO(payload))


def _texts(deck):
    return '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)


def _chart_data(deck):
    return [[(series.name, list(series.values)) for series in shape.chart.series]
            for slide in deck.slides for shape in slide.shapes if shape.has_chart]


def test_portfolio_style_is_opt_in_and_preserves_original_financial_chart_values():
    outputs, monitor = _inputs()
    before = _deck(outputs, monitor)
    after = _deck(outputs, monitor, portfolio_mode=True)
    assert len(before.slides) == len(after.slides) == 11
    assert 'Mercado de FIDCs' in _texts(before)
    assert 'Soma de FIDCs — Consolidado' in _texts(before)
    assert 'Nome do usuário 1T26' in _texts(after)
    assert 'Carteira — Consolidado' in _texts(after)
    assert _chart_data(before) == _chart_data(after)
    assert '00000000000000' in after.slides[6].notes_slide.notes_text_frame.text
    assert 'FUNDO DE INVESTIMENTO EM DIREITOS CREDITÓRIOS' in after.slides[6].notes_slide.notes_text_frame.text


def test_each_fund_uses_its_own_last_observed_competence():
    outputs, monitor = _inputs(fund_ends_before_portfolio=True)
    after = _deck(outputs, monitor, portfolio_mode=True)
    individual_text = '\n'.join(shape.text for slide in list(after.slides)[7:10]
                                for shape in slide.shapes if shape.has_text_frame)
    assert 'Data-base: jan/26' in individual_text
    assert 'Data-base: mar/26' not in individual_text


def test_zero_credit_denominator_is_explicit_and_missing_data_is_distinct():
    outputs, monitor = _inputs(credit=0)
    after = _deck(outputs, monitor, portfolio_mode=True)
    assert 'Não aplicável: NPL > 90d ex-360 é zero na janela.' in _texts(after)
    assert '0,0%' in _texts(after)
    assert any(value == 0.0 for slide in after.slides for shape in slide.shapes if shape.has_chart
               for series in shape.chart.series for value in series.values)
    outputs.consolidated_monthly['npl_over90_ex360'] = float('nan')
    outputs.fund_monthly['00000000000000']['npl_over90_ex360'] = float('nan')
    missing = _deck(outputs, monitor, portfolio_mode=True)
    assert 'N/D: série sem valor informado na janela.' in _texts(missing)


def test_portfolio_duration_keeps_every_fund_in_the_editable_chart():
    outputs, monitor = _inputs(fund_count=5)
    after = _deck(outputs, monitor, portfolio_mode=True)
    duration = next(shape.chart for shape in after.slides[4].shapes if shape.has_chart)
    assert len(duration.series) == 6
    latest = sorted(series.values[-1] for series in duration.series)
    assert latest == [4.0, 5.0, 6.0, 6.0, 7.0, 8.0]


def test_short_duration_label_does_not_change_the_original_series_lookup():
    outputs, monitor = _inputs()
    outputs.fund_monthly['00000000000000']['fund_name'] = 'MERCADO CRÉDITO I BRASIL FIDC SEGMENTO FINANCEIRO DE RESPONSABILIDADE LIMITADA'
    after = _deck(outputs, monitor, portfolio_mode=True)
    duration = next(shape.chart for shape in after.slides[4].shapes if shape.has_chart)
    assert sorted(series.values[-1] for series in duration.series) == [4.0, 6.0]


def test_portfolio_font_floor_and_credit_band_values_remain_visible():
    outputs, monitor = _inputs()
    after = _deck(outputs, monitor, portfolio_mode=True)
    sizes = [run.font.size.pt for slide in after.slides for shape in slide.shapes
             for frame in ([shape.text_frame] if shape.has_text_frame else
                           [cell.text_frame for row in shape.table.rows for cell in row.cells] if shape.has_table else [])
             for paragraph in frame.paragraphs for run in paragraph.runs if run.font.size]
    assert min(sizes) >= 11
    native_rows = [cell.text for shape in after.slides[2].shapes if shape.has_table
                   for row in shape.table.rows for cell in row.cells]
    assert 'NPL > 1d: 6,0%' in native_rows
    assert 'NPL > 30d: 5,0%' in native_rows
    assert 'NPL > 60d: 4,0%' in native_rows


def test_return_headers_are_short_and_every_native_value_is_preserved():
    outputs, monitor = _inputs()
    history = outputs.consolidated_monthly[['competencia', 'competencia_dt']].copy()
    history['class_kind'] = 'senior'
    history['class_key'] = 'senior:1'
    history['class_label'] = 'Sênior — série 1'
    history['retorno_mensal_pct'] = [1.0, 2.0, 0.0]
    outputs.fund_return_history = {'00000000000000': history}
    extra = {'monthly_cdi_rates_by_fund': {}, 'benchmark_spreads_by_fund': {}}
    before = _deck(outputs, monitor, **extra)
    after = _deck(outputs, monitor, portfolio_mode=True, **extra)
    before_tables = [[[cell.text for cell in row.cells] for row in shape.table.rows]
                     for slide in before.slides for shape in slide.shapes if shape.has_table]
    after_tables = [[[cell.text for cell in row.cells] for row in shape.table.rows]
                    for slide in after.slides for shape in slide.shapes if shape.has_table]
    assert before_tables
    for table in before_tables:
        matching = next(candidate for candidate in after_tables if candidate[1:] == table[1:])
        assert matching[0] != table[0]
        assert matching[0][0] == 'Série'
    notes = '\n'.join(slide.notes_slide.notes_text_frame.text for slide in after.slides)
    assert 'Ac. Últ. 12m (%)' in notes


def _values_only(deck):
    return [[list(series.values) for series in shape.chart.series]
            for slide in deck.slides for shape in slide.shapes if shape.has_chart]


def test_compaction_never_changes_the_shared_default_export():
    outputs, monitor = _inputs()
    before = _deck(outputs, monitor)
    guarded = _deck(outputs, monitor, compact=True)
    assert _texts(before) == _texts(guarded)
    assert _chart_data(before) == _chart_data(guarded)
    assert len(before.slides) == len(guarded.slides)


def test_compact_observed_credit_keeps_every_original_native_graph_value():
    outputs, monitor = _inputs()
    before = _deck(outputs, monitor, portfolio_mode=True)
    after = _deck(outputs, monitor, portfolio_mode=True, compact=True)
    assert len(after.slides) < len(before.slides)
    assert _values_only(before) == _values_only(after)
    assert 'NPL Over 90d ex-360 / Carteira' in '\n'.join(
        shape.chart.chart_title.text_frame.text for slide in after.slides for shape in slide.shapes if shape.has_chart)


def test_valid_zeros_are_explicit_with_complete_original_observations_in_notes():
    outputs, monitor = _inputs(credit=0)
    after = _deck(outputs, monitor, portfolio_mode=True, compact=True)
    texts = _texts(after)
    assert 'NPL > 90d ex-360: 0 em 3 observações' in texts
    assert 'não aplicável (NPL zero)' in texts
    notes = '\n'.join(slide.notes_slide.notes_text_frame.text for slide in after.slides)
    assert 'Compactação:' in notes
    assert '"npl_over90_ex360_pct": 0' in notes
    assert '"competencia": "01/2026"' in notes
    assert '"competencia": "03/2026"' in notes


def test_zero_series_with_gaps_uses_native_table_without_filling_missing_values():
    outputs, monitor = _inputs()
    for frame in [outputs.consolidated_monthly, *outputs.fund_monthly.values()]:
        for column in ('roll_61_90_m3_pct', 'roll_91_120_m4_pct', 'roll_121_150_m5_pct', 'roll_151_180_m6_pct'):
            frame[column] = [float('nan'), 0.0, 0.0]
    after = _deck(outputs, monitor, portfolio_mode=True, compact=True)
    tables = [[[cell.text for cell in row.cells] for row in shape.table.rows]
              for slide in after.slides for shape in slide.shapes if shape.has_table]
    gap = next(table for table in tables if table[0][0] == '0 / N/D na base')
    assert gap[0][1:] == ['jan/26', 'fev/26', 'mar/26']
    assert len(gap) == 5
    assert all(row[1:] == ['N/D', '0', '0'] for row in gap[1:])


def test_any_nonzero_pdd_in_the_window_keeps_a_graph_even_if_latest_is_zero():
    outputs, monitor = _inputs(credit=0)
    for frame in [outputs.consolidated_monthly, *outputs.fund_monthly.values()]:
        frame['pdd_total'] = [0.01, 0.0, 0.0]
    after = _deck(outputs, monitor, portfolio_mode=True, compact=True)
    pdd = [shape.chart for slide in after.slides for shape in slide.shapes
           if shape.has_chart and shape.chart.chart_title.text_frame.text == 'PDD total']
    assert len(pdd) == 2
    assert all(list(chart.series[0].values) == [0.01, 0.0, 0.0] for chart in pdd)


def _add_return_history(outputs, *, series_count):
    rows = []
    for series in range(series_count):
        for month in range(1, 4):
            rows.append({'competencia': f'{month:02d}/2026', 'competencia_dt': pd.Timestamp(2026, month, 1),
                         'class_kind': 'senior', 'class_key': f'senior:{series}', 'class_label': f'Sênior — série {series + 1}',
                         'retorno_mensal_pct': (month - 2) / 10.0})
    outputs.fund_return_history = {'00000000000000': pd.DataFrame(rows)}


def test_two_complementary_returns_fit_one_slide_with_all_native_cells():
    outputs, monitor = _inputs()
    _add_return_history(outputs, series_count=3)
    options = {'portfolio_mode': True, 'monthly_cdi_rates_by_fund': {}, 'benchmark_spreads_by_fund': {}}
    before = _deck(outputs, monitor, **options)
    after = _deck(outputs, monitor, compact=True, **options)
    before_tables = [[[cell.text for cell in row.cells] for row in shape.table.rows]
                     for slide in before.slides for shape in slide.shapes if shape.has_table and len(shape.table.rows) > 1]
    after_tables = [[[cell.text for cell in row.cells] for row in shape.table.rows]
                    for slide in after.slides for shape in slide.shapes if shape.has_table and len(shape.table.rows) > 1]
    for table in before_tables:
        assert table in after_tables
    return_slides = [slide for slide in after.slides if any(shape.has_text_frame and 'Rentabilidade por série' in shape.text for shape in slide.shapes)]
    assert len(return_slides) == 1
    assert sum(shape.has_table for shape in return_slides[0].shapes) == 2


def test_many_return_series_stay_paginated_without_losing_any_series():
    outputs, monitor = _inputs()
    _add_return_history(outputs, series_count=23)
    after = _deck(outputs, monitor, portfolio_mode=True, compact=True,
                  monthly_cdi_rates_by_fund={}, benchmark_spreads_by_fund={})
    returned = [shape.table for slide in after.slides for shape in slide.shapes
                if shape.has_table and shape.table.cell(0, 0).text == 'Série']
    counts = {f'Sênior — série {index + 1}': 0 for index in range(23)}
    for table in returned:
        for row in list(table.rows)[1:]:
            counts[row.cells[0].text] += 1
    assert set(counts.values()) == {2}
    assert all(len(table.rows) <= 23 for table in returned)
    slides = [slide for slide in after.slides if any(shape.has_text_frame and 'Rentabilidade por série' in shape.text for shape in slide.shapes)]
    assert len(slides) == 3


def test_empty_monthly_fund_remains_identified_with_explicit_nd_and_returns():
    outputs, monitor = _inputs()
    _add_return_history(outputs, series_count=2)
    outputs.fund_monthly['00000000000000'] = outputs.fund_monthly['00000000000000'].head(0)
    monitor.fund_monitor = outputs.fund_monthly
    after = _deck(outputs, monitor, portfolio_mode=True, compact=True,
                  monthly_cdi_rates_by_fund={}, benchmark_spreads_by_fund={})
    assert 'Indicadores mensais: N/D, sem informe utilizável na janela selecionada.' in _texts(after)
    notes = '\n'.join(slide.notes_slide.notes_text_frame.text for slide in after.slides)
    assert 'CNPJ: 00000000000000' in notes
    assert 'retornos disponíveis foram mantidos integralmente' in notes.lower()
    assert any(shape.has_table and shape.table.cell(0, 0).text == 'Série'
               for slide in after.slides for shape in slide.shapes)


def test_component_grid_rejects_overflow_instead_of_silently_dropping_a_graph():
    from services.somatorio_fidcs_ppt_export import _add_component_slide, _component

    components = [_component(str(index), pd.DataFrame({'v': [1.0]}), ('v',), lambda *_: None)
                  for index in range(5)]
    with pytest.raises(ValueError, match='nenhum gráfico pode ser descartado'):
        _add_component_slide(None, None, 'Teste', 1, 1, 'mar/26', components)


def test_compact_synthesis_explains_exact_zero_delta_and_keeps_original_in_notes():
    outputs, monitor = _inputs(credit=0)
    after = _deck(outputs, monitor, portfolio_mode=True, compact=True)
    synthesis = after.slides[-1]
    texts = '\n'.join(shape.text for shape in synthesis.shapes if shape.has_text_frame)
    assert 'estável versus mês anterior (variação zero)' in texts
    assert 'avançou 0,0 p.p.' not in texts
    assert 'Sem alerta automático adicional' not in texts
    assert 'avançou 0,0 p.p.' in synthesis.notes_slide.notes_text_frame.text


def test_dense_duration_prunes_only_annotations_and_prioritizes_consolidated():
    outputs, monitor = _inputs(fund_count=8)
    monitor.consolidated_monitor['duration_months'] = 3.0
    for index, frame in enumerate(monitor.fund_monitor.values()):
        frame['duration_months'] = 3.0 + index / 1000.0
    after = _deck(outputs, monitor, portfolio_mode=True, compact=True)
    duration = next(shape.chart for slide in after.slides for shape in slide.shapes
                    if shape.has_chart and shape.chart.chart_title.text_frame.text == 'Duration por FIDC')
    assert len(duration.series) == 9
    namespace = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    counts = {series.name: len(series._element.findall('c:dLbls/c:dLbl', namespace)) for series in duration.series}
    assert counts['Consolidado'] == 1
    assert sum(counts.values()) < len(duration.series)
    assert sorted(series.values[-1] for series in duration.series) == [3.0, 3.0, 3.001, 3.002, 3.003, 3.004, 3.005, 3.006, 3.007]
