"""Portfolio metadata must not rewrite the financial artifact it describes."""
from datetime import datetime, timezone
from io import BytesIO
import json
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table
import pytest

from services.portfolio_export_metadata import with_csv_export_context, with_xlsx_export_context


def _original_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Financeiro"
    sheet.append(["CNPJ", "Saldo", "Taxa", "Lacuna"])
    sheet.append(["00123456000100", 0, 0.0125, float("nan")])
    sheet.append(["00123456000100", 1234567.89, -0.01, None])
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(name="Calibri", bold=True)
    sheet["A2"].number_format = "@"
    sheet["B2"].number_format = '#,##0.00'
    sheet["C2"].number_format = '0.0%'
    sheet["E2"] = '=SUM(B2:B3)'
    sheet.add_table(Table(displayName="tbDados", ref="A1:D3"))
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    sheet.add_chart(chart, "G5")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_xlsx_appends_only_context_parts_and_preserves_every_native_financial_part():
    original = _original_workbook()
    context = {
        "Carteira": "Nome do usuário 1T26",
        "CNPJ": "00123456000100",
        "Fundos": 1,
        "Zero válido": 0,
        "Ausente": float("nan"),
        "Fonte": "CVM/Fundos.NET",
        "Consulta": datetime(2026, 10, 12, 19, 0, tzinfo=timezone.utc),
    }
    result = with_xlsx_export_context(original, context)
    changed_declarations = {"xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml"}
    with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(result)) as after:
        assert after.testzip() is None
        assert set(after.namelist()) - set(before.namelist()) == {"xl/worksheets/sheet2.xml"}
        for part in set(before.namelist()) - changed_declarations:
            assert after.read(part) == before.read(part), part
    before = load_workbook(BytesIO(original))
    after = load_workbook(BytesIO(result))
    assert after.sheetnames == ["Financeiro", "Contexto da exportação"]
    for actual_row, original_row in zip(after["Financeiro"], before["Financeiro"]):
        assert [(cell.value, cell.data_type, cell.number_format) for cell in actual_row] == [(cell.value, cell.data_type, cell.number_format) for cell in original_row]
    assert after["Financeiro"]["A2"].value == "00123456000100"
    assert after["Financeiro"]["B2"].value == 0
    assert after["Financeiro"]["C2"].value == 0.0125
    assert after["Financeiro"]["D2"].value is None
    assert after["Financeiro"]["E2"].data_type == "f"
    assert len(after["Financeiro"]._charts) == 1
    assert "tbDados" in after["Financeiro"].tables
    sheet = after["Contexto da exportação"]
    values = {row[0]: row[1] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert values["CNPJ"] == "00123456000100"
    assert values["Zero válido"] == 0
    assert values["Ausente"] is None
    assert values["Consulta"] == "2026-10-12T19:00:00+00:00"
    assert sheet.freeze_panes == "A2"


def test_csv_adds_context_json_without_changing_csv_bytes_or_missing_data():
    output = BytesIO()
    files = {"consolidado.csv": b'cnpj;zero;sem_dado\n00123456000100;0;NaN\n', "fundos.csv": 'CNPJ;Carteira\n00123456000100;São Paulo\n'.encode("utf-8-sig")}
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as source:
        for name, payload in files.items():
            source.writestr(name, payload)
    result = with_csv_export_context(output.getvalue(), {"CNPJ": "00123456000100", "Fundos": 1, "Ausente": None, "Janela": ["08/2025", "08/2026"]})
    with ZipFile(BytesIO(result)) as package:
        assert package.testzip() is None
        assert set(package.namelist()) == {*files, "contexto_exportacao.json"}
        for name, payload in files.items():
            assert package.read(name) == payload
        metadata = json.loads(package.read("contexto_exportacao.json"))
        assert metadata == {"CNPJ": "00123456000100", "Fundos": 1, "Ausente": None, "Janela": ["08/2025", "08/2026"]}


def test_context_never_overwrites_an_existing_sheet_or_zip_member():
    original = _original_workbook()
    result = with_xlsx_export_context(original, {"Fonte": "CVM"})
    with pytest.raises(ValueError, match="não foi sobrescrita"):
        with_xlsx_export_context(result, {"Fonte": "Outra"})
    output = BytesIO()
    with ZipFile(output, "w") as package:
        package.writestr("contexto_exportacao.json", b'{"existing": true}')
    with pytest.raises(ValueError, match="não foi sobrescrito"):
        with_csv_export_context(output.getvalue(), {})


def test_oversized_metadata_is_reported_without_truncating_the_original_file():
    original = _original_workbook()
    with pytest.raises(ValueError, match="nenhum conteúdo foi truncado"):
        with_xlsx_export_context(original, {"Campo": "x" * 32768})
    assert load_workbook(BytesIO(original))["Financeiro"]["B3"].value == 1234567.89
