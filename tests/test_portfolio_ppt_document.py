from base64 import b64decode
from io import BytesIO
from pathlib import PurePosixPath
import posixpath
from types import SimpleNamespace
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt
import pytest

from services.portfolio_ppt_presentation import finalize_portfolio_pptx
from services.portfolio_presentation_context import PortfolioPresentationContext
from services.pptx_merge import merge_pptx_bytes


NAME = "Cloudwalk (Fundos Adquirência 1T26)"
CONTEXT = PortfolioPresentationContext(
    name=NAME, fund_count=2, requested_window="10/2025 a 09/2026",
    historical_window="08/2025 a 08/2026", common_position="02/2026",
    latest_observed="08/2026", consulted_at="12/10/2026 13:34 BRT",
)
FUNDS = (
    SimpleNamespace(cnpj="42085816000105", display_name="CLOUDWALK KICK ASS I FUNDO DE INVESTIMENTO EM DIREITOS CREDITÓRIOS RESPONSABILIDADE LIMITADA"),
    SimpleNamespace(cnpj="42085830000109", display_name="CLOUDWALK AKIRA I FUNDO DE INVESTIMENTO EM DIREITOS CREDITÓRIOS RESPONSABILIDADE LIMITADA"),
)


def _text(slide, text, *, y=.2, size=22):
    shape = slide.shapes.add_textbox(Inches(.5), Inches(y), Inches(11), Inches(.6))
    shape.text = text
    shape.text_frame.paragraphs[0].font.size = Pt(size)
    return shape


def _fixture(*, full=True, b_cover_title=NAME):
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blocks = []
    if full:
        for name in ("A", "B"):
            cover = prs.slides.add_slide(prs.slide_layouts[6])
            _text(cover, NAME if name == "A" else b_cover_title, y=2.1, size=32)
            _text(cover, "Data-base original e referências da fonte", y=3.3, size=11)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            _text(slide, f"{name} · Carteira de crédito")
            _text(slide, "Data-base do fundo: 08/2026", y=.8, size=11)
            _text(slide, "Fonte: CVM · Página 1 de 99", y=7.05, size=9)
            data = CategoryChartData(); data.categories = ["jan/26", "fev/26", "mar/26"]
            data.add_series("PL", [0.0, 12.3, -2.0])
            slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.5), Inches(1.2), Inches(6), Inches(4), data)
            blocks.append((name, 2))
    for name in ("C", "D"):
        slide = prs.slides.add_slide(prs.slide_layouts[6]); _text(slide, f"{name} · Documentos")
        table = slide.shapes.add_table(3, 2, Inches(.5), Inches(1.2), Inches(10), Inches(2)).table
        for row, values in enumerate((("Critério", "Valor"), ("PL", "R$ 0,0 mm"), ("Cobertura", "N/D"))):
            for col, value in enumerate(values): table.cell(row, col).text = value
        slide.notes_slide.notes_text_frame.text = f"Leitura documental: 07/10/2026\nCompactação: {name} · repetição documental agrupada.\nFonte: data/raw/42085816000105/211003.pdf · p. 3"
        _text(slide, "Leitura documental: 07/10/2026 · 1/35", y=7.03, size=9)
        blocks.append((name, 1))
    # Keep a native image part to prove it survives package finalization.
    png = b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jC1sAAAAASUVORK5CYII=")
    prs.slides[-1].shapes.add_picture(BytesIO(png), Inches(11), Inches(1.2), width=Inches(.1))
    output = BytesIO(); prs.save(output)
    return output.getvalue(), blocks


def _visible(slide):
    return "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)


def _original_table_values(prs):
    return [[[cell.text for cell in row.cells] for row in shape.table.rows] for slide in prs.slides for shape in slide.shapes if shape.has_table and shape.table.cell(0, 0).text == "Critério"]


def _assert_no_dangling_references(payload):
    with ZipFile(BytesIO(payload)) as archive:
        paths = set(archive.namelist())
        for name in paths:
            if not name.endswith(".rels"):
                continue
            parent = PurePosixPath(name).parent
            base = str(parent.parent) if parent.name == "_rels" else str(parent)
            for relationship in etree.fromstring(archive.read(name)):
                if relationship.get("TargetMode") == "External": continue
                target = relationship.get("Target", "")
                path = posixpath.normpath(posixpath.join(base, target)) if not target.startswith("/") else target.lstrip("/")
                assert path in paths, (name, target)


def test_final_document_has_one_cover_final_numbering_and_preserved_native_data():
    before, blocks = _fixture()
    payload = finalize_portfolio_pptx(before, context=CONTEXT, blocks=blocks, funds=FUNDS)
    prs = Presentation(BytesIO(payload))
    assert len(prs.slides) == 6  # Remove one repeated cover, append one reference.
    assert sum(any(shape.has_text_frame and shape.text == NAME for shape in slide.shapes) for slide in prs.slides) == 1
    assert NAME in _visible(prs.slides[0])
    assert "Posição comum · 02/2026" in _visible(prs.slides[0])
    assert "Histórico selecionado · 08/2025 a 08/2026" in _visible(prs.slides[0])
    assert "Consulta · 12/10/2026 13:34 BRT" in _visible(prs.slides[0])
    assert "Referências da análise" in _visible(prs.slides[-1])
    assert _original_table_values(prs) == _original_table_values(Presentation(BytesIO(before)))
    assert sum(shape.has_chart for slide in prs.slides for shape in slide.shapes) == 2
    for index, slide in enumerate(prs.slides, 1):
        assert f"{index}/{len(prs.slides)}" in _visible(slide)
        assert "Página 1 de 99" not in _visible(slide)
    notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in prs.slides)
    assert "Leitura documental: 07/10/2026" in notes
    assert "C · repetição documental agrupada" in notes
    assert "D · repetição documental agrupada" in notes
    assert "capa repetida do bloco B retirada" in notes
    assert "211003.pdf · p. 3" in notes
    assert "data/raw/" not in notes
    assert all(fund.display_name in prs.slides[0].notes_slide.notes_text_frame.text for fund in FUNDS)
    assert all(fund.cnpj in prs.slides[0].notes_slide.notes_text_frame.text for fund in FUNDS)


def test_native_parts_master_layout_theme_and_images_are_unchanged_byte_for_byte():
    before, blocks = _fixture()
    payload = finalize_portfolio_pptx(before, context=CONTEXT, blocks=blocks, funds=FUNDS)
    protected = ("ppt/charts/", "ppt/embeddings/", "ppt/media/", "ppt/slideMasters/", "ppt/slideLayouts/", "ppt/theme/")
    with ZipFile(BytesIO(before)) as source, ZipFile(BytesIO(payload)) as finalized:
        for path in source.namelist():
            if path.startswith(protected):
                assert source.read(path) == finalized.read(path)
    _assert_no_dangling_references(payload)
    roundtrip = BytesIO(); Presentation(BytesIO(payload)).save(roundtrip)
    _assert_no_dangling_references(roundtrip.getvalue())
    assert _original_table_values(Presentation(BytesIO(roundtrip.getvalue()))) == _original_table_values(Presentation(BytesIO(before)))


def test_partial_document_gets_editable_cover_and_warning_without_inventing_a_date():
    before, blocks = _fixture(full=False)
    warnings = ("Blocos mensais indisponíveis: fontes inconsistentes.", "Curadoria parcial: referência não localizada.")
    payload = finalize_portfolio_pptx(before, context=CONTEXT, blocks=blocks, funds=FUNDS, warnings=warnings)
    prs = Presentation(BytesIO(payload))
    assert len(prs.slides) == 4
    assert NAME in _visible(prs.slides[0])
    assert "Exportação parcial" in _visible(prs.slides[0])
    assert "Posição comum · 02/2026" in _visible(prs.slides[0])
    assert all(warning in prs.slides[-1].notes_slide.notes_text_frame.text for warning in warnings)
    assert _original_table_values(prs) == _original_table_values(Presentation(BytesIO(before)))
    _assert_no_dangling_references(payload)


def test_b_cover_is_preserved_without_a_and_unverified_cover_is_not_deleted():
    before, blocks = _fixture()
    prs = Presentation(BytesIO(before))
    for _ in range(2):
        node = prs.slides._sldIdLst[0]; prs.part.drop_rel(node.rId); prs.slides._sldIdLst.remove(node)
    output = BytesIO(); prs.save(output)
    payload = finalize_portfolio_pptx(output.getvalue(), context=CONTEXT, blocks=blocks[1:], funds=FUNDS)
    finalized = Presentation(BytesIO(payload))
    assert len(finalized.slides) == 5
    assert NAME in _visible(finalized.slides[0])
    before, blocks = _fixture(b_cover_title="Título de outra seleção")
    payload = finalize_portfolio_pptx(before, context=CONTEXT, blocks=blocks, funds=FUNDS)
    assert "Título de outra seleção" in "\n".join(_visible(slide) for slide in Presentation(BytesIO(payload)).slides)


def test_inconsistent_block_count_fails_before_modifying_the_document():
    payload, blocks = _fixture()
    with pytest.raises(ValueError, match="contagem dos blocos"):
        finalize_portfolio_pptx(payload, context=CONTEXT, blocks=[("A", 1)], funds=FUNDS)


def test_finalization_preserves_single_compatible_notes_master_and_all_documentary_notes():
    first, _ = _fixture(full=False)
    second, _ = _fixture(full=False)
    merged = merge_pptx_bytes(first, second)
    payload = finalize_portfolio_pptx(merged, context=CONTEXT, blocks=[("C", 2), ("D", 2)], funds=FUNDS)
    presentation = Presentation(BytesIO(payload))
    assert len(presentation.slides) == 6
    assert all(slide.has_notes_slide for slide in presentation.slides)
    notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in presentation.slides)
    assert notes.count("C · repetição documental agrupada") >= 2
    assert notes.count("D · repetição documental agrupada") >= 2
    with ZipFile(BytesIO(merged)) as before, ZipFile(BytesIO(payload)) as after:
        assert len([path for path in after.namelist() if path.startswith("ppt/notesMasters/") and path.endswith(".xml")]) == 1
        masters = etree.fromstring(after.read("ppt/presentation.xml")).xpath("//p:notesMasterId", namespaces={"p": "http://schemas.openxmlformats.org/presentationml/2006/main"})
        assert len(masters) == 1
        for path in before.namelist():
            if path.startswith("ppt/notesMasters/"):
                assert before.read(path) == after.read(path)
    _assert_no_dangling_references(payload)


def test_finalizer_rejects_legacy_multiple_master_package_instead_of_hiding_it_in_a_cache():
    source, blocks = _fixture(full=False)
    output = BytesIO()
    with ZipFile(BytesIO(source)) as before, ZipFile(output, "w") as target:
        for path in before.namelist():
            data = before.read(path)
            if path == "[Content_Types].xml":
                types = etree.fromstring(data)
                etree.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", PartName="/ppt/notesMasters/notesMaster2.xml", ContentType="application/vnd.openxmlformats-officedocument.presentationml.notesMaster+xml")
                data = etree.tostring(types)
            target.writestr(path, data)
        target.writestr("ppt/notesMasters/notesMaster2.xml", before.read("ppt/notesMasters/notesMaster1.xml"))
    with pytest.raises(ValueError, match="vários notes masters"):
        finalize_portfolio_pptx(output.getvalue(), context=CONTEXT, blocks=blocks, funds=FUNDS)


def test_unique_notes_of_removed_duplicate_cover_are_preserved_in_references():
    before, blocks = _fixture()
    source = Presentation(BytesIO(before))
    unique_note = "Fonte documental exclusiva da capa B: documento 123456, p. 9. Ressalva integral."
    source.slides[2].notes_slide.notes_text_frame.text = unique_note
    output = BytesIO(); source.save(output)
    payload = finalize_portfolio_pptx(output.getvalue(), context=CONTEXT, blocks=blocks, funds=FUNDS)
    finalized = Presentation(BytesIO(payload))
    assert len(finalized.slides) == 6
    assert unique_note in finalized.slides[-1].notes_slide.notes_text_frame.text


def test_individual_snapshot_competence_stays_visible_in_the_unified_footer():
    before, blocks = _fixture(full=False)
    prs = Presentation(BytesIO(before))
    prs.slides[0].notes_slide.notes_text_frame.text += "\nPosição CVM: 08/2026"
    output = BytesIO(); prs.save(output)
    payload = finalize_portfolio_pptx(output.getvalue(), context=CONTEXT, blocks=blocks, funds=FUNDS)
    finalized = Presentation(BytesIO(payload))
    assert "CVM · 08/2026 / Fundos.NET" in _visible(finalized.slides[1])
    assert "Posição comum · 02/2026" in _visible(finalized.slides[0])


def test_cover_notes_prioritize_full_legal_names_and_keep_distinct_short_aliases():
    before, blocks = _fixture()
    funds = (
        SimpleNamespace(name=FUNDS[0].display_name, display_name="CW Kick Ass I", cnpj=FUNDS[0].cnpj),
        {"name": FUNDS[1].display_name, "display_name": "CW Akira I", "cnpj": FUNDS[1].cnpj},
    )
    payload = finalize_portfolio_pptx(before, context=CONTEXT, blocks=blocks, funds=funds)
    cover = Presentation(BytesIO(payload)).slides[0]
    notes = cover.notes_slide.notes_text_frame.text
    for fund in FUNDS:
        assert f"{fund.display_name} · CNPJ {fund.cnpj}" in notes
    assert "Nome curto: CW Kick Ass I" in notes
    assert "Nome curto: CW Akira I" in notes
    assert NAME in _visible(cover)
