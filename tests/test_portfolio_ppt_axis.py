import pandas as pd
import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

from services.somatorio_fidcs_ppt_export import _chart_pl_stacked


@pytest.mark.parametrize("senior,expected", [(2_731_000_000, 0), (0, 0), (-1_000_000, None)])
def test_stacked_pl_zero_origin_preserves_negative_values(senior, expected):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    frame = pd.DataFrame({"pl_senior": [senior], "pl_subordinada_mezz_ex360": [0]})
    _chart_pl_stacked(
        slide, (.5, 2, 6, 3), frame, ["ago/26"],
        CategoryChartData=CategoryChartData, RGBColor=RGBColor,
        XL_CHART_TYPE=XL_CHART_TYPE, XL_LEGEND_POSITION=XL_LEGEND_POSITION,
        XL_DATA_LABEL_POSITION=XL_DATA_LABEL_POSITION, Inches=Inches, Pt=Pt,
    )
    chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
    assert chart.value_axis.minimum_scale == expected
    assert chart.series[0].values[0] == senior / (1_000_000 if abs(senior) >= 1_000_000 else 1)
