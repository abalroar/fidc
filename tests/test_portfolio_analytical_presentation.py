from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO
import re
from types import SimpleNamespace
from unittest.mock import patch

from pptx import Presentation

from services.fidc_analytical_slide import (
    FidcSlideData,
    Issuance,
    build_fidc_analytical_pptx_bytes,
)


GENERATED_AT = datetime(2026, 10, 12, 13, 34)


def _presentation(items, *, legal_names=None, generated_at=GENERATED_AT, compact=False):
    funds = [SimpleNamespace(cnpj=item.cnpj, display_name=(legal_names or {}).get(item.cnpj, item.name)) for item in items]
    with patch("services.fidc_analytical_slide.build_slide_data", return_value=items):
        payload = build_fidc_analytical_pptx_bytes(funds, portfolio_name="Carteira 1T26", generated_at=generated_at, compact=compact)
    return Presentation(BytesIO(payload))


def _text(presentation):
    pieces = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                pieces.append(shape.text)
            if shape.has_table:
                pieces.extend(cell.text for row in shape.table.rows for cell in row.cells)
        if slide.has_notes_slide:
            pieces.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(pieces)


def test_no_rdb_evidence_cannot_publish_an_unrelated_conclusion_or_static_date():
    item = FidcSlideData(name="Cloudwalk Kick Ass I", cnpj="42085816000105", pl_mm=0)
    presentation = _presentation([item])
    text = _text(presentation)
    assert len(presentation.slides) == 2
    assert "RDB não localizado nos documentos consultados" in text
    assert "não comprova inexistência" in text
    assert "Não localizado" in text
    assert "R$ 0,0 mm" in text
    assert "12/10/2026" in text
    assert "14/08/2026" not in text
    assert "MT FGTS" not in text
    assert "CCB Tudo Serviços" not in text
    assert "Parati" not in text
    assert "não aplicável" not in text.lower()


def test_evidence_is_retained_and_public_notes_keep_primary_citations():
    source = "Fundos.NET ID 211003 | 01/09/2021 | Regulamento | p. 3 | data/raw/42085816000105/documento.pdf | https://fnet.bmfbovespa.com.br/fnet/publico/visualizarDocumento?id=211003"
    item = FidcSlideData(
        name="Fundo", cnpj="42085816000105", rdb_status="Explícito",
        rdb_finding="RDB Parati localizado no regulamento; uso condicionado ao contrato.",
        rdb_source="/Users/analista/projeto/data/raw/42085816000105/documento.pdf · p. 3",
        issuances=[Issuance("09/2021", "Subordinadas", "Subordinada", 5.6, "Residual", "N/D", source)],
    )
    presentation = _presentation([item])
    text = _text(presentation)
    assert "RDB Parati localizado" in text
    assert "Explícito: 1" in text
    assert "Fundos.NET ID 211003" in text
    assert "p. 3" in text
    assert "https://fnet.bmfbovespa.com.br/fnet/publico/visualizarDocumento?id=211003" in text
    assert "data/raw/" not in text
    assert "/Users/" not in text
    assert all(slide.has_notes_slide for slide in presentation.slides)


def test_cloudwalk_size_keeps_slide_count_and_separates_conclusion_from_table():
    items = [FidcSlideData(name=f"FIDC {index}", cnpj=f"{index:014d}") for index in range(11)]
    presentation = _presentation(items)
    assert len(presentation.slides) == 12
    rdb_slide = presentation.slides[0]
    table = next(shape for shape in rdb_slide.shapes if shape.has_table)
    conclusion = next(shape for shape in rdb_slide.shapes if shape.has_text_frame and "RDB não localizado nos documentos consultados" in shape.text)
    assert len(table.table.rows) == 12
    assert table.top + table.height < conclusion.top


def test_path_only_source_keeps_document_identifier_without_internal_directory():
    item = FidcSlideData(
        name="FIDC", cnpj="42085816000105",
        rdb_source="data/raw/42085816000105/211003_review_20261007.pdf · p. 3",
    )
    text = _text(_presentation([item]))
    assert "211003_review_20261007.pdf · p. 3" in text
    assert "data/raw/" not in text


def test_legal_name_stays_in_notes_while_canonical_alias_is_used_in_tables():
    legal_name = "CLOUDWALK A.I. FUNDO DE INVESTIMENTO EM DIREITOS CREDITÓRIOS RESPONSABILIDADE LIMITADA"
    item = FidcSlideData(name=legal_name[:56], cnpj="57609282000146", pl_mm=100)
    original = deepcopy(item)
    presentation = _presentation([item], legal_names={item.cnpj: legal_name})
    table = next(shape.table for shape in presentation.slides[0].shapes if shape.has_table)
    assert table.cell(1, 0).text == "CW A.I."
    assert "CW A.I. | estrutura e emissões" in _text(presentation)
    assert legal_name in presentation.slides[1].notes_slide.notes_text_frame.text
    assert item == original


def test_long_rdb_findings_continue_without_silent_word_truncation():
    finding = " ".join(f"cláusula{i}" for i in range(600))
    item = FidcSlideData(name="FIDC com texto extenso", cnpj="12345678000199", rdb_status="Parcial", rdb_finding=finding)
    presentation = _presentation([item])
    rdb_slides = [slide for slide in presentation.slides if any(shape.has_text_frame and "RDB — leitura documental" in shape.text for shape in slide.shapes)]
    assert len(rdb_slides) > 1
    cells = " ".join(shape.table.cell(row, 2).text for slide in rdb_slides for shape in slide.shapes if shape.has_table for row in range(1, len(shape.table.rows)))
    assert re.sub(r"\s+", " ", cells).strip() == finding
    for slide in rdb_slides:
        table = next(shape for shape in slide.shapes if shape.has_table)
        conclusion = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text.startswith("Resultado por fundo"))
        assert table.top + table.height < conclusion.top
    assert finding in presentation.slides[-1].notes_slide.notes_text_frame.text


def test_long_issuance_rows_remain_native_and_all_records_are_present():
    issuances = [Issuance(
        "10/2025", f"Série completa número {index} de emissão subordinada", "Subordinada", float(index),
        "Remuneração vinculada ao CDI conforme índice contratual, com pagamentos e condições descritos integralmente no documento da emissão.",
        "Amortização contratual", f"Fundos.NET ID {10000 + index} · p. 5",
    ) for index in range(12)]
    item = FidcSlideData(name="FIDC", cnpj="12345678000199", issuances=issuances)
    presentation = _presentation([item])
    emission_tables = [shape.table for slide in presentation.slides for shape in slide.shapes if shape.has_table and shape.table.cell(0, 0).text == "Data"]
    rows = [tuple(cell.text for cell in row.cells) for table in emission_tables for row in list(table.rows)[1:]]
    assert len(emission_tables) > 1
    assert len(rows) == len(issuances)
    assert {re.sub(r"\s+", " ", row[1]) for row in rows} == {item.series for item in issuances}
    assert {re.sub(r"\s+", " ", row[3]) for row in rows} == {item.remuneration for item in issuances}
    for table in emission_tables:
        for row in list(table.rows)[1:]:
            for cell in row.cells:
                assert all(run.font.size.pt == 9 for paragraph in cell.text_frame.paragraphs for run in paragraph.runs)
    for slide in presentation.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= presentation.slide_width
            assert shape.top + shape.height <= presentation.slide_height
            if shape.has_table:
                assert shape.top + shape.height < 6.95 * 914400


def test_generation_date_respects_sao_paulo_and_never_uses_portfolio_name():
    item = FidcSlideData(name="FIDC", cnpj="12345678000199")
    presentation = _presentation([item], generated_at=datetime(2026, 10, 13, 1, tzinfo=timezone.utc))
    text = _text(presentation)
    assert "Geração: 12/10/2026" in text
    assert "Geração: 13/10/2026" not in text
    assert "Carteira: Carteira 1T26" in text


def _financial_tables(presentation):
    return [[
        [re.sub(r"\s+", " ", cell.text) for cell in row.cells]
        for row in shape.table.rows
    ] for slide in presentation.slides for shape in slide.shapes if shape.has_table and shape.table.cell(0, 0).text in {"Snapshot", "Data"}]


def test_compact_all_missing_rdb_keeps_every_fund_and_financial_native_value():
    items = [FidcSlideData(
        name=f"Fundo {index}", cnpj=f"{index:014d}", competence="02/2026",
        pl_mm=0 if index == 0 else (None if index == 1 else 100.125),
        receivables_pct_pl=0, senior_pct_pl=.8, subordinated_pct_pl=.2,
        identified_issuance_volume_mm=0 if index == 0 else None,
        minimum_subordination="≥ 4,8%: subordinadas / PL, apuração diária",
        issuances=[Issuance("09/2021", "Série subordinada", "Subordinada", 0, "Residual", "N/D", "Fundos.NET ID 211003 · p. 3")],
        validations=["RDB não localizado nos documentos extraídos; ausência não prova inexistência."],
    ) for index in range(11)]
    before = deepcopy(items)
    default = _presentation(items)
    compact = _presentation(items, compact=True)
    assert len(default.slides) == 12
    assert len(compact.slides) == 11
    assert _financial_tables(default) == _financial_tables(compact)
    assert items == before
    first = compact.slides[0]
    visible = "\n".join(shape.text for shape in first.shapes if shape.has_text_frame)
    assert "RDB não localizado em 11/11 fundos nos documentos consultados" in visible
    assert "Ausência não prova inexistência" in visible
    assert "detalhamento integral nas notas" in visible
    assert all(item.cnpj in first.notes_slide.notes_text_frame.text for item in items)
    assert "R$ 0,0 mm" in _text(compact)
    assert "N/D" in _text(compact)
    assert "Cronograma quantificável: N/D" in visible


def test_compact_any_rdb_evidence_retains_integral_dedicated_comparison():
    items = [
        FidcSlideData(name="Fundo A", cnpj="11111111000111"),
        FidcSlideData(name="Fundo B", cnpj="22222222000122", rdb_status="Parcial", rdb_finding="Uso condicionado ao contrato de garantia."),
    ]
    compact = _presentation(items, compact=True)
    assert len(compact.slides) == 3
    rdb_table = next(shape.table for shape in compact.slides[0].shapes if shape.has_table)
    assert rdb_table.cell(1, 1).text == "Não localizado"
    assert rdb_table.cell(2, 1).text == "Parcial"
    assert rdb_table.cell(2, 2).text == items[1].rdb_finding
    assert "RDB não localizado em 2/2" not in _text(compact)


def test_compact_font_floor_and_long_validation_visibility():
    validations = [f"Controle documental {index}: " + " ".join(f"ressalva{word}" for word in range(70)) for index in range(8)]
    item = FidcSlideData(name="Fundo A", cnpj="11111111000111", validations=validations)
    compact = _presentation([item], compact=True)
    assert len(compact.slides) > 1
    visible = "\n".join(shape.text for slide in compact.slides for shape in slide.shapes if shape.has_text_frame)
    for index in range(8):
        assert f"Controle documental {index}" in visible
    for slide in compact.slides:
        for shape in slide.shapes:
            assert shape.top + shape.height <= compact.slide_height
            assert shape.left + shape.width <= compact.slide_width
            frames = [shape.text_frame] if shape.has_text_frame else []
            if shape.has_table:
                frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
            for frame in frames:
                for paragraph in frame.paragraphs:
                    for run in paragraph.runs:
                        assert run.font.name == "Calibri"
                        assert run.font.size.pt >= (9 if shape.top / 914400 >= 7 else 10)


def test_compact_long_issuance_table_paginates_at_readable_font_without_lost_records():
    issuances = [Issuance(
        "10/2025", f"Série subordinada identificada integralmente {index}", "Subordinada", float(index),
        "Remuneração vinculada ao CDI com pagamento contratual sem alteração dos dados da emissão.",
        "Texto documental integral de amortização", f"Fundos.NET ID {10000 + index}",
    ) for index in range(15)]
    item = FidcSlideData(name="Fundo A", cnpj="11111111000111", issuances=issuances)
    compact = _presentation([item], compact=True)
    tables = [shape.table for slide in compact.slides for shape in slide.shapes if shape.has_table and shape.table.cell(0, 0).text == "Data"]
    assert len(tables) > 1
    records = [list(row.cells) for table in tables for row in list(table.rows)[1:]]
    assert len(records) == 15
    assert {re.sub(r"\s+", " ", row[1].text) for row in records} == {item.series for item in issuances}
    assert {re.sub(r"\s+", " ", row[3].text) for row in records} == {item.remuneration for item in issuances}
    assert all(shape.top + shape.height < 6.64 * 914400 for slide in compact.slides for shape in slide.shapes if shape.has_table)


def test_compact_extensive_subordination_rule_remains_native_and_explicit():
    rule = "Subordinação mínima de 15%, com apuração diária. " + " ".join(f"condição{index}" for index in range(700))
    item = FidcSlideData(name="Fundo A", cnpj="11111111000111", minimum_subordination=rule)
    compact = _presentation([item], compact=True)
    assert len(compact.slides) > 1
    visible = "\n".join(shape.text for shape in compact.slides[0].shapes if shape.has_text_frame)
    assert "Regra extensa — leitura integral na continuação e nas notas." in visible
    tables = [shape.table for slide in compact.slides for shape in slide.shapes if shape.has_table and shape.table.cell(0, 0).text == "Regra documental integral"]
    assert len(tables) > 1
    rendered = " ".join(table.cell(row, 0).text for table in tables for row in range(1, len(table.rows)))
    assert re.sub(r"\s+", " ", rendered).strip() == rule
    assert rule in compact.slides[0].notes_slide.notes_text_frame.text
