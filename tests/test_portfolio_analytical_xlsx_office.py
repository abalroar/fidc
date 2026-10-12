"""Excel filter ownership must preserve native financial workbook contents."""
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

from lxml import etree
from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest

from services import fidc_analytical_slide as analytical
from services.portfolio_export_metadata import with_xlsx_export_context


_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


@pytest.fixture
def selected_funds(monkeypatch, tmp_path):
    funds = [
        SimpleNamespace(cnpj="00123456000100", display_name="FUNDO UM RAZÃO SOCIAL INTEGRAL"),
        SimpleNamespace(cnpj="22222222000191", display_name="FUNDO DOIS RAZÃO SOCIAL INTEGRAL"),
    ]
    data = [
        analytical.FidcSlideData(
            name="Fundo um", cnpj=funds[0].cnpj, competence="02/2026",
            pl_mm=123.456789, receivables_pct_pl=1.0759,
            senior_pct_pl=0.88, subordinated_pct_pl=0.12,
            identified_issuance_volume_mm=47.125,
            issuances=[analytical.Issuance(
                date="02/2026", series="Série A", quota_class="Sênior",
                volume_mm=47.125, remuneration="CDI + 2%",
                amortization="Mensal", source="Fonte documental",
            )],
            validations=["Controle documental existente"],
        ),
        analytical.FidcSlideData(
            name="Fundo dois", cnpj=funds[1].cnpj, competence="02/2026",
            pl_mm=0.0, receivables_pct_pl=0.0,
            senior_pct_pl=None, subordinated_pct_pl=None,
            identified_issuance_volume_mm=0.0,
        ),
    ]
    evidence_path = tmp_path / "rdb.csv"
    pd.DataFrame([
        [funds[0].cnpj, "Fundo um", "Parcial", "Cláusula documental"],
        [funds[1].cnpj, "Fundo dois", "Não localizado", "N/D"],
    ], columns=["CNPJ", "Fundo", "Status", "Achado regulamento"]).to_csv(evidence_path, index=False)
    monkeypatch.setattr(analytical, "build_slide_data", lambda selected: data)
    monkeypatch.setattr(analytical, "RDB_EVIDENCE", evidence_path)
    return funds


def _build(funds, portfolio_mode):
    return analytical.build_fidc_analytical_xlsx_bytes(
        funds, portfolio_name="Nome do usuário 1T26", portfolio_mode=portfolio_mode,
    )


@pytest.mark.parametrize("portfolio_mode", [False, True])
def test_native_tables_own_their_filters_and_plain_sheets_keep_worksheet_filters(selected_funds, portfolio_mode):
    payload = with_xlsx_export_context(
        _build(selected_funds, portfolio_mode), {"Fonte": "CVM/Fundos.NET", "Fundos": 2},
    )
    workbook = load_workbook(BytesIO(payload))
    for name in ("Snapshot", "Emissões", "RDB", "Validações"):
        sheet = workbook[name]
        assert len(sheet.tables) == 1
        table = next(iter(sheet.tables.values()))
        assert table.autoFilter.ref == table.ref
        assert sheet.auto_filter.ref is None
    assert workbook["Metodologia"].auto_filter.ref == workbook["Metodologia"].dimensions
    assert workbook["Contexto da exportação"].auto_filter.ref == "A1:B3"
    with ZipFile(BytesIO(payload)) as package:
        root = etree.fromstring(package.read("xl/workbook.xml"))
        filter_names = root.findall(f"{{{_MAIN}}}definedNames/{{{_MAIN}}}definedName")
        assert [name.get("localSheetId") for name in filter_names if name.get("name") == "_xlnm._FilterDatabase"] == ["4"]


@pytest.mark.parametrize("portfolio_mode", [False, True])
def test_filter_fix_preserves_original_cells_formats_styles_and_native_table_parts(selected_funds, portfolio_mode, monkeypatch):
    fixed_payload = _build(selected_funds, portfolio_mode)
    fixed_write = analytical._write_sheet

    def reproduce_legacy_overlap(workbook, title, rows, table_name):
        fixed_write(workbook, title, rows, table_name)
        sheet = workbook[title]
        sheet.auto_filter.ref = sheet.dimensions

    monkeypatch.setattr(analytical, "_write_sheet", reproduce_legacy_overlap)
    legacy_payload = _build(selected_funds, portfolio_mode)
    before, after = load_workbook(BytesIO(legacy_payload)), load_workbook(BytesIO(fixed_payload))
    assert after.sheetnames == before.sheetnames
    for name in before.sheetnames:
        assert [
            (cell.coordinate, cell.value, cell.data_type, cell.number_format)
            for row in after[name] for cell in row
        ] == [
            (cell.coordinate, cell.value, cell.data_type, cell.number_format)
            for row in before[name] for cell in row
        ]
    assert after["Snapshot"]["B2"].value == "00123456000100"
    assert after["Snapshot"]["B2"].data_type == "s"
    assert after["Snapshot"]["E3"].value == 0.0
    assert after["Snapshot"]["E3"].data_type == "n"
    assert after["Snapshot"]["G3"].value is None
    with ZipFile(BytesIO(legacy_payload)) as legacy, ZipFile(BytesIO(fixed_payload)) as fixed:
        assert set(legacy.namelist()) == set(fixed.namelist())
        allowed = {"xl/workbook.xml", "docProps/core.xml"}
        for part in legacy.namelist():
            if part.startswith("xl/worksheets/sheet") and part.endswith(".xml"):
                original = etree.fromstring(legacy.read(part))
                updated = etree.fromstring(fixed.read(part))
                assert etree.tostring(original.find(f"{{{_MAIN}}}sheetData")) == etree.tostring(updated.find(f"{{{_MAIN}}}sheetData"))
            elif part not in allowed:
                assert legacy.read(part) == fixed.read(part), part


def test_header_only_sheet_keeps_its_worksheet_filter_without_creating_a_table():
    workbook = Workbook()
    workbook.remove(workbook.active)
    analytical._write_sheet(workbook, "Sem registros", [["CNPJ", "Valor"]], "tbSemRegistros")
    output = BytesIO()
    workbook.save(output)
    sheet = load_workbook(BytesIO(output.getvalue()))["Sem registros"]
    assert not sheet.tables
    assert sheet.auto_filter.ref == "A1:B1"
    assert list(sheet.iter_rows(values_only=True)) == [("CNPJ", "Valor")]
