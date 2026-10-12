from dataclasses import replace
from io import BytesIO
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

from services.deep_dive_models import DeepDiveManifest
from services.deep_dive_ppt_export import build_document_comparison_pptx_bytes
from services.document_curation_comparison import DocumentComparisonPage
from services.portfolio_document_display import DOCUMENT_GROUPING_CAPTION, compact_document_frame, portfolio_document_warnings
from services.pptx_merge import merge_pptx_bytes


def _manifest():
    return DeepDiveManifest.from_dict({
        "title": "Curadoria Carteira Cloudwalk 1T26",
        "generated_at": "2026-10-07T11:08:04-03:00",
        "source": "CVM / Fundos.NET",
        "funds": [{
            "cnpj": "42085816000105", "short_name": "CW Kick Ass I",
            "name": "CLOUDWALK KICK ASS I FUNDO DE INVESTIMENTO EM DIREITOS CREDITÓRIOS RESPONSABILIDADE LIMITADA",
        }],
    }, package_dir=Path("unused"))


def _page():
    return DocumentComparisonPage(
        "Recebíveis e elegibilidade",
        pd.DataFrame({"Critério": ["Devedor", "Taxa", "Quantidade"], "CW Kick Ass I": ["Emissores dos instrumentos de pagamento", None, 0]}, dtype=object),
        notes=("A leitura se limita aos documentos consultados.",),
        sources=("Regulamento ID Fundos.NET 211003 · p. 3 · https://fnet.bmfbovespa.com.br/fnet/publico/visualizarDocumento?id=211003",),
    )


def _presentation(*, portfolio_mode=False, pages=None, compact=False):
    return Presentation(BytesIO(build_document_comparison_pptx_bytes(_manifest(), pages or [_page()], portfolio_mode=portfolio_mode, compact=compact)))


def _tables(presentation):
    return [shape.table for slide in presentation.slides for shape in slide.shapes if shape.has_table]


def _values(presentation):
    return [[[cell.text for cell in row.cells] for row in table.rows] for table in _tables(presentation)]


def test_portfolio_mode_is_opt_in_and_preserves_native_content_and_slide_count():
    default = _presentation()
    portfolio = _presentation(portfolio_mode=True)
    assert len(default.slides) == len(portfolio.slides) == 1
    assert _values(default) == _values(portfolio)
    assert _tables(portfolio)[0].cell(2, 1).text == "-"
    assert _tables(portfolio)[0].cell(3, 1).text == "0"
    assert _tables(default)[0].cell(1, 1).text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
    assert _tables(portfolio)[0].cell(1, 1).text_frame.paragraphs[0].alignment == PP_ALIGN.LEFT


def test_portfolio_title_and_footer_are_short_with_complete_identification_in_notes():
    slide = _presentation(portfolio_mode=True).slides[0]
    visible = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
    notes = slide.notes_slide.notes_text_frame.text
    assert "Recebíveis e elegibilidade" in visible
    assert _manifest().title not in visible
    assert "Documentos · Leitura: 07/10/2026 · Fontes nas notas" in visible
    assert _manifest().title in notes
    assert _manifest().funds[0]["name"] in notes
    assert "CNPJ 42085816000105" in notes
    assert _page().sources[0] in notes
    assert _page().notes[0] in notes


def test_default_document_title_footer_and_alignment_remain_available():
    slide = _presentation().slides[0]
    visible = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
    assert f"{_manifest().title} – {_page().title}" in visible
    assert "Leitura documental: 07/10/2026. Fontes e referências nas notas do slide." in visible


def test_portfolio_mode_keeps_pagination_for_long_rows_and_many_funds():
    frame = pd.DataFrame({"Critério": [f"Regra {index}" for index in range(30)]})
    for index in range(7):
        frame[f"Fundo {index}"] = [f"Valor contratual fundo {index}, regra {row}" for row in range(30)]
    page = DocumentComparisonPage("Regras extensas", frame)
    default = _presentation(pages=[page])
    portfolio = _presentation(portfolio_mode=True, pages=[page])
    assert len(default.slides) == len(portfolio.slides)
    assert _values(default) == _values(portfolio)


def test_compact_document_frame_groups_only_identical_full_columns_and_keeps_order():
    frame = pd.DataFrame({
        "Critério": ["Taxa", "Devedor"],
        "Fundo A": ["5%", "Emissor A"],
        "Fundo B": ["5%", "Emissor A"],
        "Fundo C": ["5%", "Emissor C"],
    }, index=[4, 9])
    before = frame.copy(deep=True)
    compact, mappings = compact_document_frame(frame)
    assert compact.index.tolist() == [4, 9]
    assert compact.iloc[:, 0].tolist() == ["Taxa", "Devedor"]
    assert len(compact.columns) == 3
    assert compact.iloc[:, 1].tolist() == ["5%", "Emissor A"]
    assert compact.iloc[:, 2].tolist() == ["5%", "Emissor C"]
    assert "Fundo A" in mappings[0] and "Fundo B" in mappings[0]
    assert "Fundo C" in mappings[1]
    pd.testing.assert_frame_equal(frame, before)


def test_compact_document_frame_zero_is_distinct_from_missing_and_close_values_are_distinct():
    frame = pd.DataFrame({
        "Critério": ["Quantidade", "Taxa"],
        "Fundo A": [0, "1,0%"],
        "Fundo B": [None, "1,0%"],
        "Fundo C": [0, "1,00%"],
    }, dtype=object)
    compact, _ = compact_document_frame(frame)
    assert len(compact.columns) == 4
    assert compact.iloc[0, 1] == "0"
    assert compact.iloc[0, 2] == "-"
    assert compact.iloc[1, 1] == "1,0%"
    assert compact.iloc[1, 3] == "1,00%"


def test_compact_document_frame_all_funds_is_explicit_and_has_full_identity_mapping():
    legal_name = _manifest().funds[0]["name"]
    funds = [*_manifest().funds, {"name": "FUNDO B RAZÃO SOCIAL", "short_name": "Fundo B", "cnpj": "22222222000122"}]
    frame = pd.DataFrame({"Critério": ["Devedor"], "CW Kick Ass I": ["Emissores"], "Fundo B": ["Emissores"]})
    compact, mappings = compact_document_frame(frame, funds=funds)
    assert compact.columns.tolist() == ["Critério", "Todos os fundos"]
    assert legal_name in mappings[0]
    assert "CNPJ 42085816000105" in mappings[0]
    assert "FUNDO B RAZÃO SOCIAL" in mappings[0]
    assert "CNPJ 22222222000122" in mappings[0]


def test_compact_document_frame_single_fund_and_no_fund_columns_stay_unmerged():
    single = pd.DataFrame({"Critério": ["Taxa"], "Fundo A": [0]})
    compact, mappings = compact_document_frame(single)
    assert compact.columns.tolist() == single.columns.tolist()
    assert compact.iloc[0, 1] == "0"
    assert mappings == ()
    criteria_only = pd.DataFrame({"Critério": ["Taxa"]})
    compact, mappings = compact_document_frame(criteria_only)
    pd.testing.assert_frame_equal(compact, criteria_only)
    assert mappings == ()


def test_compact_document_frame_large_group_labels_do_not_collide_and_retain_all_members():
    frame = pd.DataFrame({"Critério": ["Regra"]})
    for index in range(6):
        frame[f"Fundo {index}"] = ["A" if index < 3 else "B"]
    compact, mappings = compact_document_frame(frame)
    assert compact.columns.tolist() == ["Critério", "3 fundos", "3 fundos (2)"]
    assert len(mappings) == 2
    assert all(f"Fundo {index}" in "\n".join(mappings) for index in range(6))
    assert compact.iloc[0, 1:].tolist() == ["A", "B"]


def test_compact_ppt_is_opt_in_and_explains_grouping_in_caption_and_native_notes():
    frame = pd.DataFrame({"Critério": ["Devedor", "Taxa"], "CW Kick Ass I": ["Emissores", 0], "Fundo B": ["Emissores", 0]})
    page = DocumentComparisonPage("Regras iguais", frame, notes=("Qualificação documental preservada.",), sources=("Regulamento, p. 3",))
    default = _presentation(portfolio_mode=True, pages=[page])
    compact = _presentation(portfolio_mode=True, pages=[page], compact=True)
    assert len(_tables(default)[0].columns) == 3
    table = _tables(compact)[0]
    assert len(table.columns) == 2
    assert table.cell(0, 1).text == "Todos os fundos"
    assert table.cell(2, 1).text == "0"
    slide = compact.slides[0]
    visible = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
    notes = slide.notes_slide.notes_text_frame.text
    assert DOCUMENT_GROUPING_CAPTION in visible
    assert "Todos os fundos →" in notes
    assert "Fundo B" in notes
    assert _manifest().funds[0]["name"] in notes
    assert page.notes[0] in notes and page.sources[0] in notes
    caption = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text == DOCUMENT_GROUPING_CAPTION)
    table_shape = next(shape for shape in slide.shapes if shape.has_table)
    assert caption.top + caption.height < table_shape.top


def test_compact_ppt_retains_long_clauses_and_continues_without_overflow():
    clause = " ".join(f"cláusula{index}" for index in range(600))
    page = DocumentComparisonPage("Texto extenso", pd.DataFrame({"Critério": ["Amortização"], "Fundo A": [clause], "Fundo B": [clause]}))
    compact = _presentation(portfolio_mode=True, pages=[page], compact=True)
    assert len(compact.slides) > 1
    text = "\n".join(table.cell(row, 1).text for table in _tables(compact) for row in range(1, len(table.rows)))
    assert all(f"cláusula{index}" in text for index in range(600))
    assert "..." not in text and "…" not in text
    for slide in compact.slides:
        for shape in slide.shapes:
            assert shape.top + shape.height <= compact.slide_height
            assert shape.left + shape.width <= compact.slide_width


def test_portfolio_documentary_ppt_keeps_all_manifest_caveats_once_in_first_slide_notes_only():
    warnings = tuple(f"RESSALVA_{index:02}: lacuna documental {index}." for index in range(14)) + (
        "RESSALVA_14: arquivo /Users/analista/Documentos/consulta.pdf. "
        + "Texto integral preservado. " * 40 + "FINAL_DA_RESSALVA.",
    )
    manifest = replace(_manifest(), warnings=warnings)
    pages = [_page(), replace(_page(), title="Proteções")]
    payload = build_document_comparison_pptx_bytes(manifest, pages, portfolio_mode=True, compact=True)
    portfolio = Presentation(BytesIO(payload))
    default = Presentation(BytesIO(build_document_comparison_pptx_bytes(manifest, pages)))
    assert len(portfolio.slides) == len(default.slides) == 2
    assert _values(portfolio) == _values(default)
    notes = [slide.notes_slide.notes_text_frame.text for slide in portfolio.slides]
    for warning in portfolio_document_warnings(warnings):
        assert warning in notes[0]
        assert "\n".join(notes).count(warning) == 1
    assert "Ressalvas da leitura:" in notes[0]
    assert "Ressalvas da leitura:" not in notes[1]
    assert "/Users/analista" not in "\n".join(notes)
    assert "FINAL_DA_RESSALVA." in notes[0]
    visible = "\n".join(shape.text for slide in portfolio.slides for shape in slide.shapes if shape.has_text_frame)
    assert "RESSALVA_" not in visible
    assert all("RESSALVA_" not in slide.notes_slide.notes_text_frame.text for slide in default.slides)
    assert manifest.warnings == warnings
    prefix = build_document_comparison_pptx_bytes(_manifest(), [_page()])
    merged = Presentation(BytesIO(merge_pptx_bytes(prefix, payload)))
    merged_notes = [slide.notes_slide.notes_text_frame.text for slide in merged.slides]
    assert len(merged.slides) == 3
    assert "Ressalvas da leitura:" not in merged_notes[0]
    assert "Ressalvas da leitura:" not in merged_notes[2]
    for warning in portfolio_document_warnings(warnings):
        assert warning in merged_notes[1]
        assert "\n".join(merged_notes).count(warning) == 1


def test_portfolio_caveat_sanitization_retains_duplicates_and_signals_private_only_content():
    warnings = ["Lacuna documental.", "Lacuna documental.", "Imagem fornecida", "", None]
    output = portfolio_document_warnings(warnings)
    assert output[:2] == ("Lacuna documental.", "Lacuna documental.")
    assert len(output) == 3
    assert output[2] == "Ressalva com referência privada; conteúdo omitido na apresentação pública."
