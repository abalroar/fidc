from __future__ import annotations

from io import BytesIO
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
import pytest

from scripts.patch_industry_workbook_operational_cache import (
    SHEET_NS,
    patch_workbook_bytes,
    validate_workbook_caches,
)


def _fixture(*, grouped_columns: bool = False) -> tuple[bytes, list[tuple[float | None, float | None]]]:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Base competência-CNPJ"
    for address, label in {
        "G4": "Carteira DC", "H4": "Inadimplência reportada",
        "AC4": "Ajustado (fórmula)", "AD4": "Excesso (fórmula)",
    }.items():
        worksheet[address] = label
    worksheet.column_dimensions["G"].width = 13
    worksheet.column_dimensions["H"].width = 14
    if grouped_columns:
        dimension = worksheet.column_dimensions["AB"]
        dimension.min, dimension.max, dimension.width = 28, 32, 10
    else:
        worksheet.column_dimensions["AC"].width = 14
    records = [(0, 0), (0, 25), (-10, 25), (100, -5), (100, 50), (None, 25), (100, None), (None, None), (10.25, 13.75), (10_000_000_000, 100_000_000_000)]
    for row, (portfolio, delinquency) in enumerate(records, 5):
        worksheet[f"A{row}"] = "2026-08"
        worksheet[f"B{row}"] = f"00.000.000/0001-{row:02}"
        worksheet[f"G{row}"] = portfolio
        worksheet[f"H{row}"] = delinquency
        worksheet[f"AC{row}"] = f'=IF(AND(G{row}<>"",H{row}<>""),MIN(MAX(H{row},0),MAX(G{row},0)),"")'
        worksheet[f"AD{row}"] = f'=IF(AC{row}="","",MAX(H{row}-AC{row},0))'
        worksheet[f"G{row}"].number_format = "R$ #,##0.00"
        worksheet[f"H{row}"].number_format = "R$ #,##0.00"
        worksheet[f"AD{row}"].number_format = "R$ #,##0.00"
    workbook.create_sheet("Preserved detail")["A1"] = "Fonte documental preservada"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), records


def _parts(payload: bytes) -> dict[str, bytes]:
    with ZipFile(BytesIO(payload)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _mutate(payload: bytes, name: str, mutation) -> bytes:
    output = BytesIO()
    with ZipFile(BytesIO(payload)) as source, ZipFile(output, "w") as destination:
        for info in source.infolist():
            content = source.read(info.filename)
            destination.writestr(info, mutation(content) if info.filename == name else content)
    return output.getvalue()


def _cells(raw: bytes) -> dict[str, bytes]:
    root = ET.fromstring(raw)
    return {cell.attrib["r"]: ET.tostring(cell) for cell in root.iter(f"{{{SHEET_NS}}}c")}


@pytest.mark.parametrize("grouped_columns", [False, True])
def test_repairs_every_semantic_result_preserving_reported_zero_missing_and_negative(grouped_columns) -> None:
    original, records = _fixture(grouped_columns=grouped_columns)
    with pytest.raises(ValueError, match="cache disagrees"):
        validate_workbook_caches(original)
    patched = patch_workbook_bytes(original)
    validate_workbook_caches(patched)
    assert patch_workbook_bytes(patched) == patched

    values = load_workbook(BytesIO(patched), data_only=True)
    formulas = load_workbook(BytesIO(patched), data_only=False)
    for row, (portfolio, delinquency) in enumerate(records, 5):
        sheet = values["Base competência-CNPJ"]
        assert sheet[f"G{row}"].value == portfolio
        assert sheet[f"H{row}"].value == delinquency
        cap = None if portfolio is None or delinquency is None else min(max(delinquency, 0), max(portfolio, 0))
        excess = None if portfolio is None or delinquency is None else max(delinquency - max(portfolio, 0), 0)
        assert sheet[f"AC{row}"].value == cap
        assert sheet[f"AD{row}"].value == excess
        assert formulas["Base competência-CNPJ"][f"AC{row}"].value == f'=IF(AND(G{row}<>"",H{row}<>""),MIN(MAX(H{row},0),MAX(G{row},0)),"")'
        assert formulas["Base competência-CNPJ"][f"AD{row}"].value == f'=IF(AC{row}="","",MAX(H{row}-AC{row},0))'
    assert formulas.calculation.calcMode == "auto"
    assert formulas.calculation.fullCalcOnLoad and formulas.calculation.forceFullCalc
    values.close()
    formulas.close()

    before, after = _parts(original), _parts(patched)
    assert before.keys() == after.keys()
    assert {name for name in before if before[name] != after[name]} == {"xl/workbook.xml", "xl/worksheets/sheet1.xml"}
    source_cells, final_cells = _cells(before["xl/worksheets/sheet1.xml"]), _cells(after["xl/worksheets/sheet1.xml"])
    for address, cell in source_cells.items():
        if address.startswith(("AC", "AD")) and int(address.lstrip("ACD")) >= 5:
            continue
        assert final_cells[address] == cell, address
    width_nodes = ET.fromstring(after["xl/worksheets/sheet1.xml"]).findall(f"{{{SHEET_NS}}}cols/{{{SHEET_NS}}}col")
    width = lambda column: next(float(node.attrib["width"]) for node in width_nodes if int(node.attrib["min"]) <= column <= int(node.attrib["max"]))
    assert width(6) >= 32 and width(7) >= 32 and width(8) >= 32
    assert width(29) >= 26 and width(30) >= 26
    if grouped_columns:
        assert width(28) == width(31) == width(32) == 10


def test_fails_closed_on_a_changed_formula_without_touching_input() -> None:
    original, _ = _fixture()
    bad = _mutate(original, "xl/worksheets/sheet1.xml", lambda raw: raw.replace(b"MAX(H5,0)", b"MAX(H5,1)", 1))
    with pytest.raises(ValueError, match="Unexpected operational formula AC5"):
        patch_workbook_bytes(bad)
    assert _parts(original)["xl/worksheets/sheet1.xml"] != _parts(bad)["xl/worksheets/sheet1.xml"]


def test_validator_rejects_a_stale_positive_cache_and_source_errors() -> None:
    original, _ = _fixture()
    patched = patch_workbook_bytes(original)
    bad = _mutate(patched, "xl/worksheets/sheet1.xml", lambda raw: re.sub(rb'(<c r="AD6"[^>]*>.*?<v>)25(</v>)', rb'\g<1>999\g<2>', raw, count=1))
    with pytest.raises(ValueError, match="cache disagrees"):
        validate_workbook_caches(bad)
    nonfinite = _mutate(original, "xl/worksheets/sheet1.xml", lambda raw: raw.replace(b'<c r="G5" s="1" t="n"><v>0</v>', b'<c r="G5" s="1" t="n"><v>NaN</v>', 1))
    with pytest.raises(ValueError, match="Nonfinite numeric source"):
        patch_workbook_bytes(nonfinite)


def test_zero_cache_must_remain_exactly_zero() -> None:
    original, _ = _fixture()
    patched = patch_workbook_bytes(original)
    bad = _mutate(patched, "xl/worksheets/sheet1.xml", lambda raw: re.sub(rb'(<c r="AC5"[^>]*>.*?<v>)0(</v>)', rb'\g<1>0.0000001\g<2>', raw, count=1))
    with pytest.raises(ValueError, match="cache disagrees.*AC5"):
        validate_workbook_caches(bad)
    validate_workbook_caches(patch_workbook_bytes(bad))


@pytest.mark.parametrize("source_column", [6, 7, 8])
def test_source_monetary_width_only_patch_preserves_all_cells_and_caches(source_column: int) -> None:
    original, _ = _fixture()
    corrected = patch_workbook_bytes(original)
    def narrow_portfolio(raw: bytes) -> bytes:
        columns = ET.fromstring(raw).findall(f"{{{SHEET_NS}}}cols/{{{SHEET_NS}}}col")
        column_id = str(source_column)
        column = next(node for node in columns if int(node.attrib["min"]) <= source_column <= int(node.attrib["max"]))
        assert column.attrib["min"] == column.attrib["max"] == column_id
        width = column.attrib["width"].encode()
        for match in re.finditer(rb"<col\b[^>]*>", raw):
            node = ET.fromstring(match.group())
            if node.attrib.get("min") == node.attrib.get("max") == column_id:
                changed = match.group().replace(b'width="' + width + b'"', b'width="13"', 1)
                return raw[:match.start()] + changed + raw[match.end():]
        raise AssertionError("portfolio column not found")

    narrow = _mutate(corrected, "xl/worksheets/sheet1.xml", narrow_portfolio)
    assert narrow != corrected
    with pytest.raises(ValueError, match=f"column {source_column} is too narrow"):
        validate_workbook_caches(narrow)
    restored = patch_workbook_bytes(narrow)
    assert restored == corrected
    before, after = _parts(narrow), _parts(restored)
    assert {name for name in before if before[name] != after[name]} == {"xl/worksheets/sheet1.xml"}
    assert _cells(before["xl/worksheets/sheet1.xml"]) == _cells(after["xl/worksheets/sheet1.xml"])
    assert patch_workbook_bytes(restored) == restored
