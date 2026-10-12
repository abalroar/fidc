from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pandas as pd
import pytest
from openpyxl import load_workbook

from services import fidc_analytical_slide as analytical


RDB_COLUMNS = [
    "CNPJ", "Fundo", "Status", "Achado regulamento", "Função documentada",
    "Limite", "Documento", "Página", "Proveniência",
]


@pytest.fixture
def selected_data(monkeypatch, tmp_path):
    funds = [
        SimpleNamespace(cnpj="11111111000191", display_name="FUNDO UM RAZÃO SOCIAL INTEGRAL"),
        SimpleNamespace(cnpj="22222222000191", display_name="FUNDO DOIS RAZÃO SOCIAL INTEGRAL"),
    ]
    data = [
        analytical.FidcSlideData(
            name="Fundo um", cnpj=funds[0].cnpj, competence="02/2026",
            pl_mm=123.456789, receivables_pct_pl=1.0759,
            senior_pct_pl=0.88, subordinated_pct_pl=0.12,
            identified_issuance_volume_mm=47.125,
            first_issuance="01/2025", last_issuance="02/2026",
            minimum_subordination="12%",
            issuances=[analytical.Issuance(
                date="02/2026", series="Série A", quota_class="Sênior",
                volume_mm=47.125, remuneration="CDI + 2%", amortization="Mensal",
                source="Documento selecionado",
            )],
            validations=["Controle documental existente"],
        ),
        analytical.FidcSlideData(
            name="Fundo dois", cnpj=funds[1].cnpj, competence="02/2026",
            pl_mm=0.0, receivables_pct_pl=0.0,
            senior_pct_pl=None, subordinated_pct_pl=None,
            identified_issuance_volume_mm=0.0,
            issuances=[analytical.Issuance(
                date="01/2026", series="Série B", quota_class="Subordinada",
                volume_mm=0.0, remuneration="N/D", amortization="N/D", source="N/D",
            )],
        ),
    ]
    evidence = pd.DataFrame([
        ["11.111.111/0001-91", "Nome documental original", "Explícito", "Cláusula A", "Liquidez", "R$ 1 milhão", "a.pdf", "7", "Fundos.NET/B3"],
        ["99999999000191", funds[1].display_name, "Explícito", "Fora da cesta", "Outra função", "R$ 2 milhões", "alheio.pdf", "9", "Fundos.NET/B3"],
        [funds[0].cnpj, "Outro documento do mesmo CNPJ", "Parcial", "Cláusula B", "N/D", "N/D", "b.pdf", "11", "Fundos.NET/B3"],
    ], columns=RDB_COLUMNS)
    evidence_path = tmp_path / "rdb.csv"
    evidence.to_csv(evidence_path, index=False)
    original_source = evidence_path.read_bytes()
    received_funds = []

    def slide_data(selected):
        received_funds.append(tuple(selected))
        return data

    monkeypatch.setattr(analytical, "build_slide_data", slide_data)
    monkeypatch.setattr(analytical, "RDB_EVIDENCE", evidence_path)
    return funds, data, evidence, evidence_path, original_source, received_funds


def _xlsx(funds, **kwargs):
    return analytical.build_fidc_analytical_xlsx_bytes(funds, portfolio_name="Nome do usuário 1T26", **kwargs)


def _rows(sheet):
    return list(sheet.iter_rows(values_only=True))


def test_portfolio_rdb_uses_exact_cnpjs_preserves_evidence_and_reports_each_missing_fund(selected_data):
    funds, _, evidence, evidence_path, original_source, received_funds = selected_data
    workbook = load_workbook(BytesIO(_xlsx(iter(funds), portfolio_mode=True)))
    rows = _rows(workbook["RDB"])
    assert rows[0] == tuple(RDB_COLUMNS)
    assert rows[1:3] == [tuple(evidence.iloc[0]), tuple(evidence.iloc[2])]
    missing = dict(zip(RDB_COLUMNS, rows[3]))
    assert missing["CNPJ"] == funds[1].cnpj
    assert missing["Fundo"] == funds[1].display_name
    assert missing["Status"] == "Não localizado"
    assert missing["Achado regulamento"] == "Nenhum registro de RDB localizado no cadastro documental consultado."
    assert missing["Função documentada"] == "N/D; não inferido."
    assert all(missing[column] == "N/D" for column in ["Limite", "Documento", "Página", "Proveniência"])
    assert len(rows) == 4
    assert workbook["RDB"].tables["tbRDB"].ref == "A1:I4"
    assert received_funds == [tuple(funds)]
    assert evidence_path.read_bytes() == original_source
    methodology = dict(_rows(workbook["Metodologia"])[1:])
    assert "1 registro(s) de outros CNPJs excluído(s)" in methodology["RDB — perímetro"]
    assert "1 fundo(s) sem registro" in methodology["RDB — lacunas"]
    assert "não comprova inexistência de RDB" in methodology["RDB — lacunas"]


def test_default_rdb_keeps_global_registry_and_original_methodology(selected_data):
    funds, _, evidence, _, _, _ = selected_data
    default = load_workbook(BytesIO(_xlsx(funds)))
    explicit_default = load_workbook(BytesIO(_xlsx(funds, portfolio_mode=False)))
    expected_rows = [tuple(RDB_COLUMNS)] + [tuple(row) for row in evidence.values.tolist()]
    assert _rows(default["RDB"]) == expected_rows
    assert _rows(explicit_default["RDB"]) == expected_rows
    assert _rows(default["Metodologia"]) == _rows(explicit_default["Metodologia"])
    assert default["Metodologia"].max_row == 6
    assert "99999999000191" in [row[0] for row in _rows(default["RDB"])[1:]]


def test_portfolio_mode_keeps_financial_worksheets_and_native_numbers_identical(selected_data):
    funds, _, _, _, _, _ = selected_data
    default_payload = _xlsx(funds)
    portfolio_payload = _xlsx(funds, portfolio_mode=True)
    with ZipFile(BytesIO(default_payload)) as before, ZipFile(BytesIO(portfolio_payload)) as after:
        for number in [1, 2, 4]:
            path = f"xl/worksheets/sheet{number}.xml"
            assert before.read(path) == after.read(path)
    workbook = load_workbook(BytesIO(portfolio_payload), data_only=False)
    assert workbook.sheetnames == ["Snapshot", "Emissões", "RDB", "Validações", "Metodologia"]
    assert workbook["Snapshot"]["A2"].value == "Nome do usuário 1T26"
    assert workbook["Snapshot"]["E2"].value == 123.456789
    assert workbook["Snapshot"]["F2"].value == 1.0759
    assert workbook["Snapshot"]["E3"].value == 0.0
    assert workbook["Snapshot"]["E3"].data_type == "n"
    assert workbook["Snapshot"]["G3"].value is None
    assert workbook["Emissões"]["F3"].value == 0.0
    assert workbook["Emissões"]["F3"].data_type == "n"


def test_portfolio_without_rdb_records_keeps_one_explicit_row_per_selected_fund(selected_data):
    funds, _, evidence, evidence_path, _, _ = selected_data
    evidence.iloc[[1]].to_csv(evidence_path, index=False)
    workbook = load_workbook(BytesIO(_xlsx(funds, portfolio_mode=True)))
    rows = _rows(workbook["RDB"])[1:]
    assert [row[0] for row in rows] == [fund.cnpj for fund in funds]
    assert [row[1] for row in rows] == [fund.display_name for fund in funds]
    assert all(row[2] == "Não localizado" for row in rows)
    assert all("N/D" == row[7] for row in rows)
    assert len(rows) == len(funds)


def test_portfolio_rdb_refuses_a_registry_without_cnpj_instead_of_guessing_names(selected_data):
    funds, _, evidence, evidence_path, _, _ = selected_data
    evidence.drop(columns=["CNPJ"]).to_csv(evidence_path, index=False)
    with pytest.raises(ValueError, match="coluna CNPJ"):
        _xlsx(funds, portfolio_mode=True)
