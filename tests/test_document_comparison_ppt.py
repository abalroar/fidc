from __future__ import annotations

from contextlib import nullcontext
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

import pandas as pd
from pptx import Presentation
from pptx.util import Inches
import pytest

from services.deep_dive_models import DeepDiveManifest
from services.deep_dive_ppt_export import build_document_comparison_pptx_bytes
from services.document_curation_comparison import DocumentComparisonPage
from services.portfolio_store import PortfolioFund, PortfolioRecord, portfolio_basket_signature
from services.pptx_merge import merge_pptx_bytes
from tabs import portfolio_page


def _manifest() -> DeepDiveManifest:
    return DeepDiveManifest.from_dict(
        {
            "deep_dive_id": "documentos",
            "title": "FIDC Auto BV",
            "portfolio_id": "carteira-auto",
            "generated_at": "2026-10-07T11:08:04-03:00",
            "source": "CVM / Fundos.NET",
            "funds": [{"cnpj": "57532556000146", "name": "FIDC Auto I", "short_name": "FIDC Auto I"}],
        },
        package_dir=Path("unused"),
    )


def _page() -> DocumentComparisonPage:
    return DocumentComparisonPage(
        "Critérios e estrutura",
        pd.DataFrame({"Critério": ["LTV individual", "Índice de perdas 90"], "FIDC Auto I": ["≤ 80%", "≤ 5%"]}),
        notes=("Perdas contratuais usam o valor de face histórico das CCBs.",),
        sources=("Regulamento, ID CVM 101, 12/03/2025, cláusula 9.1, página 18. https://fnet.bmfbovespa.com.br/fnet/publico/exibirDocumento?id=101",),
    )


def _blank_deck(title: str) -> bytes:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1)).text = title
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _native_tables(presentation: Presentation) -> list:
    return [shape.table for slide in presentation.slides for shape in slide.shapes if shape.has_table]


def test_document_comparison_matches_reference_as_native_editable_table() -> None:
    payload = build_document_comparison_pptx_bytes(_manifest(), [_page()])
    presentation = Presentation(BytesIO(payload))
    assert len(presentation.slides) == 1
    table = _native_tables(presentation)[0]
    assert table.cell(0, 0).text == "Critério"
    assert tuple(table.cell(0, 1).fill.fore_color.rgb) == (236, 112, 0)
    assert tuple(table.cell(1, 1).fill.fore_color.rgb) == (255, 255, 255)
    assert tuple(table.cell(0, 1).text_frame.paragraphs[0].runs[0].font.color.rgb) == (31, 31, 31)
    assert table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.bold
    assert table.cell(1, 1).text == "≤ 80%"
    slide = presentation.slides[0]
    all_text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
    assert "07/10/2026" in all_text
    assert "valor de face histórico" in all_text
    assert _page().sources[0] in slide.notes_slide.notes_text_frame.text
    assert "CNPJ 57532556000146" in slide.notes_slide.notes_text_frame.text
    with ZipFile(BytesIO(payload)) as archive:
        assert "<a:tbl>" in archive.read("ppt/slides/slide1.xml").decode()
        assert not any(path.startswith("ppt/media/") for path in archive.namelist())


def test_document_comparison_paginates_all_funds_and_rows() -> None:
    frame = pd.DataFrame({"Critério": [f"Critério {row}" for row in range(45)]})
    for fund in range(7):
        frame[f"Fundo {fund}"] = [f"Valor {fund}-{row}" for row in range(45)]
    payload = build_document_comparison_pptx_bytes(_manifest(), [DocumentComparisonPage("Comparativo", frame)])
    tables = _native_tables(Presentation(BytesIO(payload)))
    assert len(tables) >= 6
    observed = set()
    for table in tables:
        assert len(table.columns) <= 5
        assert table.cell(0, 0).text == "Critério"
        for row in range(1, len(table.rows)):
            for col in range(1, len(table.columns)):
                observed.add(table.cell(row, col).text)
    assert observed == {f"Valor {fund}-{row}" for fund in range(7) for row in range(45)}


def test_document_comparison_retains_long_payment_schedule_without_overflow() -> None:
    schedule = "\n".join(f"10/{month:02d}/2027: pagamento contratual {month} conforme waterfall" for month in range(1, 13))
    schedule += "\n" + " ".join(f"cláusula{i}" for i in range(200))
    page = DocumentComparisonPage("Calendário", pd.DataFrame({"Critério": ["Amortização"], "Fundo A": [schedule]}))
    presentation = Presentation(BytesIO(build_document_comparison_pptx_bytes(_manifest(), [page])))
    texts = "\n".join(table.cell(row, 1).text for table in _native_tables(presentation) for row in range(1, len(table.rows)))
    assert "10/12/2027" in texts
    assert all(f"cláusula{i}" in texts for i in range(200))
    assert "..." not in texts and "…" not in texts
    assert len(presentation.slides) > 1
    for slide in presentation.slides:
        for shape in slide.shapes:
            assert shape.top + shape.height <= presentation.slide_height
            assert shape.left + shape.width <= presentation.slide_width


def test_document_comparison_missing_is_explicit_and_zero_remains_zero() -> None:
    page = DocumentComparisonPage("Lacunas", pd.DataFrame({"Critério": ["Taxa", "Quantidade"], "Fundo A": [None, 0]}, dtype=object))
    table = _native_tables(Presentation(BytesIO(build_document_comparison_pptx_bytes(_manifest(), [page]))))[0]
    assert table.cell(1, 1).text == "Não localizado"
    assert table.cell(2, 1).text == "0"
    with pytest.raises(ValueError, match="Não há tabelas"):
        build_document_comparison_pptx_bytes(_manifest(), [])


def test_merge_preserves_documentary_source_notes_and_office_roundtrip() -> None:
    first = build_document_comparison_pptx_bytes(_manifest(), [_page()])
    second_page = replace(_page(), title="Reservas", sources=("Regulamento ID CVM 202, página 30, cláusula 11",))
    second = build_document_comparison_pptx_bytes(_manifest(), [second_page])
    merged = merge_pptx_bytes(_blank_deck("IME"), first, second)
    presentation = Presentation(BytesIO(merged))
    assert len(presentation.slides) == 3
    assert _page().sources[0] in presentation.slides[1].notes_slide.notes_text_frame.text
    assert second_page.sources[0] in presentation.slides[2].notes_slide.notes_text_frame.text
    with ZipFile(BytesIO(merged)) as archive:
        assert "notesMasterIdLst" in archive.read("ppt/presentation.xml").decode()
        assert "notesSlide" in archive.read("ppt/slides/_rels/slide2.xml.rels").decode()
    saved = BytesIO()
    presentation.save(saved)
    reopened = Presentation(BytesIO(saved.getvalue()))
    assert _page().sources[0] in reopened.slides[1].notes_slide.notes_text_frame.text
    assert second_page.sources[0] in reopened.slides[2].notes_slide.notes_text_frame.text


def _portfolio() -> PortfolioRecord:
    return PortfolioRecord(
        id="carteira-auto", name="FIDC AUTO BV",
        funds=(PortfolioFund(cnpj="57532556000146", display_name="FIDC Auto I"),),
        created_at="2026-10-07", updated_at="2026-10-07",
    )


def test_documentary_export_selects_exact_current_basket_before_newer_fallback() -> None:
    portfolio = _portfolio()
    signature = portfolio_basket_signature(portfolio.funds)
    exact = replace(_manifest(), portfolio_signature=signature)
    fallback = replace(exact, portfolio_id="outra-carteira", generated_at="2026-10-08")
    stale = replace(exact, portfolio_signature="99999999000199", generated_at="2026-10-09")
    with (
        patch("services.deep_dive_store.list_deep_dives", return_value=[stale, fallback, exact]),
        patch("services.document_curation_comparison.build_document_comparison_pages", return_value=[_page()]) as load,
    ):
        payload = portfolio_page._build_documentary_portfolio_deck(portfolio)
    load.assert_called_once_with(exact)
    assert _native_tables(Presentation(BytesIO(payload)))[0].cell(1, 1).text == "≤ 80%"
    with patch("services.deep_dive_store.list_deep_dives", return_value=[stale]):
        assert portfolio_page._build_documentary_portfolio_deck(portfolio) is None


def test_complete_portfolio_download_contains_comparison_and_documentary_notes() -> None:
    portfolio = _portfolio()
    cnpj = portfolio.funds[0].cnpj
    outputs = SimpleNamespace(fund_monthly={cnpj: SimpleNamespace()})
    analysis = portfolio_page.PortfolioAnalysisData(
        scopes=(portfolio_page.PortfolioAnalysisScope(value=cnpj, label="FIDC A", kind="fund", cnpj=cnpj, dashboard=SimpleNamespace()),),
        aggregate_bundle=None, outputs=outputs, monitor_outputs=SimpleNamespace(),
        research_outputs=SimpleNamespace(), verification_report=None,
        dashboard_errors={}, load_errors={},
    )
    manifest = replace(_manifest(), portfolio_signature=portfolio_basket_signature(portfolio.funds))
    with (
        patch("tabs.portfolio_page.credit_tab.resolve_fund_return_export_inputs", return_value=({}, {})),
        patch("services.somatorio_fidcs_ppt_export.build_somatorio_fidcs_pptx_bytes", return_value=_blank_deck("IME")),
        patch("services.fidc_analytical_slide.build_fidc_analytical_pptx_bytes", return_value=_blank_deck("Estrutura")),
        patch("services.fidc_analytical_slide.build_fidc_analytical_xlsx_bytes", return_value=b"xlsx"),
        patch("services.deep_dive_store.list_deep_dives", return_value=[manifest]),
        patch("services.document_curation_comparison.build_document_comparison_pages", return_value=[_page()]),
        patch("tabs.portfolio_page.build_consolidated_snapshot_excel_bytes", return_value=b"xlsx"),
        patch("tabs.portfolio_page.build_full_variable_excel_export_bytes", return_value=b"xlsx"),
        patch("tabs.portfolio_page.build_full_variable_csv_zip_bytes", return_value=b"zip"),
        patch("tabs.portfolio_page.st.expander", return_value=nullcontext()),
        patch("tabs.portfolio_page.st.download_button") as download,
    ):
        portfolio_page._render_unified_portfolio_download(analysis=analysis, selected_portfolio=portfolio, period=SimpleNamespace(label="09/2026"))
    call = next(call for call in download.call_args_list if call.kwargs.get("key", "").startswith("portfolio_unified_pptx::"))
    presentation = Presentation(BytesIO(call.kwargs["data"]))
    assert len(presentation.slides) == 3
    assert _native_tables(presentation)[0].cell(2, 1).text == "≤ 5%"
    assert _page().sources[0] in presentation.slides[-1].notes_slide.notes_text_frame.text
