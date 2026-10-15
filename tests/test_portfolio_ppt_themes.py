"""Thematic layout rules of the portfolio deck, valid for every basket."""
from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

import pandas as pd
from pptx import Presentation
from pptx.util import Pt

from services.portfolio_ppt_style import DEFAULT_STYLE

from tests.test_portfolio_ppt_export import _build, _inputs, _spec, _tables, _visible_text


def _titles(payload):
    presentation = Presentation(BytesIO(payload))
    return [next((shape.text for shape in slide.shapes if shape.has_text_frame and shape.top < 914400), "")
            for slide in presentation.slides]


def _dashboard(months=13, *, subordination=None, coverage=None):
    dates = pd.date_range("2025-06-01", periods=months, freq="MS").strftime("%m/%Y")
    return SimpleNamespace(
        latest_competencia=dates[-1], summary={"ativos_totais": 1_200_000_000.0, "direitos_creditorios": 900_000_000.0},
        subordination_history_df=pd.DataFrame({"competencia": dates, "subordinacao_pct": subordination if subordination is not None else [30.0] * months}),
        default_history_df=pd.DataFrame({"competencia": dates, "inadimplencia_pct": [1.0 + i / 10 for i in range(months)],
                                         "provisao_pct_direitos": [0.5 + i / 20 for i in range(months)],
                                         "cobertura_pct": coverage if coverage is not None else [100.0] * months}),
    )


def _first(titles, name):
    return next(index for index, title in enumerate(titles) if name in title.split(" · ") or title.startswith(name))


def test_deck_reads_snapshot_themes_appendix_documents_then_coverage():
    kwargs = _inputs(months=13)
    kwargs["dashboard"] = _dashboard()
    titles = [title for title in _titles(_build(kwargs)) if title]
    order = ["Snapshot", "Estrutura", "Atraso e provisão", "Migração e safras", "Rentabilidade",
             "Apêndice", "Documental", "Referências e cobertura"]
    positions = [_first(titles, name) for name in order]
    assert positions == sorted(positions)
    # A theme is contiguous: its charts and reading table never split apart.
    for theme in order[1:5]:
        indexes = [index for index, title in enumerate(titles) if theme in title.split(" · ")]
        assert indexes == list(range(indexes[0], indexes[-1] + 1))
    assert "Séries mensais" not in titles and "Séries IME" not in titles


def test_series_tables_form_one_appendix_in_theme_order():
    spec = _spec(_inputs(months=13))
    titles = [page.title for page in spec.slides]
    appendix = [page for page in spec.slides if page.section == "Apêndice · tabelas"]
    assert appendix and {page.title for page in appendix} == {"Apêndice · séries"}
    # After the analysis, before the documents; each table says its theme.
    assert titles.index("Apêndice · séries") > titles.index("Rentabilidade")
    assert titles.index("Apêndice · séries") < titles.index("Documentos completos")
    themes = [block.title.split(" · ")[0] for page in appendix for block in page.blocks]
    order = ["Estrutura", "Atraso e provisão", "Migração e safras", "Rentabilidade"]
    assert themes == sorted(themes, key=order.index)
    charts = [page for page in spec.slides if page.title in {"Estrutura", "Atraso e provisão"}
              and any(hasattr(block, "series") for block in page.blocks)]
    assert charts and all(page.layout == "panels" for page in charts)


def test_inline_placement_keeps_series_after_each_theme():
    spec = _spec({**_inputs(months=13), "table_placement": "inline"})
    assert all(page.section != "Apêndice · tabelas" for page in spec.slides)
    assert any(page.title == "Estrutura" and page.layout == "flow" for page in spec.slides)


def test_reported_series_equal_to_derived_is_shown_once():
    kwargs = _inputs(months=13)
    kwargs["dashboard"] = _dashboard(subordination=[30.0] * 13)
    kwargs["table_placement"] = "inline"
    spec = _spec(kwargs)
    labels = [str(row.iloc[0]) for page in spec.slides for block in page.blocks
              if hasattr(block, "frame") and page.title == "Estrutura" for _, row in block.frame.iterrows()]
    assert "Subordinação ex-360" in labels and "Subordinação reportada" not in labels
    charts = [block for page in spec.slides for block in page.blocks if getattr(block, "title", "") == "Subordinação · %"]
    assert all(len(chart.series) == 1 for chart in charts if hasattr(chart, "series"))
    assert "coincide" in _notes_text(_build(kwargs))


def test_different_reported_series_shares_the_derived_chart():
    kwargs = _inputs(months=13)
    kwargs["dashboard"] = _dashboard(subordination=[25.0 + i / 10 for i in range(13)])
    chart = next(block for page in _spec(kwargs).slides for block in page.blocks
                 if getattr(block, "title", "") == "Subordinação · %")
    assert [series.label for series in chart.series] == ["Subordinação ex-360", "Subordinação reportada (IME)"]


def test_constant_series_become_a_sentence_not_a_chart():
    kwargs = _inputs(months=13)
    dashboard = _dashboard()
    dashboard.default_history_df["inadimplencia_pct"] = 1.0
    dashboard.default_history_df["provisao_pct_direitos"] = 0.5
    kwargs["dashboard"] = dashboard
    spec = _spec(kwargs)
    charts = [block.title for page in spec.slides for block in page.blocks if hasattr(block, "series")]
    assert "Vencidos e provisão / direitos creditórios (IME) · %" not in charts
    assert "Vencidos: 1,0% em todo o período." in _visible_text(_build(kwargs))


def test_ime_coverage_is_in_the_snapshot_not_a_chart():
    kwargs = _inputs(months=13)
    kwargs["dashboard"] = _dashboard()
    kwargs["dashboard"].summary["cobertura_pct"] = 100.0
    spec = _spec(kwargs)
    snapshot = next(page for page in spec.slides if page.title == "Snapshot").blocks[0]
    assert any("Cobertura provisão / vencidos · IME" in label for label in snapshot.frame["Indicador"])
    assert not any("Cobertura provisão / vencidos (IME)" in block.title for page in spec.slides for block in page.blocks if hasattr(block, "series"))


def test_snapshot_drops_comparison_columns_without_data_and_negative_zero():
    kwargs = _inputs(months=13)
    cnpj = kwargs["funds"][0].cnpj
    for frame in (kwargs["outputs"].fund_monthly[cnpj], kwargs["monitor_outputs"].fund_monitor[cnpj]):
        frame["carteira_ex360_yoy_pct"] = float("nan")
        frame["subordinacao_total_ex360_pct"] = [30.0] * 12 + [29.99]
    snapshot = next(page for page in _spec(kwargs).slides if page.title == "Snapshot").blocks[0]
    assert "Var. anual" not in snapshot.frame.columns
    assert not any("-0,0" in str(value) for value in snapshot.frame.to_numpy().ravel())


def _notes_text(payload):
    presentation = Presentation(BytesIO(payload))
    return "\n".join(slide.notes_slide.notes_text_frame.text for slide in presentation.slides)


def test_long_documentary_text_leaves_the_cell_and_keeps_full_clause_in_the_notes():
    payload = _build(_inputs(long_document=True))
    cells = [cell.text for table in _tables(payload) for row in table.rows for cell in row.cells]
    assert any("texto integral em C1 (notas do slide)" in cell for cell in cells)
    assert not any("Cláusulas integrais" in title for title in _titles(payload))
    assert not any("cláusula219" in cell for cell in cells) and "cláusula219" in _notes_text(payload)


def test_continued_table_says_cont_and_never_repeats_a_subtitle_on_a_slide():
    kwargs = _inputs(months=13, series=23)
    presentation = Presentation(BytesIO(_build(kwargs)))
    for slide in presentation.slides:
        subtitles = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.name == "Título"
                     and 1.3 * 914400 < shape.top < 6.6 * 914400]
        assert len(subtitles) == len(set(subtitles))


def test_spike_caps_the_axis_and_is_reported():
    from services.portfolio_ppt_export import _axis_cap

    values = [340, 314, 700, 247, 437, 276, 256, 247, 261, 262, 245, 246, 274]
    cap = _axis_cap(values)
    # Only the 700% spike is far out; 437% stays visible under the cap.
    assert cap is not None and 437 < cap < 700
    assert _axis_cap([1, 2, 3, 4, 5, 6, 7, 8, 9]) is None


def test_dividers_only_in_long_decks_and_before_sections_with_three_slides():
    from services.portfolio_ppt_export import DIVIDER_MIN_SLIDES, _section_dividers

    long_counts = [("", 1), ("Estrutura", 18), ("Atraso e provisão", 20), ("Rentabilidade", 2)]
    assert _section_dividers(long_counts, "auto") == {1: ("Estrutura", 18), 2: ("Atraso e provisão", 20)}
    short_counts = [("", 1), ("Estrutura", 5), ("Atraso e provisão", 5), ("Rentabilidade", 2)]
    assert 1 + sum(count for _, count in short_counts) <= DIVIDER_MIN_SLIDES
    assert _section_dividers(short_counts, "auto") == {}
    assert _section_dividers(short_counts, "always") == {1: ("Estrutura", 5), 2: ("Atraso e provisão", 5)}
    assert _section_dividers(long_counts, "never") == {}


def test_multi_fund_keeps_fund_detail_in_a_per_fund_annex():
    spec = _spec(_inputs(months=13, funds=3))
    sections = [page.section for page in spec.slides]
    assert "Anexo por fundo" in sections
    first_annex = sections.index("Anexo por fundo")
    assert sections.index("Referências e cobertura") > first_annex
    assert all(page.context.startswith("Consolidado") for page in spec.slides
               if page.title in {"Estrutura", "Atraso e provisão"})


def test_two_spikes_in_a_short_series_still_cap_the_axis():
    from services.portfolio_ppt_export import _axis_cap

    # A coverage ratio over a near-zero NPL: two far-out points in eight.
    values = [804_685.9, 1_146_011.0, 20_905.1, 23_791.8, 3_027.7, 3_165.0, 3_016.4, 1_736.8]
    cap = _axis_cap(values)
    assert cap is not None and 23_791.8 < cap < 804_685.9


def test_series_with_fewer_than_three_points_become_values_not_lines():
    from services.portfolio_ppt_content import ChartBlock, ChartSeries, _is_trivial, _trivial_sentence

    chart = ChartBlock("Vencidos · %", ("jan/26", "fev/26", "mar/26"),
                       (ChartSeries("a", "Vencidos", (0.0, 0.7, None)), ChartSeries("b", "Provisão", (None, 3.4, None))))
    assert _is_trivial(chart)
    text = _trivial_sentence(chart).text
    assert "Vencidos: 0,0% (jan/26) → 0,7% (fev/26)." in text
    assert "Provisão: 3,4% em fev/26 (1 observação)." in text


def test_four_integer_digits_drop_decimals_in_cells():
    from services.portfolio_ppt_export import format_cell

    assert format_cell(10_971.84, "number:2") == "10.972"
    assert format_cell(1_146_010.95, "percent:2") == "1.146.011%"
    assert format_cell(999.5, "number:2") == "999,50"
    assert format_cell(10_971_840_000.0, "money_value:2") == "10.972"


def test_curated_absence_with_a_reason_is_absence():
    from services.portfolio_ppt_content import _is_missing

    assert _is_missing("N/D — texto extraído sem associação confiável")
    assert _is_missing("Não localizado nas cláusulas revisadas; requer documento aplicável")
    assert _is_missing("Não informado nas fontes consultadas")
    assert not _is_missing("Não previsto")


def test_rows_and_columns_without_values_leave_series_tables():
    from services.portfolio_ppt_content import TableBlock, _prune_empty

    frame = pd.DataFrame({"Indicador": ["A", "B"], "jan/26": [1.0, None], "fev/26": [None, None]})
    block, said = _prune_empty(TableBlock("T · %", frame, numeric_columns=("jan/26", "fev/26")))
    assert list(block.frame.columns) == ["Indicador", "jan/26"] and list(block.frame["Indicador"]) == ["A"]
    assert any("B" in sentence for sentence in said) and any("fev/26" in sentence for sentence in said)


def test_base_warnings_are_one_line_per_warning_with_canonical_names():
    kwargs = _inputs(months=13, funds=2)
    cnpj = kwargs["funds"][0].cnpj
    kwargs["outputs"].warnings_df = pd.DataFrame({
        "scope_name": ["FUNDO 1 LEGAL", "FUNDO 1 LEGAL", "Carteira"], "cnpj": [cnpj, cnpj, "CONSOLIDADO"],
        "competencia": ["07/2025", "08/2025", "07/2025"],
        "warning": ["denominador zero/ausente", "denominador zero/ausente", "consolidado parcial na competência"]})
    cells = [cell.text for table in _tables(_build(kwargs)) for row in table.rows for cell in row.cells]
    assert "scope_name" not in cells and "warning" not in cells
    row = cells.index("Aviso da base · denominador zero/ausente")
    # Canonical fund name (not the legal one), months compressed in one cell.
    assert cells[row + 1].endswith(": jul/25, ago/25.") and "LEGAL" not in cells[row + 1]


def test_large_basket_summarizes_the_annex_in_one_row_per_fund():
    from services.portfolio_ppt_content import FUND_ANNEX_MAX_FUNDS

    funds = FUND_ANNEX_MAX_FUNDS + 1
    spec = _spec(_inputs(months=13, funds=funds))
    annex = [page for page in spec.slides if page.section == "Anexo por fundo"]
    assert len(annex) == 1 and len(annex[0].blocks[0].frame) == funds
    snapshot = next(page for page in spec.slides if page.title == "Snapshot")
    assert all(block.title != "Comparação entre fundos" for block in snapshot.blocks)
    full = _spec({**_inputs(months=13, funds=funds), "fund_annex": "full"})
    assert sum(page.section == "Anexo por fundo" for page in full.slides) == funds


def test_small_basket_annex_compares_funds_metric_by_metric():
    spec = _spec(_inputs(months=13, funds=3))
    annex = [page for page in spec.slides if page.section == "Anexo por fundo"]
    assert annex and all(page.layout == "grid" and page.title.startswith("Por fundo · ") for page in annex)
    names = [block.title for block in annex[0].blocks]
    assert len(names) == len(set(names)) == 3 and all([block.title for block in page.blocks] == names for page in annex)
    # Ratios share one scale; volumes keep each fund's own.
    ratio = next(page for page in annex if "NPL > 90d" in page.title)
    assert ratio.title.endswith("escala comum") and len({block.axis_range for block in ratio.blocks}) == 1
    volume = next(page for page in annex if page.title.startswith("Por fundo · PL total"))
    assert volume.title.endswith("escala própria") and all(block.axis_range == () for block in volume.blocks)
    # A metric constant in every fund is a sentence, not a page.
    assert not any("Subordinação ex-360" in page.title for page in annex)
    full = _spec({**_inputs(months=13, funds=3), "fund_annex": "full"})
    assert [page.title for page in full.slides if page.section == "Anexo por fundo"] == names


def test_fund_without_monthly_data_leaves_comparisons_and_is_named_in_coverage():
    kwargs = _inputs(months=13, funds=3)
    missing = kwargs["funds"][2].cnpj
    kwargs["outputs"].fund_monthly[missing] = pd.DataFrame()
    kwargs["monitor_outputs"].fund_monitor[missing] = pd.DataFrame()
    spec = _spec(kwargs)
    comparison = next(block for page in spec.slides for block in page.blocks
                      if getattr(block, "title", "") == "Comparação entre fundos")
    assert len(comparison.frame) == 2
    annex = [page for page in spec.slides if page.section == "Anexo por fundo"]
    assert annex and all(len(page.blocks) == 2 for page in annex)
    coverage = next(block for page in spec.slides for block in page.blocks
                    if getattr(block, "title", "") == "Fontes, cortes e cobertura").frame
    row = coverage[coverage["Referência"] == "Sem informe mensal"]
    assert len(row) == 1 and "2 de 3 fundos" in row["Conteúdo"].iloc[0]


def _sheet_item(cnpj):
    return SimpleNamespace(cnpj=cnpj, name="Fundo 1", competence="2026-08-01", classification="Multicedente",
                           pl_mm=512.0, receivables_pct_pl=0.96, senior_pct_pl=0.64, subordinated_pct_pl=0.36,
                           identified_issuance_volume_mm=None, first_issuance=None, last_issuance=None,
                           minimum_subordination="Sênior: 30% (mez.+jr.)/PL", rdb_status=None, rdb_finding=None,
                           rdb_source=None, validations=(), issuances=())


def test_snapshot_has_no_sparklines_and_sits_beside_the_fund_sheet():
    kwargs = _inputs(months=13)
    kwargs["analytical_data"] = (_sheet_item(kwargs["funds"][0].cnpj),)
    snapshot = next(page for page in _spec(kwargs).slides if page.title == "Snapshot")
    assert snapshot.layout == "panels" and [block.title for block in snapshot.blocks][1] == "Ficha do fundo"
    assert "Evolução" not in snapshot.blocks[0].frame.columns
    presentation = Presentation(BytesIO(_build(kwargs)))
    second = presentation.slides[1]
    assert not any(shape.has_chart for shape in second.shapes)
    tables = [shape for shape in second.shapes if shape.has_table]
    # Side by side: same top, the sheet to the right of the snapshot.
    assert len(tables) == 2 and tables[0].top == tables[1].top and tables[1].left > tables[0].left


def test_charts_take_half_the_slide_and_never_stretch_beyond_their_neighbour():
    from services.portfolio_ppt_content import ChartBlock, ChartSeries, SlideSpec, TableBlock
    from services.portfolio_ppt_export import _chart_height, _items, _pack, _place

    months = tuple(f"{month:02d}/25" for month in range(1, 13))
    line = lambda title, step: ChartBlock(title, months, (ChartSeries(title, title, tuple(1.0 + step * i for i in range(12))),))
    two = SlideSpec("Estrutura", "c", (line("PL", .1), line("Subordinação", .2)), layout="panels", section="Estrutura")
    placed = _place(_pack(list(_items(two, DEFAULT_STYLE)), DEFAULT_STYLE)[0], DEFAULT_STYLE)
    # Side by side, half the content height: they do not grow into free space.
    assert len(placed) == 2 and placed[0].y == placed[1].y and placed[0].width == placed[1].width
    assert placed[0].height == placed[1].height == _chart_height(DEFAULT_STYLE)
    short = TableBlock("Safras · %", pd.DataFrame({"Safra": ["jan/26", "fev/26"], "M1": [1.0, 2.0]}),
                       {"M1": "percent:1"}, numeric_columns=("M1",), compact=True)
    beside = SlideSpec("Migração", "c", (line("Roll", .1), short), layout="panels", section="Migração")
    chart, table = _place(_pack(list(_items(beside, DEFAULT_STYLE)), DEFAULT_STYLE)[0], DEFAULT_STYLE)
    # Beside a short table the chart keeps its own height, never more.
    assert chart.height == _chart_height(DEFAULT_STYLE) and table.height < chart.height and chart.y == table.y
    # Never a chart across the whole width.
    assert all(item.width == (DEFAULT_STYLE.content_width - DEFAULT_STYLE.gap) / 2 for item in (chart, *placed))


def test_small_themes_share_one_slide_named_after_both():
    from services.portfolio_ppt_content import ChartBlock, ChartSeries, SlideSpec, TextBlock
    from services.portfolio_ppt_export import _merge_small_themes, _groups

    sentence = lambda title: TextBlock(title, "Constante em 0,0% no período.")
    pages = (SlideSpec("Atraso e provisão", "c", (sentence("NPL"),), layout="panels", section="Atraso e provisão"),
             SlideSpec("Migração e safras", "c", (sentence("Roll"),), layout="panels", section="Migração e safras"))
    groups = _merge_small_themes(_groups(SimpleNamespace(slides=pages)), DEFAULT_STYLE)
    assert len(groups) == 1 and {item.group for item in groups[0]["items"] if item.block is not None} == {"Atraso e provisão", "Migração e safras"}


def test_quality_states_read_as_month_ranges_in_the_reference_table():
    from services.portfolio_ppt_content import _state_runs

    labels = ["set/25", "out/25", "nov/25", "dez/25", "jan/26"]
    states = ["Sem dados", "Sem dados", "Válido", "Não calculável", "Válido"]
    assert _state_runs(labels, states) == "Sem dados: set/25–out/25; Válido: nov/25, jan/26; Não calculável: dez/25"


def test_single_fund_references_have_no_peer_row_and_no_cover_repeats():
    spec = _spec(_inputs(months=13))
    frame = next(block for page in spec.slides for block in page.blocks
                 if getattr(block, "title", "") == "Fontes, cortes e cobertura").frame
    labels = set(frame["Referência"])
    assert not any("pares" in label.casefold() for label in labels)
    assert not labels & {"Consulta solicitada", "Período solicitado", "Histórico selecionado", "Consulta", "Fonte mensal"}


def test_axis_cap_note_keeps_the_decimals_a_small_cap_needs():
    from services.portfolio_ppt_content import ChartBlock, ChartSeries
    from services.portfolio_ppt_export import _chart_notes

    values = (0.1, 0.12, 0.11, 0.13, 0.1, 0.12, 0.11, 0.1, 2.4, 2.5)
    chart = ChartBlock("Cobertura · %", tuple(f"m{i}" for i in range(len(values))), (ChartSeries("cov", "Cobertura", values),), kind="column")
    note = next(line for line in _chart_notes(chart)[1] if line.startswith("Eixo limitado"))
    assert "a 0,15%" in note


def _varying_inputs(**kwargs):
    inputs = _inputs(months=13, **kwargs)
    for mapping in (inputs["outputs"].fund_monthly, inputs["monitor_outputs"].fund_monitor):
        for frame in mapping.values():
            frame["subordinacao_total_ex360_pct"] = [28.0 + month / 4 for month in range(13)]
    return inputs


def _charts_by_title(payload):
    presentation = Presentation(BytesIO(payload))
    return {shape.chart.chart_title.text_frame.text: shape for slide in presentation.slides
            for shape in slide.shapes if shape.has_chart and shape.chart.has_title}


def test_theme_charts_use_columns_and_lines_with_labels_on_every_point():
    from pptx.oxml.ns import qn

    charts = _charts_by_title(_build(_varying_inputs()))
    pl = charts["PL por classe · R$ milhões"].chart
    plot_tags = [element.tag for element in pl._chartSpace.xpath(".//c:plotArea/*")]
    # Stacked parts, plus an invisible line whose labels are the totals.
    assert qn("c:barChart") in plot_tags and qn("c:lineChart") in plot_tags
    assert pl._chartSpace.xpath(".//c:barChart/c:grouping/@val") == ["stacked"]
    assert charts["Subordinação · %"].chart._chartSpace.xpath(".//c:plotArea/c:lineChart")
    carteira = charts["Carteira bruta ex-360 · R$ milhões"].chart
    assert carteira._chartSpace.xpath(".//c:barChart/c:grouping/@val") == ["clustered"]
    # History in grey, latest month in orange.
    assert carteira._chartSpace.xpath(".//c:ser/c:spPr/a:solidFill/a:srgbClr/@val") == ["8B9199"]
    assert carteira._chartSpace.xpath(".//c:ser/c:dPt/c:spPr/a:solidFill/a:srgbClr/@val") == ["EC7000"]
    for title, shape in charts.items():
        space = shape.chart._chartSpace
        if len(shape.chart.plots[0].series) > 2 and space.xpath(".//c:plotArea/c:lineChart") and not space.xpath(".//c:barChart"):
            # More than two lines: each legend entry carries the latest value.
            assert all(": " in series.name for series in shape.chart.plots[0].series), title
            continue
        assert space.xpath(".//c:dLbls/c:showVal/@val"), title
        assert set(space.xpath(".//a:latin/@typeface")) == {"Calibri"}, title


def test_latest_period_label_is_bold_larger_and_filled_with_the_series_colour():
    charts = _charts_by_title(_build(_varying_inputs()))
    space = charts["Subordinação · %"].chart._chartSpace
    last = space.xpath(".//c:ser/c:dLbls/c:dLbl")[-1]
    assert last.xpath("./c:idx/@val") == ["12"]
    assert last.xpath("./c:spPr/a:solidFill/a:srgbClr/@val") == ["EC7000"]
    assert last.xpath("./c:txPr//a:defRPr/@b") == ["1"]
    others = space.xpath(".//c:ser/c:dLbls/c:txPr//a:defRPr/@sz")
    assert int(last.xpath("./c:txPr//a:defRPr/@sz")[0]) > int(others[0])


def test_no_chart_spans_the_whole_slide_width():
    presentation = Presentation(BytesIO(_build(_inputs(months=13, funds=3))))
    widths = [shape.width for slide in presentation.slides for shape in slide.shapes if shape.has_chart]
    assert widths and max(widths) <= Pt((DEFAULT_STYLE.content_width - DEFAULT_STYLE.gap) / 2) + 1
