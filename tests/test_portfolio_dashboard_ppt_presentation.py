from dataclasses import fields
from datetime import datetime, timezone
from io import BytesIO
from zipfile import ZipFile

import pandas as pd
from lxml import etree
from pptx import Presentation

from services.fund_return_disclosures import CVM_RETURN_REINVESTMENT_NOTE
from services.fundonet_dashboard import FundonetDashboardData
from services.fundonet_ppt_export import build_dashboard_pptx_bytes


GENERATED = datetime(2026, 10, 12, 16, 0, tzinfo=timezone.utc)
NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart", "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def _dashboard(*, returns=False):
    kwargs = {field.name: pd.DataFrame() for field in fields(FundonetDashboardData) if field.name.endswith("_df")}
    kwargs.update(
        competencias=["01/2026", "02/2026"],
        latest_competencia="02/2026",
        fund_info={"nome_fundo": "Carteira escolhida 1T26", "aggregation_scope": "portfolio", "periodo_analisado": "Jan/26 → Fev/26"},
        summary={"ativos_totais": 1_100_000_000.0, "direitos_creditorios": 950_000_000.0, "pl_total": 1_000_000_000.0, "inadimplencia_total": 0.0, "cobertura_pct": None, "subordinacao_pct": 10.0},
        methodology_notes=["Limite da fonte: dado reportado no Informe Mensal."],
        subordination_history_df=pd.DataFrame({"competencia": ["01/2026", "02/2026"], "subordinacao_pct": [11.0, 10.0]}),
        quota_pl_history_df=pd.DataFrame([
            {"competencia": month, "class_macro_label": label, "pl": value}
            for month in ["01/2026", "02/2026"]
            for label, value in [("Sênior", 900_000_000.0), ("Mezzanino", 75_000_000.0), ("Subordinada", 25_000_000.0)]
        ]),
    )
    if returns:
        kwargs["return_history_df"] = pd.DataFrame([
            {"competencia": month, "competencia_dt": pd.Timestamp(date), "label": "Sênior · Série 1", "retorno_mensal_pct": value}
            for month, date, value in [("01/2026", "2026-01-31", 0.0), ("02/2026", "2026-02-28", -0.2)]
        ])
        kwargs["return_summary_df"] = pd.DataFrame([{"label": "Sênior · Série 1", "retorno_ano_pct": -0.2, "retorno_periodo_pct": -0.2, "retorno_periodo_label": "2 meses"}])
    return FundonetDashboardData(**kwargs)


def _payload(dashboard, **kwargs):
    return build_dashboard_pptx_bytes(dashboard, generated_at=GENERATED, requested_period_label="12 meses", **kwargs)


def _chart_values_and_style(payload):
    with ZipFile(BytesIO(payload)) as archive:
        charts = {}
        for name in archive.namelist():
            if name.startswith("ppt/charts/chart") and name.endswith(".xml"):
                xml = etree.fromstring(archive.read(name))
                # Portfolio mode changes PL label policy and explicit font
                # sizes; datasets, scales, colors and native legend stay intact.
                for node in xml.xpath("//c:dLbls", namespaces=NS):
                    node.getparent().remove(node)
                for node in xml.xpath("//a:rPr | //a:defRPr | //a:endParaRPr", namespaces=NS):
                    node.attrib.pop("sz", None)
                charts[name] = etree.tostring(xml)
        return charts


def _cells_and_geometry(presentation):
    return [
        (slide_index, shape.left, shape.top, shape.width, shape.height, tuple(tuple(cell.text for cell in row.cells) for row in shape.table.rows))
        for slide_index, slide in enumerate(presentation.slides)
        for shape in slide.shapes
        if shape.has_table
    ]


def test_dashboard_default_is_unchanged_when_portfolio_option_is_false():
    dashboard = _dashboard(returns=True)
    default = _payload(dashboard)
    explicit = _payload(dashboard, portfolio_mode=False)
    with ZipFile(BytesIO(default)) as before, ZipFile(BytesIO(explicit)) as after:
        names = [name for name in before.namelist() if name.endswith(".xml") or name.endswith(".rels")]
        assert names == [name for name in after.namelist() if name.endswith(".xml") or name.endswith(".rels")]
        assert all(before.read(name) == after.read(name) for name in names)
    assert "COBERTURA DE PROVISÃO" in "\n".join(shape.text for slide in Presentation(BytesIO(default)).slides for shape in slide.shapes if shape.has_text_frame)


def test_portfolio_cards_use_readable_labels_keep_zero_and_explain_in_notes():
    dashboard = _dashboard()
    presentation = Presentation(BytesIO(_payload(dashboard, portfolio_mode=True)))
    slide = presentation.slides[1]
    shapes = [shape for shape in slide.shapes if shape.has_text_frame]
    text = "\n".join(shape.text for shape in shapes)
    for label in ["ATIVO", "DIREITOS CREDITÓRIOS", "PL", "VENCIDOS", "PDD / VENCIDOS", "SUBORDINAÇÃO"]:
        shape = next(shape for shape in shapes if shape.text == label)
        assert min(run.font.size.pt for paragraph in shape.text_frame.paragraphs for run in paragraph.runs) >= 10
        assert shape.height / 914400 >= 0.18
    assert "R$ 0,00" in text
    assert "N/D" in text
    assert "Posição mensal: fev-26" in text
    assert dashboard.fund_info["nome_fundo"] not in text
    notes = slide.notes_slide.notes_text_frame.text
    for expected in ["Subordinação reportada", "Mezzanino + subordinada residual / PL", "Provisão / créditos vencidos", "Janela solicitada: 12 meses", "Limite da fonte"]:
        assert expected in notes


def test_portfolio_pl_removes_rounded_internal_labels_preserves_native_series_and_total():
    dashboard = _dashboard()
    before = _payload(dashboard)
    after = _payload(dashboard, portfolio_mode=True)
    assert _chart_values_and_style(before) == _chart_values_and_style(after)
    old = Presentation(BytesIO(before))
    new = Presentation(BytesIO(after))
    assert len(old.slides) == len(new.slides) == 2
    assert [(shape.left, shape.top, shape.width, shape.height) for shape in old.slides[1].shapes if shape.has_chart] == [(shape.left, shape.top, shape.width, shape.height) for shape in new.slides[1].shapes if shape.has_chart]
    pl = next(shape.chart for shape in new.slides[1].shapes if shape.has_chart and shape.chart.chart_title.text_frame.text == "PL por tipo de cota")
    assert pl.has_legend
    assert not pl._chartSpace.xpath("//c:dLbl")
    assert "fev-26 · PL total: R$ 1.000,00 mm" in "\n".join(shape.text for shape in new.slides[1].shapes if shape.has_text_frame)
    assert "Compactação: rótulos internos" in new.slides[1].notes_slide.notes_text_frame.text


def test_portfolio_return_table_preserves_cells_and_moves_full_disclosure_to_notes():
    dashboard = _dashboard(returns=True)
    before = Presentation(BytesIO(_payload(dashboard)))
    after = Presentation(BytesIO(_payload(dashboard, portfolio_mode=True)))
    assert len(before.slides) == len(after.slides)
    assert _cells_and_geometry(before) == _cells_and_geometry(after)
    slide = after.slides[2]
    text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
    assert CVM_RETURN_REINVESTMENT_NOTE not in text
    assert "Metodologia de reinvestimento e ressalvas nas notas do slide." in text
    assert CVM_RETURN_REINVESTMENT_NOTE in slide.notes_slide.notes_text_frame.text
    table = next(shape.table for shape in slide.shapes if shape.has_table)
    assert all(run.font.size.pt >= 10 for row in table.rows for cell in row.cells for paragraph in cell.text_frame.paragraphs for run in paragraph.runs)


def test_portfolio_missing_subordination_keeps_visible_notice_and_note_definitions():
    dashboard = _dashboard()
    dashboard.subordination_history_df.drop(dashboard.subordination_history_df.index, inplace=True)
    presentation = Presentation(BytesIO(_payload(dashboard, portfolio_mode=True)))
    slide = presentation.slides[1]
    text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
    assert "histórico de subordinação sem dados nesta seleção" in text
    assert "Definição original — Subordinação reportada" in slide.notes_slide.notes_text_frame.text


def test_portfolio_native_chart_text_has_ten_point_floor_without_changing_scales_or_values():
    dashboard = _dashboard()
    months = pd.date_range("2025-02-28", periods=13, freq="ME")
    dashboard.default_aging_history_df["competencia"] = months.strftime("%m/%Y")
    dashboard.default_aging_history_df["faixa"] = "Até 30 dias"
    dashboard.default_aging_history_df["percentual_inadimplencia"] = 100.0
    before = _payload(dashboard)
    after = _payload(dashboard, portfolio_mode=True)
    assert _chart_values_and_style(before) == _chart_values_and_style(after)
    presentation = Presentation(BytesIO(after))
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not shape.has_chart:
                continue
            font_styles = shape.chart._chartSpace.xpath(".//a:rPr | .//a:defRPr | .//a:endParaRPr")
            assert font_styles
            assert all(int(style.get("sz")) >= 1000 for style in font_styles if style.get("sz") is not None)
    aging = next(shape.chart for slide in presentation.slides for shape in slide.shapes if shape.has_chart and shape.chart.chart_title.text_frame.text == "Aging")
    assert len(aging.plots[0].categories) == 13
    assert not aging._chartSpace.xpath("//c:dLbl")
    assert "Compactação: rótulos internos do Aging retirados" in "\n".join(slide.notes_slide.notes_text_frame.text for slide in presentation.slides if slide.has_notes_slide)


def test_portfolio_over_zero_label_is_explained_and_zero_series_remains_native():
    dashboard = _dashboard()
    dashboard.default_over_history_df["competencia"] = ["01/2026", "02/2026", "01/2026", "02/2026"]
    dashboard.default_over_history_df["serie"] = ["Over 90", "Over 90", "Over 360", "Over 360"]
    dashboard.default_over_history_df["percentual"] = [8.0, 7.0, 0.0, 0.0]
    dashboard.default_over_history_df["calculo_status"] = "calculado"
    before = _payload(dashboard)
    after = _payload(dashboard, portfolio_mode=True)
    assert _chart_values_and_style(before) == _chart_values_and_style(after)
    presentation = Presentation(BytesIO(after))
    slide = presentation.slides[2]
    chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
    zeros = next(series for series in chart.series if series.name == "Over 360")
    assert tuple(zeros.values) == (0.0, 0.0)
    assert chart.has_legend
    assert not zeros._element.xpath("./c:dLbls/c:dLbl")
    nonzero = next(series for series in chart.series if series.name == "Over 90")
    assert nonzero._element.xpath("./c:dLbls/c:dLbl")
    assert "Compactação: rótulos finais zerados" in slide.notes_slide.notes_text_frame.text
    assert "Over 360" in slide.notes_slide.notes_text_frame.text


def test_portfolio_over_small_positive_value_keeps_informative_label_instead_of_rounding_to_zero():
    dashboard = _dashboard()
    dashboard.default_over_history_df["competencia"] = ["01/2026", "02/2026", "01/2026", "02/2026"]
    dashboard.default_over_history_df["serie"] = ["Over 90", "Over 90", "Over 360", "Over 360"]
    dashboard.default_over_history_df["percentual"] = [8.0, 7.0, 0.023089636972648372, 0.00968490389584029]
    dashboard.default_over_history_df["calculo_status"] = "calculado"
    before = _payload(dashboard)
    after = _payload(dashboard, portfolio_mode=True)
    assert _chart_values_and_style(before) == _chart_values_and_style(after)
    presentation = Presentation(BytesIO(after))
    chart = next(shape.chart for shape in presentation.slides[2].shapes if shape.has_chart)
    series = next(series for series in chart.series if series.name == "Over 360")
    assert tuple(series.values)[-1] == 0.00968490389584029
    assert series._element.xpath("./c:dLbls/c:dLbl//a:t/text()") == ["0,01%"]
